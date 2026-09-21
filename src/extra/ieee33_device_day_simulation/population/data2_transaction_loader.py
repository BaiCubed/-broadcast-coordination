from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook

from .device_day_loader import DeviceDay, DeviceDayPool


FIELDS = (
    "transaction_id",
    "begin_time",
    "end_time",
    "total_charging_kwh",
    "current_soc",
    "out_power",
)

_POOL_CACHE: dict[tuple[Any, ...], DeviceDayPool] = {}


def clear_data2_transaction_pool_cache() -> None:
    _POOL_CACHE.clear()


def _pool_cache_key(workbook: Path, config: dict[str, Any]) -> tuple[Any, ...]:
    population = config["population"]
    return (
        str(workbook.resolve()),
        int(config["source_device_map"].get("assignment_seed", 0)),
        str(population.get("energy_input_source", "none")),
        str(population.get("load_source", "observed_out_power")),
        float(population.get("power_scale", 1.0)),
        float(population.get("energy_input_scale", 1.0)),
        float(population.get("fallback_soc_pct", 50.0)),
        float(population.get("fallback_capacity_kwh", 60.0)),
    )


def _as_datetime(value: Any) -> datetime:
    if isinstance(value, datetime):
        return value
    return datetime.fromisoformat(str(value))


def _assignment(index: int, config: dict[str, Any], rng: np.random.Generator) -> tuple[str, int]:
    mapping = config["source_device_map"]
    zones = config["zones"]["zones"]
    order = list(mapping["zone_order"])
    zone = order[index % len(order)]
    buses = list(zones[zone]["buses"])
    return zone, int(buses[int(rng.integers(0, len(buses)))])


def _finalize_session(
    rows: list[dict[str, Any]],
    config: dict[str, Any],
    index: int,
    rng: np.random.Generator,
) -> DeviceDay:
    steps = int(config["simulation"]["steps_per_day"])
    power_bins: list[list[float]] = [[] for _ in range(steps)]
    soc_bins: list[list[float]] = [[] for _ in range(steps)]
    for row in rows:
        end = _as_datetime(row["end_time"])
        slot = min(steps - 1, (end.hour * 3600 + end.minute * 60 + end.second) // 300)
        power_bins[slot].append(float(row["out_power"]))
        soc_bins[slot].append(float(row["current_soc"]))
    power_scale = float(config["population"].get("power_scale", 1.0))
    if power_scale <= 0:
        raise ValueError("data2 population.power_scale must be positive")
    power = np.asarray([float(np.mean(values)) if values else 0.0 for values in power_bins], dtype=float) * power_scale
    observed_soc = np.asarray([float(values[-1]) if values else np.nan for values in soc_bins], dtype=float)
    valid_soc = np.flatnonzero(np.isfinite(observed_soc))
    if len(valid_soc):
        observed_soc[: valid_soc[0]] = observed_soc[valid_soc[0]]
        for position in range(valid_soc[0] + 1, steps):
            if not np.isfinite(observed_soc[position]):
                observed_soc[position] = observed_soc[position - 1]
    else:
        observed_soc.fill(float(config["population"].get("fallback_soc_pct", 50.0)))

    first_soc = float(observed_soc[valid_soc[0]]) if len(valid_soc) else 50.0
    last_soc = float(observed_soc[valid_soc[-1]]) if len(valid_soc) else first_soc
    first_energy = float(rows[0]["total_charging_kwh"])
    last_energy = float(rows[-1]["total_charging_kwh"])
    delta_soc = last_soc - first_soc
    delta_energy = last_energy - first_energy
    if delta_soc > 3.0 and delta_energy > 1.0:
        inferred_capacity = delta_energy / (delta_soc / 100.0)
    else:
        inferred_capacity = float(config["population"].get("fallback_capacity_kwh", 60.0))
    capacity = float(np.clip(inferred_capacity, 20.0, 150.0))
    peak = float(np.clip(np.max(power), 1.0, 150.0))
    zone_id, bus_id = _assignment(index, config, rng)
    transaction_id = str(rows[0]["transaction_id"])
    timestamp = np.arange(steps, dtype=np.int64) * int(config["simulation"]["time_step_seconds"])
    input_source = str(config["population"].get("energy_input_source", "none"))
    load_source = str(config["population"].get("load_source", "observed_out_power"))
    input_scale = float(config["population"].get("energy_input_scale", 1.0))
    if input_scale <= 0:
        raise ValueError("data2 population.energy_input_scale must be positive")
    if input_source == "observed_out_power":
        energy_input = power.copy() * input_scale
    elif input_source == "none":
        energy_input = np.zeros(steps, dtype=float)
    else:
        raise ValueError(f"unsupported data2 energy_input_source: {input_source}")
    if load_source == "observed_out_power":
        load = power.copy()
    elif load_source == "zero":
        load = np.zeros(steps, dtype=float)
    else:
        raise ValueError(f"unsupported data2 load_source: {load_source}")
    return DeviceDay(
        device_id=f"data2_transaction_{index:05d}",
        source_device_id=transaction_id,
        day_id=0,
        zone_id=zone_id,
        bus_id=bus_id,
        load_kw=load,
        pv_kw=np.zeros(steps, dtype=float),
        energy_input_kw=energy_input,
        baseline_battery_kw=power,
        soc_kwh=np.clip(observed_soc / 100.0, 0.05, 0.98) * capacity,
        capacity_kwh=capacity,
        peak_power_kw=peak,
        c_rate=float(np.clip(peak / capacity, 0.05, 2.0)),
        timestamp=timestamp,
    )


def load_data2_transaction_pool(config: dict[str, Any]) -> DeviceDayPool:
    workbook = Path(config["population"]["data_workbook"])
    if not workbook.is_absolute():
        workbook = Path.cwd() / workbook
    if not workbook.exists():
        raise FileNotFoundError(workbook)
    cache_key = _pool_cache_key(workbook, config)
    if cache_key in _POOL_CACHE:
        return _POOL_CACHE[cache_key]
    selected: list[DeviceDay] = []
    seen: set[str] = set()
    duplicate_rows = 0
    rng = np.random.default_rng(int(config["source_device_map"].get("assignment_seed", 0)))
    wb = load_workbook(workbook, read_only=True, data_only=True)
    for worksheet in wb.worksheets:
        iterator = worksheet.iter_rows(values_only=True)
        header = next(iterator)
        columns = {str(value): index for index, value in enumerate(header)}
        missing = [field for field in FIELDS if field not in columns]
        if missing:
            raise ValueError(f"{worksheet.title} is missing data2 fields: {missing}")
        current_id: str | None = None
        session: list[dict[str, Any]] = []

        def flush() -> None:
            nonlocal session, current_id, duplicate_rows
            if not session or current_id is None:
                return
            if current_id in seen:
                duplicate_rows += len(session)
            else:
                seen.add(current_id)
                selected.append(_finalize_session(session, config, len(selected), rng))
            session = []

        for raw in iterator:
            transaction_id = str(raw[columns["transaction_id"]])
            if current_id is not None and transaction_id != current_id:
                flush()
            current_id = transaction_id
            session.append({field: raw[columns[field]] for field in FIELDS})
        flush()
    expected = config["population"].get("transaction_count_expected")
    if expected is not None and len(selected) != int(expected):
        raise ValueError(f"data2 transaction count changed: expected {expected}, found {len(selected)}")
    pool = DeviceDayPool(
        records=selected,
        metadata_count=len(selected),
        source_count=len(selected),
        filled_values=duplicate_rows,
    )
    _POOL_CACHE[cache_key] = pool
    return pool
