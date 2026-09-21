import os
import re

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'Supplementary_Tables_2-4.docx')
FONT = 'Times New Roman'
CAP_PT, BODY_PT, NOTE_PT = 12, 10, 9

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.top_margin = sec.bottom_margin = Cm(1.9)
TEXT_W = 17.0

st = doc.styles['Normal']
st.font.name = FONT
st.font.size = Pt(BODY_PT)
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

TOK = re.compile(r'(\*[^*]+\*|_\{[^}]+\}|\^\{[^}]+\})')


def add_runs(p, text, size, bold=False):
    for tok in TOK.split(text):
        if not tok:
            continue
        it = sub = sup = False
        if tok.startswith('*'):
            tok, it = tok[1:-1], True
        elif tok.startswith('_{'):
            tok, sub = tok[2:-1], True
        elif tok.startswith('^{'):
            tok, sup = tok[2:-1], True
        r = p.add_run(tok)
        r.font.name = FONT
        r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = it
        r.font.subscript = sub
        r.font.superscript = sup
        r.font.color.rgb = RGBColor(0, 0, 0)


def para(text, size, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0, spacing=1.5, keep=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = spacing
    pf.keep_with_next = keep
    add_runs(p, text, size, bold)
    return p


def set_borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{e}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '6')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '000000')
        b.append(el)
    tblPr.append(b)
    mar = OxmlElement('w:tblCellMar')
    for side, v in (('top', 28), ('bottom', 28), ('left', 70), ('right', 70)):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:w'), str(v))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tblPr.append(mar)


def table(header, rows, widths, merges=(), left_cols=()):
    ncol = len(header)
    t = doc.add_table(rows=1 + len(rows), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_borders(t)
    lay = OxmlElement('w:tblLayout')
    lay.set(qn('w:type'), 'fixed')
    t._tbl.tblPr.append(lay)
    for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(w / 2.54 * 1440)))
    grid = [header] + rows
    for r, line in enumerate(grid):
        row = t.rows[r]
        row.height = Cm(0.62)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement('w:cantSplit')
        cs.set(qn('w:val'), 'true')
        trPr.append(cs)
        if r == 0:
            h = OxmlElement('w:tblHeader')
            h.set(qn('w:val'), 'true')
            trPr.append(h)
        for c in range(ncol):
            cell = row.cells[c]
            cell.width = Cm(widths[c])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            txt = line[c]
            if txt is None:
                continue
            p = cell.paragraphs[0]
            pf = p.paragraph_format
            pf.alignment = (WD_ALIGN_PARAGRAPH.LEFT if (r > 0 and c in left_cols)
                            else WD_ALIGN_PARAGRAPH.CENTER)
            pf.space_before = pf.space_after = Pt(0)
            pf.line_spacing = 1.05
            add_runs(p, txt, BODY_PT, bold=(r == 0))
    for r0, c0, r1, c1 in merges:
        a = t.cell(r0 + 1, c0)
        a.merge(t.cell(r1 + 1, c1))
        a.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        ps = a.paragraphs
        for extra in ps[1:]:
            if not extra.text.strip():
                extra._element.getparent().remove(extra._element)
    return t


def caption(num, title):
    para(f'Supplementary Table {num}. {title}', CAP_PT, bold=True, before=0, after=6)


def part(label, text):
    p = para('', 11, before=10, after=4, spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_runs(p, label + ' ', 11, bold=True)
    add_runs(p, text, 11)


def note(text):
    para(text, NOTE_PT, before=4, after=0, spacing=1.15, keep=False)


NEFF = '*N*_{eff}'
NSTAR = '*N*^{*}_{eff}'

caption(2, 'Experimental design space and simulation coverage.')
H2 = ['Experiment\nmodule', 'Varied factor', 'Tested levels / range', 'Populations and replication', 'Purpose']
R2 = [
    ['Population scale\n(Result 1)', 'Physical population size *N*',
     '29 levels, 2–3,000; capped at the available source count of each population',
     '15 published + 119 mixed populations; 48 operating conditions × 130 repetitions (10 calibration, 120 held-out)',
     'Locate prediction and control transitions'],
    [None, 'Coupling arm', 'Data-coupled; decoupled',
     'Paired capacities, availability, controller draws and noise seeds', 'Isolate empirical cross-device dependence'],
    [None, 'Effective population size',
     f'{NEFF} from residual inter-device correlation; Γ = {NEFF}/{NSTAR}',
     'Estimated for every population–scale cell', 'Normalize population scale'],
    [None, 'Feeder representation', 'Reference feeder; IEEE-33; IEEE-69; IEEE-123',
     '119 mixed populations; 2,130 cells per feeder; 120 repetitions', f'Test topology dependence of {NSTAR}'],
    ['Response structure\n(Result 2)', 'Local response mechanism',
     'Open-loop; SOC-threshold; probabilistic; hybrid; price response; random delay',
     '15 published populations; *N* = 2–3,000; both coupling arms; 1,644 cells; 120 repetitions',
     f'Mechanism dependence of {NSTAR} and *β*'],
    [None, 'Broadcast waveform', 'Step; ramp; periodic; rapidly changing',
     '15 published populations; *N* = 6 and largest available size; both coupling arms; 4,176 cells; 24 repetitions',
     'Temporal compatibility'],
    [None, 'Response delay', 'Fixed; uniform; lognormal', None, 'Waveform–delay interaction'],
    [None, 'Controller homogeneity *c*', '0, 0.25, 0.50, 0.75, 0.90, 1.00', None, 'Collective response similarity'],
    [None, 'Broadcast period and duty cycle',
     'Period 15, 30, 60, 120 min; duty 0.25, 0.50, 0.75; *c* = 0, 0.9; *N* = 200, 600, 1,200',
     '14 published populations; 776 cells; 8 repetitions; 100 surrogates per window',
     'Phase coherence under common forcing'],
    ['Population evolution\n(Result 3)', 'Availability structure\n(primary family)',
     '5 structures; *p* = 0.30–1.00 (8 levels)', '15 published populations; 600 cells',
     'Effective scale under structured absence'],
    [None, 'Availability structure\n(second family)', '9 structures; *p* = 0.40–1.00 (5 levels)',
     '15 published (675 cells) + 119 mixed populations (5,355 cells)', f'Replicate the loss of {NEFF}'],
    [None, 'Controller drift', 'Abrupt, gradual; affected fraction 0.2, 0.5, 0.8',
     '14 published (672 runs) + 104 mixed populations (4,992 runs); 8 seeds', 'Detection and aggregate-only recovery'],
    ['Network realization\n(Result 4)', 'Feeder model', 'IEEE-33; IEEE-69; IEEE-123',
     '5,000 storage devices (5–20 kWh); 9 coordination and reference methods; dataset-level evaluation',
     'Physical deliverability'],
    [None, 'Network condition',
     'M0 nominal; M1 deep-feeder stress; M2 reverse-flow bottleneck; M3 derating and error; M4 uniform placement; '
     'M5 50% feeder concentration; M6 80% node concentration', None, 'Network stress and spatial placement'],
    [None, 'Constraint limits', 'Branch loading; transformer capacity; minimum and maximum voltage bounds (swept)',
     None, 'Separate response fidelity from feasibility'],
]
M2 = [(0, 0, 3, 0), (4, 0, 8, 0), (5, 3, 7, 3), (9, 0, 11, 0), (12, 0, 14, 0), (12, 3, 14, 3)]
table(H2, R2, [2.4, 2.8, 4.6, 4.6, 3.0], M2)
note('*N*, physical population size; ' + NEFF + ', effective population size; ' + NSTAR +
     ', effective control threshold; Γ, effective margin; *β*, fitted scaling exponent of conditional variability; '
     '*p*, participation rate. Mixed populations are composition stress tests, not independent empirical datasets.')

doc.add_page_break()

caption(3, 'Population evolution and perturbation protocols.')
part('a,', 'Availability perturbation')
H3a = ['Family', 'Availability structure', 'Temporal organization', 'Cross-device organization', 'Participation rate *p*']
R3a = [
    ['Primary family\n(Fig. 4a–c;\nSupplementary Fig. S14)', 'Independent', 'Independent draws', 'None',
     '0.30, 0.50, 0.70, 0.80, 0.85, 0.90, 0.95, 1.00'],
    [None, 'Persistent (Markov)', 'Persistent on/off states', 'None', None],
    [None, 'Uncoordinated empirical', 'Empirical trajectories', 'Alignment removed', None],
    [None, 'Neighbourhood-correlated', 'Shared absence episodes', 'Shared within neighbourhoods', None],
    [None, 'Empirically aligned', 'Empirical trajectories', 'Empirical alignment preserved', None],
    ['Second family\n(Supplementary\nFigs. S15, S16)', 'Independent absence', 'Independent draws', 'None',
     '0.40, 0.60, 0.80, 0.90, 1.00'],
    [None, 'Time-of-day, shifted', 'Diurnal schedule', 'Staggered across devices', None],
    [None, 'Time-of-day, class-specific', 'Diurnal schedule', 'Shared within device class', None],
    [None, 'Time-of-day, opposed', 'Diurnal schedule', 'Anti-phase between groups', None],
    [None, 'Shared outage schedules', 'Common outage schedule', '100, 20, 5 or 1 independent groups', None],
    [None, 'Weather-driven block', 'Contiguous absence blocks', 'Population-wide', None],
]
M3a = [(0, 0, 4, 0), (0, 4, 4, 4), (5, 0, 10, 0), (5, 4, 10, 4)]
table(H3a, R3a, [3.4, 3.8, 3.5, 3.9, 2.8], M3a)

part('b,', 'Controller drift')
H3b = ['Drift mode', 'Affected-controller fraction', 'Change profile', 'Runs per condition', 'Model-maintenance strategies']
R3b = [
    ['Abrupt', '0.2', 'Step change in local response parameters at the injection window',
     '112 published\n(14 × 8 seeds);\n832 mixed\n(104 × 8 seeds)',
     'Frozen pre-drift calibration; aggregate-only recalibration; full-information end-state retraining (reference)'],
    [None, '0.5', None, None, None],
    [None, '0.8', None, None, None],
    ['Gradual', '0.2', 'Progressive shift in local response parameters after injection', None, None],
    [None, '0.5', None, None, None],
    [None, '0.8', None, None, None],
]
M3b = [(0, 0, 2, 0), (0, 2, 2, 2), (3, 0, 5, 0), (3, 2, 5, 2), (0, 3, 5, 3), (0, 4, 5, 4)]
table(H3b, R3b, [2.4, 2.9, 4.3, 3.1, 4.7], M3b)
note('Drift was injected at window 28 of 86; each window contains 50 dispatch conditions (~33 h). The alarm threshold '
     'was the 99th percentile of window-level NRMSE on a drift-free validation stream and was frozen before evaluation. '
     'Outcomes: detection time *T*_{detect}, recovery time *T*_{recover}, checkpoints to 90% recovery *M*_{90} and '
     'cumulative excess loss. The full-information arm is a retraining reference, not a guaranteed upper bound.')

doc.add_page_break()

caption(4, 'Control strategies and information assumptions.')
part('a,', 'Local response mechanisms in the population-scale experiments (Result 2)')
H4a = ['Mechanism', 'Response type', 'Operator-side online information', 'Device-side information', 'Role in analysis']
R4a = [
    ['Open-loop', 'Direct response',
     'Common broadcast command only; no device-level uplink during routine dispatch',
     'Device power and energy limits', 'Direct-response reference; no response fraction defined'],
    ['SOC-threshold', 'Deterministic rule', None, 'Own SOC relative to local thresholds', 'Local heuristic'],
    ['Probabilistic', 'Randomized rule', None, 'Own state; local random draw', 'Randomized coordination'],
    ['Hybrid', 'Rule-based + randomized', None, 'Own SOC; local random draw', 'Combined practical strategy'],
    ['Price response', 'Economic response', None, 'Own state; local price sensitivity', 'Market-style response'],
    ['Random delay', 'Delayed response', None, 'Local random response delay', 'Timing-disorder stress test'],
]
table(H4a, R4a, [2.6, 3.0, 4.1, 4.1, 3.6], [(0, 2, 5, 2)])

part('b,', 'Coordination and reference methods in the network evaluation (Result 4)')
H4b = ['Method', 'Control type', 'Information used during routine dispatch', 'Online complexity', 'Role']
R4b = [
    ['No coordination', 'None', 'None', '*O*(0)', 'Lower reference for curtailment reduction'],
    ['Local SOC rules', 'Local heuristic', 'Own device SOC only; no operator messages', '*O*(1) per device',
     'Decentralized baseline'],
    ['Model predictive control', 'Receding-horizon optimization', 'Device states and forecasts over horizon *H*',
     '*O*(*HN*)', 'Optimization baseline'],
    ['Mean-field control', 'Distribution-level coordination', 'Population state distribution from device reports',
     '*O*(*N*)', 'Mean-field baseline'],
    ['Virtual battery', 'Aggregate state-space model', 'Device SOC and power limits aggregated into a virtual battery',
     '*O*(*N*)', 'Aggregate-model baseline'],
    ['Packetized energy management', 'Request–grant protocol', 'Device energy-packet requests',
     '*O*(*N* log *N*)', 'Asynchronous coordination baseline'],
    ['Transactive control', 'Market clearing', 'Device bids', '*O*(*N* log *N*)', 'Market-based baseline'],
    ['EPS', 'Broadcast coordination',
     'Common broadcast score and aggregate measurements; device parameters used only at registration and offline calibration',
     '*O*(1)', 'Reduced-observability architecture'],
    ['Centralized greedy', 'Device-level dispatch', 'Full device-level states (SOC, availability, power limits)',
     '*O*(*N*)', 'Reference implementation, not a strict upper bound'],
]
table(H4b, R4b, [3.0, 3.1, 5.7, 2.3, 3.3])
note('Complexity refers to operator-side routine online coordination with respect to population size *N*; '
     '*H*, prediction horizon. Registration, commissioning, offline calibration and occasional model maintenance are '
     'excluded. Methods request different response magnitudes, so the comparison characterizes information assumptions '
     'rather than algorithmic superiority. The same network-feasibility procedure was applied where required for like-for-like network comparisons.')

doc.save(OUT)
print('saved', OUT)
