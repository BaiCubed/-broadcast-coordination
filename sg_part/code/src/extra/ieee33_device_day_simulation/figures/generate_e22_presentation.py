"""生成E22图片汇总演示文稿。"""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[4]
E22 = ROOT / "results" / "E22"
OUTPUT = E22 / "E22_figures.pptx"

NAVY = RGBColor(31, 50, 72)
TEAL = RGBColor(20, 133, 130)
ORANGE = RGBColor(216, 103, 62)
TEXT = RGBColor(42, 49, 56)
MUTED = RGBColor(96, 105, 115)
PALE = RGBColor(239, 244, 245)
WHITE = RGBColor(255, 255, 255)


def add_text(slide, left, top, width, height, text, size, color=TEXT, bold=False):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.06)
    frame.margin_right = Inches(0.06)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Microsoft YaHei"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return box


def add_image(slide, path, left, top, width, height):
    with Image.open(path) as source:
        ratio = source.width / source.height
    area_ratio = width / height
    if ratio >= area_ratio:
        image_width = width
        image_height = width / ratio
        image_left = left
        image_top = top + (height - image_height) / 2
    else:
        image_height = height
        image_width = height * ratio
        image_left = left + (width - image_width) / 2
        image_top = top
    slide.shapes.add_picture(
        str(path), Inches(image_left), Inches(image_top), Inches(image_width), Inches(image_height)
    )


def add_slide_header(slide, title, section):
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.72)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = NAVY
    header.line.fill.background()
    add_text(slide, 0.45, 0.12, 10.9, 0.42, title, 21, WHITE, True)
    add_text(slide, 11.4, 0.16, 1.45, 0.3, section, 12, WHITE, True)


def add_note(slide, title, body, accent=TEAL):
    band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(5.52), Inches(12.43), Inches(1.45)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = PALE
    band.line.color.rgb = RGBColor(208, 220, 222)
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(5.52), Inches(0.08), Inches(1.45)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    add_text(slide, 0.68, 5.68, 2.2, 0.28, title, 11.5, accent, True)
    add_text(slide, 2.68, 5.62, 9.85, 1.18, body, 9.2, TEXT)


def figure_info(path):
    relative = path.relative_to(E22).as_posix()
    name = path.stem
    if path.parent.name == "by_dataset":
        dataset = name.removesuffix("_r2_vs_n").replace("_", " ")
        return (
            f"R²随设备规模变化：{dataset}",
            "单个数据集的规模实验。横坐标为设备数量N，纵坐标为聚合响应的R²；用于观察达到高拟合度所需的设备规模。",
            "R²曲线",
        )
    if name == "e22_completed_r2_group":
        return (
            "五个数据集的R²规模审计",
            "该图仅包含已完成R²规模扫描的BDG1、BDG2、CEC、Danish和EU-Rural，不代表全部17个数据集。横坐标为物理设备数量N，纵坐标为独立测试R²；红线为0.95，绿线为对数插值得到的N95。它回答聚合预测需要多大规模，不直接回答Reviewer #4的配网安全问题。",
            "规模审计",
        )
    main = {
        "topology_effect_dumbbell": (
            "IEEE-33与IEEE-69拓扑效果对比",
            "Reviewer回应：设备、数据、场景和seed成对固定，只将馈线从IEEE-33改为IEEE-69。横坐标为网络约束后的可交付弃电降低率，连线长度表示拓扑变化造成的效果差异。Centralized greedy UB未经过网络裁剪，只作设备层参考。",
            "网络复杂度验证",
        ),
        "ieee69_effect_safety_pareto": (
            "IEEE-69效果与原始请求安全性",
            "Reviewer回应：每个子图对应M0-M6。横坐标是约束后效果，纵坐标是在安全层介入前，算法原始请求不触发线路、电压或主变违规的时间比例；右上角才表示效果高且原始请求较安全。Greedy横坐标仅为设备层理论值。",
            "安全性分离验证",
        ),
        "retention_by_dataset_algorithm": (
            "不同数据集的拓扑效果保持率",
            "Reviewer回应：逐数据集计算IEEE-69效果除以IEEE-33效果，检验结论是否只来自少数数据集。横坐标为算法，纵坐标为17个数据集，颜色为七场景平均保持率。它支持跨数据集稳健性，不单独证明约束满足；Greedy仅作设备层参考。",
            "跨数据集稳健性",
        ),
        "ieee69_stress_profiles": (
            "IEEE-69网络与空间压力响应",
            "Reviewer回应：M0-M3加入深馈线高负载、双瓶颈反向潮流、线路降额和预测误差；M4-M6固定网络参数，只改变设备空间分布。曲线下降说明聚合层效果不能脱离网络压力和空间位置讨论。",
            "压力场景验证",
        ),
        "ieee69_stress_profiles_comparison": (
            "IEEE-69全部算法压力场景对比",
            "一张图统一比较14种算法，包括传统基线、E20/E21迁移EPS、EPS in-domain和三种IEEE-69直接训练消融。横坐标都是M0-M6，纵坐标都是17个数据集平均可交付弃电降低率；曲线从M0到M3的变化反映网络压力，从M4到M6的变化反映空间集中。",
            "单图压力对比",
        ),
        "ieee69_network_acceptance": (
            "IEEE-69网络可接受响应比例",
            "Reviewer回应：网络接受率是通过线路热限、0.95-1.05 p.u.电压范围和6 MVA主变容量校核后的响应，除以算法原始请求。低值表示广播请求与设备空间位置或网络余量不匹配；Greedy行是其理论请求中通过网络校核的比例。",
            "约束可交付性",
        ),
        "ieee69_branch_time_heatmap": (
            "IEEE-69支路约束的时空位置",
            "Reviewer回应：选取CEC、M2双瓶颈、E21 mixed EPS和第一个seed，显示约束具体发生在何时、哪条支路。横坐标为288个五分钟步，纵坐标为68条支路，1.0为热限。它是馈线热限的单案例证据，不单独证明电压和主变安全。",
            "馈线热限案例",
        ),
        "spatial_concentration_retention": (
            "设备空间集中造成的效果损失",
            "Reviewer回应：M4-M6固定负荷、能源输入、设备参数、可用率、控制随机数和网络参数，只改变设备母线位置。低于1表示5000台设备集中在社区馈线、充电站或工业园后形成局部瓶颈；Greedy仅作设备层参考。",
            "空间集中验证",
        ),
        "trained_all_algorithms_comparison": (
            "IEEE-69全部算法平均效果对比",
            "一张横向柱状图统一比较14种算法，横坐标是IEEE-69上跨17个数据集和M0-M6的平均可交付弃电降低率，纵坐标给出算法和复杂度。EPS方法保持O(1)在线广播，Centralized greedy UB是O(N)设备层理论上界，不是网络可执行结果。",
            "单图全算法对比",
        ),
        "ieee69_constraint_component_audit": (
            "IEEE-69约束分项审计",
            "Reviewer直接证据：四个面板汇总17个数据集和30个paired seeds，分别显示执行后最大支路负载率、最低节点电压、最大主变负载率，以及安全层使无违规时间比例提升的百分点。阈值为支路1.0、电压0.95 p.u.、主变1.0。它仍是LinDistFlow标准馈线仿真，不是现场安全认证。",
            "Reviewer #4直接回应",
        ),
    }
    if name in main:
        return main[name]
    if name.startswith("trained_") or "trained_" in name:
        is_direct = "trained_eps_ieee69_direct" in relative
        experiment = "直接训练消融" if is_direct else "迁移与域内训练"
        trained = {
            "trained_vs_transfer_dataset_heatmap_ieee69": "按数据集比较训练方式",
            "trained_vs_transfer_stress_profiles_ieee69": "按压力场景比较训练方式",
        }
        if name == "trained_eps_all_algorithms_ieee69":
            title = "IEEE-69直接训练全算法对比" if is_direct else "IEEE-69迁移与域内训练全算法对比"
        elif name == "trained_gain_by_dataset_ieee69":
            title = "直接训练相对迁移的增益" if is_direct else "域内训练相对迁移的增益"
        else:
            title = trained.get(name, name.replace("_", " "))
        if is_direct and name == "trained_eps_all_algorithms_ieee69":
            body = (
                "该图按跨17个数据集和M0-M6的平均可交付弃电降低率比较全部算法。EPS部分包含五种含义不同的控制："
                "E20 pooled EPS使用跨数据集 pooled 模型，E21 mixed EPS使用混合场景模型；"
                "EPS formula-only只使用解析公式，EPS learned-only只使用IEEE-69直接训练模型，"
                "EPS IEEE-69 fused融合解析公式、学习模型和闭环修正。后三种是直接训练消融，在线通信复杂度均为O(1)。"
                "图中其他算法作为基线，Centralized greedy UB是O(N)理论设备层上界。"
            )
        elif is_direct and name == "trained_vs_transfer_dataset_heatmap_ieee69":
            body = (
                "逐数据集比较E20 pooled迁移、E21 mixed迁移，以及IEEE-69直接训练后的formula-only、learned-only和fused控制；颜色和数字为M0-M6平均可交付弃电降低率。该图回答训练分布是否影响效果，不是独立的网络安全证据。"
            )
        elif is_direct and name == "trained_vs_transfer_stress_profiles_ieee69":
            body = (
                "比较两种迁移EPS与三种IEEE-69直接训练消融在M0-M6下的效果。横坐标为压力场景，纵坐标为17个数据集平均效果。它判断直接训练是否改善压力场景适配；网络安全结论仍应结合接受率和约束分项审计。"
            )
        elif is_direct and name == "trained_gain_by_dataset_ieee69":
            body = (
                "横坐标为IEEE-69直接训练fused EPS相对E21 mixed迁移模型的效果差值，纵坐标为数据集；正值表示直接训练更好，负值表示迁移模型更好。该图检验训练收益差异，不证明模型原始请求天然满足配网约束。"
            )
        elif not is_direct and name == "trained_eps_all_algorithms_ieee69":
            body = (
                "比较固定基线、E20/E21迁移EPS、目标数据域内训练EPS和Centralized greedy UB在IEEE-69上的平均效果。EPS均保持O(1)在线广播；Greedy为O(N)设备层理论上界。该图比较效果和信息复杂度，不单独回答线路、电压和主变安全问题。"
            )
        elif not is_direct and name == "trained_vs_transfer_dataset_heatmap_ieee69":
            body = (
                "逐数据集比较E20 pooled、E21 mixed迁移和目标数据域内训练EPS，颜色为M0-M6平均可交付弃电降低率。它回答迁移是否损失数据集适配能力；所有方法使用同一IEEE-69测试协议，网络约束不是该图的自变量。"
            )
        elif not is_direct and name == "trained_vs_transfer_stress_profiles_ieee69":
            body = (
                "比较迁移EPS与域内训练EPS在M0-M6下的效果。横坐标为压力场景，纵坐标为17数据集平均可交付弃电降低率。是否安全仍要查看原始请求无违规比例、网络接受率和约束分项审计。"
            )
        elif not is_direct and name == "trained_gain_by_dataset_ieee69":
            body = (
                "展示域内训练EPS相对E21 mixed迁移EPS的逐数据集效果增益。正值表示域内训练改善，负值表示迁移模型更好。该图说明训练收益不是所有数据集一致，也不等价于现场可实施性或网络安全提升。"
            )
        else:
            body = (
                f"IEEE-69上的{experiment}结果。图中比较迁移模型与目标数据训练模型；具体横纵坐标随图中标注，重点观察训练方式对可交付弃电降低率的影响。"
            )
        return title, body, f"{experiment}对比"
    return name.replace("_", " "), f"E22实验图片。图片文件：{relative}。横纵坐标和颜色含义以图中标注为准。", "E22"


def create_presentation():
    excluded = {
        "figures/ieee69_stress_profiles.png",
        "trained_eps_ieee69/figures/trained_vs_transfer_stress_profiles_ieee69.png",
        "trained_eps_ieee69_direct/figures/trained_vs_transfer_stress_profiles_ieee69.png",
        "trained_eps_ieee69/figures/trained_vs_transfer_dataset_heatmap_ieee69.png",
        "trained_eps_ieee69/figures/trained_eps_all_algorithms_ieee69.png",
        "trained_eps_ieee69_direct/figures/trained_eps_all_algorithms_ieee69.png",
    }
    images = sorted(
        path
        for path in E22.rglob("*.png")
        if path.parent.name != "by_dataset"
        and path.relative_to(E22).as_posix() not in excluded
    )
    if not images:
        raise FileNotFoundError(f"未找到E22图片：{E22}")
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = "E22实验图片汇总"
    presentation.core_properties.subject = "IEEE-33/IEEE-69空间分布与配网约束验证"
    presentation.core_properties.author = "Codex"

    for index, image in enumerate(images, start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = WHITE
        title, body, section = figure_info(image)
        add_slide_header(slide, title, section)
        add_image(slide, image, 0.45, 0.86, 12.43, 4.52)
        add_note(slide, "实验回答", f"{body} 来源：{image.relative_to(E22).as_posix()}。")
        add_text(slide, 11.85, 7.08, 0.95, 0.2, f"{index}/{len(images)}", 9, MUTED, False)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(OUTPUT)
    print(f"已生成：{OUTPUT}")
    print(f"页面数：{len(images)}")


if __name__ == "__main__":
    create_presentation()
