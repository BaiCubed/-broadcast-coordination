"""从IEEE-69直接训练结果生成E23严格失效边界实验。"""

from __future__ import annotations

import csv
import json
import shutil
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from tools.ncstyle import SEQ
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches as PptInches


ROOT = Path(__file__).resolve().parents[4]
SOURCE = ROOT / "results/E22/trained_eps_ieee69_direct"
OUTPUT = ROOT / "results/E23"
OVERALL_SOURCE = SOURCE / "data/combined_e22_overall.csv"
SUMMARY_SOURCE = SOURCE / "data/combined_e22_summary.csv"

SCENARIO_ORDER = ("M0", "M1", "M2", "M3", "M4", "M5", "M6")
ALGORITHM_ORDER = (
    "no_coordination",
    "local_rules",
    "mpc_optimal",
    "mean_field_control",
    "virtual_battery",
    "packetized_energy_management",
    "transactive_control",
    "eps_ieee69_fused",
    "centralized_optimal",
)
ALGORITHM_COLORS = {
    "no_coordination": "#9c9c9c",
    "local_rules": "#ed7d31",
    "mpc_optimal": "#7057ff",
    "mean_field_control": "#4e79a7",
    "virtual_battery": "#59a14f",
    "packetized_energy_management": "#af7aa1",
    "transactive_control": "#d55e00",
    "eps_ieee69_fused": "#1f5aa6",
    "centralized_optimal": "#3a9d66",
}
ALGORITHM_NAMES: dict[str, str] = {}

matplotlib.rcParams["font.family"] = "SimSun"
matplotlib.rcParams["axes.unicode_minus"] = False

FIGURES = (
    ("ieee69_effect_under_stress", "IEEE-69严格场景下的算法效果"),
    ("ieee69_failure_boundary_dashboard", "IEEE-69失效边界总览"),
    ("ieee69_constraint_margins", "IEEE-69配网约束余量"),
    ("ieee69_mean_effect", "IEEE-69算法平均效果"),
)


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        raise ValueError(f"没有可写入的数据：{path}")
    fields = list(rows[0])
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def _filter_rows(path: Path) -> list[dict[str, str]]:
    rows = _read_csv(path)
    return [
        row
        for row in rows
        if row.get("topology") == "ieee69" and row.get("algorithm") in ALGORITHM_ORDER
    ]


def _prepare_data() -> tuple[list[dict[str, str]], list[dict[str, str]], list[dict[str, Any]]]:
    overall = _filter_rows(OVERALL_SOURCE)
    summary = _filter_rows(SUMMARY_SOURCE)
    if len(overall) != len(ALGORITHM_ORDER) * len(SCENARIO_ORDER):
        raise ValueError(f"IEEE-69总体结果数量异常：{len(overall)}")
    if not summary:
        raise ValueError("IEEE-69逐数据集结果为空")

    boundary_rows: list[dict[str, Any]] = []
    for algorithm in ALGORITHM_ORDER:
        for stress_mode in SCENARIO_ORDER:
            selected = [
                row
                for row in summary
                if row["algorithm"] == algorithm and row["stress_mode"] == stress_mode
            ]
            overall_row = next(
                row
                for row in overall
                if row["algorithm"] == algorithm and row["stress_mode"] == stress_mode
            )
            mean_branch = float(np.mean([float(row["maximum_branch_loading_mean"]) for row in selected]))
            mean_voltage = float(np.mean([float(row["minimum_voltage_pu_mean"]) for row in selected]))
            mean_transformer = float(
                np.mean([float(row["maximum_transformer_loading_mean"]) for row in selected])
            )
            boundary_rows.append(
                {
                    "algorithm": algorithm,
                    "algorithm_label": overall_row["algorithm_label"],
                    "complexity": overall_row["complexity"],
                    "stress_mode": stress_mode,
                    "curtailment_reduction_pct": overall_row["curtailment_reduction_pct"],
                    "network_acceptance_pct": 100.0 * float(overall_row["network_acceptance_ratio"]),
                    "requested_violation_free_pct": overall_row["requested_violation_free_pct"],
                    "executed_violation_free_pct": overall_row["executed_violation_free_pct"],
                    "mean_branch_loading": mean_branch,
                    "branch_margin_to_limit": 1.0 - mean_branch,
                    "mean_minimum_voltage_pu": mean_voltage,
                    "voltage_margin_to_limit": mean_voltage - 0.95,
                    "mean_transformer_loading": mean_transformer,
                    "transformer_margin_to_limit": 1.0 - mean_transformer,
                    "raw_request_below_95pct_safe": float(overall_row["requested_violation_free_pct"]) < 95.0,
                    "physical_boundary_reached": (
                        mean_branch > 1.0 or mean_voltage < 0.95 or mean_transformer > 1.0
                    ),
                }
            )
    return overall, summary, boundary_rows


def _save_figure(figure: Any, name: str) -> None:
    output = OUTPUT / "figures" / f"{name}.png"
    output.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(output, dpi=220, bbox_inches="tight")
    figure.savefig(output.with_suffix(".pdf"), bbox_inches="tight")
    plt.close(figure)


def _labels(overall: list[dict[str, str]]) -> dict[str, tuple[str, str]]:
    return {
        algorithm: next(
            (row["algorithm_label"], row["complexity"])
            for row in overall
            if row["algorithm"] == algorithm
        )
        for algorithm in ALGORITHM_ORDER
    }


def _matrix(rows: list[dict[str, Any]], field: str, scale: float = 1.0) -> np.ndarray:
    return np.asarray(
        [
            [
                scale
                * float(
                    next(
                        row[field]
                        for row in rows
                        if row["algorithm"] == algorithm and row["stress_mode"] == stress_mode
                    )
                )
                for stress_mode in SCENARIO_ORDER
            ]
            for algorithm in ALGORITHM_ORDER
        ]
    )


def _plot_effect(overall: list[dict[str, str]]) -> None:
    labels = _labels(overall)
    figure, axis = plt.subplots(figsize=(14.5, 7.2))
    x = np.arange(len(SCENARIO_ORDER))
    markers = ("o", "s", "^", "D", "v", "P", "X", "*", "h")
    for index, algorithm in enumerate(ALGORITHM_ORDER):
        values = [
            float(
                next(
                    row["curtailment_reduction_pct"]
                    for row in overall
                    if row["algorithm"] == algorithm and row["stress_mode"] == stress_mode
                )
            )
            for stress_mode in SCENARIO_ORDER
        ]
        axis.plot(
            x,
            values,
            marker=markers[index],
            linewidth=2.2,
            markersize=6,
            color=ALGORITHM_COLORS[algorithm],
            label=f"{labels[algorithm][0]} ({labels[algorithm][1]})",
        )
    axis.set_xticks(x, SCENARIO_ORDER)
    axis.set_ylim(0, 110)
    axis.set_xlabel("IEEE-69网络与设备空间压力场景")
    axis.set_ylabel("可交付弃电降低率 (%)")
    axis.set_title("IEEE-69严格条件下的算法效果边界")
    axis.grid(axis="y", alpha=0.25)
    axis.legend(loc="upper center", bbox_to_anchor=(0.5, -0.12), ncol=3, fontsize=8)
    figure.tight_layout()
    _save_figure(figure, "ieee69_effect_under_stress")


def _heatmap(
    figure: Any,
    axis: Any,
    values: np.ndarray,
    title: str,
    colorbar_label: str,
    vmin: float,
    vmax: float,
    threshold: float | None = None,
) -> None:
    image = axis.imshow(values, aspect="auto", cmap=SEQ, vmin=vmin, vmax=vmax)
    axis.set_title(title, fontsize=12)
    axis.set_xticks(np.arange(len(SCENARIO_ORDER)), SCENARIO_ORDER)
    axis.set_yticks(np.arange(len(ALGORITHM_ORDER)), [ALGORITHM_NAMES[a] for a in ALGORITHM_ORDER])
    for row_index in range(values.shape[0]):
        for column_index in range(values.shape[1]):
            value = values[row_index, column_index]
            axis.text(
                column_index,
                row_index,
                f"{value:.1f}",
                ha="center",
                va="center",
                fontsize=7,
                color="white" if value > (vmax - vmin) * 0.58 + vmin else "black",
            )
    if threshold is not None:
        axis.contour(values, levels=[threshold], colors="#d1495b", linewidths=1.6)
    colorbar = figure.colorbar(image, ax=axis, pad=0.015, fraction=0.035)
    colorbar.set_label(colorbar_label)


def _plot_boundary_dashboard(boundary_rows: list[dict[str, Any]]) -> None:
    figure, axes = plt.subplots(3, 1, figsize=(14.0, 13.0), sharex=True)
    _heatmap(
        figure,
        axes[0],
        _matrix(boundary_rows, "curtailment_reduction_pct"),
        "实际可交付效果",
        "弃电降低率 (%)",
        0.0,
        100.0,
    )
    _heatmap(
        figure,
        axes[1],
        _matrix(boundary_rows, "network_acceptance_pct"),
        "网络接受率",
        "接受率 (%)",
        0.0,
        100.0,
        threshold=95.0,
    )
    _heatmap(
        figure,
        axes[2],
        _matrix(boundary_rows, "requested_violation_free_pct"),
        "安全层介入前的原始请求无违规比例",
        "无违规比例 (%)",
        0.0,
        100.0,
        threshold=95.0,
    )
    axes[-1].set_xlabel("IEEE-69严格压力场景；红色等值线为95%参考线")
    figure.suptitle("E23：IEEE-69算法失效边界总览", y=0.995, fontsize=16)
    figure.tight_layout(rect=(0, 0, 1, 0.985))
    _save_figure(figure, "ieee69_failure_boundary_dashboard")


def _plot_constraint_margins(boundary_rows: list[dict[str, Any]]) -> None:
    figure, axes = plt.subplots(3, 1, figsize=(14.0, 13.0), sharex=True)
    _heatmap(
        figure,
        axes[0],
        _matrix(boundary_rows, "mean_branch_loading"),
        "平均最大支路负载率",
        "负载率；1.0为热限",
        0.0,
        2.0,
        threshold=1.0,
    )
    _heatmap(
        figure,
        axes[1],
        _matrix(boundary_rows, "mean_minimum_voltage_pu"),
        "平均最低节点电压",
        "电压 (p.u.)；0.95为下限",
        0.90,
        1.05,
        threshold=0.95,
    )
    _heatmap(
        figure,
        axes[2],
        _matrix(boundary_rows, "mean_transformer_loading"),
        "平均最大主变负载率",
        "负载率；1.0为容量上限",
        0.0,
        1.2,
        threshold=1.0,
    )
    axes[-1].set_xlabel("IEEE-69严格压力场景；红色等值线为物理约束阈值")
    figure.suptitle("E23：配网物理约束的失效边界", y=0.995, fontsize=16)
    figure.tight_layout(rect=(0, 0, 1, 0.985))
    _save_figure(figure, "ieee69_constraint_margins")


def _plot_mean_effect(overall: list[dict[str, str]]) -> None:
    labels = _labels(overall)
    values = []
    for algorithm in ALGORITHM_ORDER:
        values.append(
            float(
                np.mean(
                    [
                        float(row["curtailment_reduction_pct"])
                        for row in overall
                        if row["algorithm"] == algorithm
                    ]
                )
            )
        )
    y = np.arange(len(ALGORITHM_ORDER))
    figure, axis = plt.subplots(figsize=(12.0, 7.6))
    axis.barh(y, values, color=[ALGORITHM_COLORS[a] for a in ALGORITHM_ORDER])
    axis.set_yticks(y, [f"{labels[a][0]}  {labels[a][1]}" for a in ALGORITHM_ORDER])
    axis.invert_yaxis()
    axis.set_xlim(0, 110)
    for index, value in enumerate(values):
        axis.text(value + 0.6, index, f"{value:.2f}%", va="center", fontsize=8)
    axis.set_xlabel("M0-M6平均可交付弃电降低率 (%)")
    axis.set_title("E23：IEEE-69下baseline与EPS fused的平均效果")
    axis.grid(axis="x", alpha=0.22)
    figure.tight_layout()
    _save_figure(figure, "ieee69_mean_effect")


def _copy_direct_training_artifacts() -> None:
    (OUTPUT / "models").mkdir(parents=True, exist_ok=True)
    for source_path in (SOURCE / "models").glob("*.pt"):
        shutil.copy2(source_path, OUTPUT / "models" / source_path.name)
    for source_path in (SOURCE / "models").glob("*.json"):
        shutil.copy2(source_path, OUTPUT / "models" / source_path.name)
    metadata = json.loads((SOURCE / "data/training_metadata.json").read_text(encoding="utf-8"))
    for record in metadata:
        dataset = str(record.get("dataset", record.get("dataset_id", "")))
        if dataset:
            record["model_path"] = f"results/E23/models/{dataset}.pt"
    (OUTPUT / "data/training_metadata.json").write_text(
        json.dumps(metadata, indent=2), encoding="utf-8"
    )


def _write_documents(boundary_rows: list[dict[str, Any]]) -> None:
    (OUTPUT / "README.md").write_text(
        """# E23：IEEE-69严格失效边界

本实验只比较IEEE-69网络下的传统baseline、Centralized greedy UB和在IEEE-69条件下直接训练的EPS fused。

每个点汇总17个数据集、M0-M6场景和30个paired seeds，所有EPS结果均来自IEEE-69网络约束数据上的直接训练。

运行入口：`python src/extra/ieee33_device_day_simulation/figures/run_e23_ieee69_failure_boundary.py`

失效边界指：原始广播请求开始频繁触发物理约束，或网络接受率低于95%参考线。物理约束包括支路热限1.0、节点最低电压0.95 p.u.和主变负载率1.0。
""",
        encoding="utf-8",
    )
    (OUTPUT / "EXPERIMENT_DESIGN.md").write_text(
        """# E23实验设计

## 研究问题

在IEEE-69更长、更深的配电网中，广播控制从高效果区域进入网络失效区域的边界在哪里？

## 对比方法

|方法|在线复杂度|说明|
|-|-|-|
|No coordination|O(0)|不发送协调信号。|
|Local SOC rules|O(1) per device|每台设备只根据自身SOC和本地规则响应。|
|MPC|O(HN)|使用滚动预测安排设备。|
|Mean-field control|O(N)|根据设备群体平均状态控制。|
|Virtual battery|O(N)|将设备群近似为聚合电池。|
|Packetized Energy Management|O(N log N)|按能量包申请和接受响应。|
|Transactive control|O(N log N)|通过价格和功率意愿进行清算。|
|EPS IEEE-69 fused|O(1)|在IEEE-69网络约束数据上直接训练，并由解析公式、学习模型和闭环修正融合控制。|
|Centralized greedy UB|O(N)|设备层理论上界，不代表网络可执行结果。|

## 固定设置

- 网络：IEEE-69 radial feeder。
- 场景：M0-M6；M0-M3增加网络压力，M4-M6改变设备空间集中。
- 数据：17个公开数据集，每个场景使用相同设备数量和相同paired seeds。
- 约束：支路负载率不超过1.0，节点电压不低于0.95 p.u.，主变负载率不超过1.0。
- 统计：按17个数据集汇总，并保留逐数据集CSV用于审计。

## 失效边界定义

E23不把弃电降低率下降单独称为物理失效。物理失效由支路、电压或主变约束是否越过阈值判断；广播请求边界另外用网络接受率和原始请求无违规比例表示。图中的95%线是运行可接受性的审计参考线，不是物理约束本身。

## 产出

- `data/ieee69_overall.csv`：9种方法在IEEE-69和M0-M6下的总体结果。
- `data/ieee69_by_dataset.csv`：逐数据集结果。
- `data/ieee69_boundary_metrics.csv`：效果、接受率、原始安全率和三类物理约束指标。
- `figures/`：失效边界曲线、总览热图、约束热图和平均效果图。
""",
        encoding="utf-8",
    )
    _write_figure_descriptions()
    _write_docx()
    _write_pptx()


def _write_figure_descriptions() -> None:
    descriptions = [
        (
            "ieee69_effect_under_stress",
            "IEEE-69严格场景下的算法效果",
            "横坐标为M0-M6网络与设备空间压力场景，纵坐标为17个数据集汇总后的可交付弃电降低率。每条线是一种baseline或IEEE-69直接训练的EPS fused。",
            "用于观察算法效果何时开始下降；M0-M3主要表示网络压力，M4-M6主要表示设备集中位置变化。",
        ),
        (
            "ieee69_failure_boundary_dashboard",
            "IEEE-69失效边界总览",
            "三层热图分别显示实际效果、网络接受率和安全层介入前原始请求无违规比例。横坐标为M0-M6，纵坐标为9种方法。",
            "网络接受率或原始请求安全率低于95%时，表示广播请求已进入需要频繁修正的边界区域；这不等同于执行结果一定违规。",
        ),
        (
            "ieee69_constraint_margins",
            "IEEE-69配网约束余量",
            "三层热图分别显示平均最大支路负载率、平均最低节点电压和平均最大主变负载率。红色等值线分别对应1.0、0.95 p.u.和1.0。",
            "该图直接回答IEEE-69中馈线、节点电压和主变约束何时越过物理边界。指标先按数据集和seed汇总，再进行17个数据集平均。",
        ),
        (
            "ieee69_mean_effect",
            "IEEE-69算法平均效果",
            "横坐标为M0-M6平均可交付弃电降低率，纵坐标列出9种方法及在线复杂度。",
            "该图只比较整体效果，不能替代失效边界热图和物理约束热图。Centralized greedy UB是设备层上界。",
        ),
    ]
    lines = [
        "# E23图片说明",
        "",
        "E23只包含IEEE-69下的baseline和IEEE-69直接训练EPS fused。",
        "",
    ]
    for index, (name, title, axes, reading) in enumerate(descriptions, start=1):
        lines.extend(
            [
                f"## {index}. {title}",
                "",
                f"![{title}](figures/{name}.png)",
                "",
                "文件：",
                "",
                f"- PNG：`figures/{name}.png`",
                f"- PDF：`figures/{name}.pdf`",
                "- 数据：`data/ieee69_boundary_metrics.csv`",
                "- 生成脚本：`src/extra/ieee33_device_day_simulation/figures/run_e23_ieee69_failure_boundary.py`",
                "",
                "### 横纵坐标和内容",
                "",
                f"- {axes}",
                "",
                "### 如何理解",
                "",
                reading,
                "",
            ]
        )
    (OUTPUT / "FIGURE_DESCRIPTIONS.md").write_text("\n".join(lines), encoding="utf-8")


def _write_docx() -> None:
    document = Document()
    document.sections[0].top_margin = Inches(0.65)
    document.sections[0].bottom_margin = Inches(0.65)
    document.sections[0].left_margin = Inches(0.7)
    document.sections[0].right_margin = Inches(0.7)
    title = document.add_heading("E23：IEEE-69严格失效边界", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "本实验只比较IEEE-69下的baseline和IEEE-69直接训练的EPS fused。每页对应一张最终图片。"
    )
    descriptions = [
        ("ieee69_effect_under_stress", "IEEE-69严格场景下的算法效果", "横坐标为M0-M6，纵坐标为可交付弃电降低率。曲线下降表示网络压力或设备空间集中开始限制实际效果。"),
        ("ieee69_failure_boundary_dashboard", "IEEE-69失效边界总览", "三层分别是效果、网络接受率和原始请求无违规比例。95%线用于标记广播请求开始频繁需要安全层修正的区域。"),
        ("ieee69_constraint_margins", "IEEE-69配网约束余量", "三层分别检查支路热限、电压下限和主变容量上限，红线是对应物理阈值。"),
        ("ieee69_mean_effect", "IEEE-69算法平均效果", "横坐标为M0-M6平均可交付弃电降低率，纵坐标标出方法和在线复杂度。"),
    ]
    for index, (name, heading, body) in enumerate(descriptions, start=1):
        document.add_heading(f"{index}. {heading}", level=1)
        document.add_picture(str(OUTPUT / "figures" / f"{name}.png"), width=Inches(6.7))
        document.add_paragraph(body)
    document.save(OUTPUT / "FIGURE_DESCRIPTIONS.docx")


def _write_pptx() -> None:
    presentation = Presentation()
    presentation.slide_width = PptInches(13.333)
    presentation.slide_height = PptInches(7.5)
    presentation.core_properties.title = "E23 IEEE-69严格失效边界"
    for index, (name, title) in enumerate(FIGURES, start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(0), PptInches(0), PptInches(13.333), PptInches(0.7))
        header.fill.solid()
        header.fill.fore_color.rgb = RGBColor(31, 50, 72)
        header.line.fill.background()
        text_box = slide.shapes.add_textbox(PptInches(0.45), PptInches(0.12), PptInches(11.7), PptInches(0.4))
        text_box.text_frame.paragraphs[0].text = title
        text_box.text_frame.paragraphs[0].font.size = Pt(20)
        text_box.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        image_path = OUTPUT / "figures" / f"{name}.png"
        with Image.open(image_path) as image:
            image_ratio = image.width / image.height
        area_width = 12.4
        area_height = 5.75
        if image_ratio >= area_width / area_height:
            image_width = area_width
            image_height = area_width / image_ratio
            image_left = 0.45
            image_top = 0.88 + (area_height - image_height) / 2
        else:
            image_height = area_height
            image_width = area_height * image_ratio
            image_left = 0.45 + (area_width - image_width) / 2
            image_top = 0.88
        slide.shapes.add_picture(
            str(image_path),
            PptInches(image_left),
            PptInches(image_top),
            width=PptInches(image_width),
            height=PptInches(image_height),
        )
        note = slide.shapes.add_textbox(PptInches(0.55), PptInches(6.78), PptInches(11.8), PptInches(0.45))
        note.text_frame.paragraphs[0].text = f"E23｜IEEE-69严格网络条件｜第{index}/4张"
        note.text_frame.paragraphs[0].font.size = Pt(10)
        note.text_frame.paragraphs[0].font.color.rgb = RGBColor(90, 90, 90)
    presentation.save(OUTPUT / "E23_figures.pptx")


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    overall, summary, boundary_rows = _prepare_data()
    _write_csv(OUTPUT / "data/ieee69_overall.csv", overall)
    _write_csv(OUTPUT / "data/ieee69_by_dataset.csv", summary)
    _write_csv(OUTPUT / "data/ieee69_boundary_metrics.csv", boundary_rows)
    _copy_direct_training_artifacts()
    labels = _labels(overall)
    global ALGORITHM_NAMES
    ALGORITHM_NAMES = {algorithm: labels[algorithm][0] for algorithm in ALGORITHM_ORDER}
    _plot_effect(overall)
    _plot_boundary_dashboard(boundary_rows)
    _plot_constraint_margins(boundary_rows)
    _plot_mean_effect(overall)
    _write_documents(boundary_rows)
    manifest = {
        "protocol": "E23_ieee69_failure_boundary_direct_only",
        "topology": "ieee69",
        "algorithms": list(ALGORITHM_ORDER),
        "scenario_modes": list(SCENARIO_ORDER),
        "dataset_count": 17,
        "seed_count": 30,
        "source_protocol": "E22_ieee69_direct_training_ablation_v2",
        "comparison_scope": "ieee69_direct_only",
    }
    (OUTPUT / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    print(f"已生成E23：{OUTPUT}")
    print(f"总体结果：{len(overall)}行；逐数据集结果：{len(summary)}行；图片：{len(FIGURES)}张")


if __name__ == "__main__":
    main()
