"""Generate the cross-dataset figure, PPTX, and Chinese figure report.

The script reads the immutable experiment outputs below results/ and writes only
the E13 output directory plus the two requested report files.
"""

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "results" / "nc_excel_experiments_new"
E7 = OUT / "E7_empirical_scaling_new"
E13 = OUT / "E13_cross_dataset_comparison_new"
E13_FIG = E13 / "Figs" / "figure_E13_cross_dataset_comparison.png"
PPTX_PATH = OUT / "nc_excel_experiments_new_实验图片总结.pptx"
MD_PATH = OUT / "nc_excel_experiments_new_实验图片总结.md"

DATASET_NAMES = {
    "bdg1_building_data_genome": "BDG1",
    "low_carbon_london": "LCL",
    "bdg2_building_data_genome": "BDG2",
    "danish_smart_heat_meters": "Danish",
    "smart_grid_smart_city": "SGSC",
    "heapo_heat_pumps": "HEAPO",
    "goiener_smart_meters": "Goiener",
    "european_lv_urban_8087": "LV-Urban-8087",
    "european_lv_rural_2731": "LV-Rural-2731",
    "european_lv_urban_35297": "LV-Urban-35297",
}
DATASET_ORDER = list(DATASET_NAMES)


def read_e13_data():
    scaling = pd.read_csv(E7 / "dataset_summary.csv")
    summary = pd.read_csv(E7 / "scaling_summary.csv")
    at_500 = summary[summary["N"] == 500].copy()
    at_500["abs_response_per_resource"] = (
        at_500["mean_response_kw"].abs() / at_500["N"]
    )
    rows = []
    for dataset in DATASET_ORDER:
        charge = scaling[(scaling.dataset == dataset) & (scaling.direction == "charge")].iloc[0]
        discharge = scaling[
            (scaling.dataset == dataset) & (scaling.direction == "discharge")
        ].iloc[0]
        q = at_500[(at_500.dataset == dataset) & (at_500.direction == "charge")].iloc[0]
        d = at_500[(at_500.dataset == dataset) & (at_500.direction == "discharge")].iloc[0]
        max_fleet = float(scaling[scaling.dataset == dataset]["maximum_unique_fleet"].iloc[0])
        rows.append(
            {
                "dataset": dataset,
                "name": DATASET_NAMES[dataset],
                "charge_beta": charge.beta,
                "charge_beta_ci_lower": charge.beta_ci_lower,
                "charge_beta_ci_upper": charge.beta_ci_upper,
                "discharge_beta": discharge.beta,
                "discharge_beta_ci_lower": discharge.beta_ci_lower,
                "discharge_beta_ci_upper": discharge.beta_ci_upper,
                "charge_cv_N500": q.cv,
                "discharge_cv_N500": d.cv,
                "charge_abs_response_per_resource_kw": q.abs_response_per_resource,
                "discharge_abs_response_per_resource_kw": d.abs_response_per_resource,
                "N90": int(charge.N_90),
                "N95": int(charge.N_95),
                "maximum_unique_fleet": max_fleet,
                "headroom_ratio": max_fleet / float(charge.N_90),
            }
        )
    return pd.DataFrame(rows)


def make_e13_figure(data):
    E13_FIG.parent.mkdir(parents=True, exist_ok=True)
    names = data["name"].tolist()
    x = np.arange(len(data))
    width = 0.36
    colors = {"charge": "#167c80", "discharge": "#d4772c"}
    cjk_font = Path("/usr/share/fonts/truetype/lyx/SimSun.ttf")
    if cjk_font.exists():
        font_manager.fontManager.addfont(str(cjk_font))
        cjk_name = font_manager.FontProperties(fname=str(cjk_font)).get_name()
    else:
        cjk_name = "DejaVu Sans"
    plt.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": [cjk_name, "DejaVu Sans"],
        "axes.unicode_minus": False,
        "axes.titlesize": 13,
        "axes.labelsize": 10,
        "xtick.labelsize": 8,
        "ytick.labelsize": 9,
    })
    fig, axes = plt.subplots(2, 2, figsize=(17, 11), constrained_layout=True)
    fig.patch.set_facecolor("white")

    ax = axes[0, 0]
    for offset, direction, label in [(-width / 2, "charge", "充电"), (width / 2, "discharge", "放电")]:
        beta = data[f"{direction}_beta"]
        lower = beta - data[f"{direction}_beta_ci_lower"]
        upper = data[f"{direction}_beta_ci_upper"] - beta
        ax.errorbar(x + offset, beta, yerr=[lower, upper], fmt="o", capsize=3,
                    color=colors[direction], label=label, markersize=5)
    ax.axhline(-0.5, color="#555555", linestyle="--", linewidth=1, label="独立聚合基准 beta=-0.5")
    ax.set_title("A 规模律斜率 beta（bootstrap 95% CI）")
    ax.set_ylabel("beta = slope(log N, log CV)")
    ax.set_xticks(x, names, rotation=35, ha="right")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(fontsize=8, loc="lower left")

    ax = axes[0, 1]
    ax.bar(x - width / 2, data["charge_cv_N500"] * 100, width, color=colors["charge"], label="充电")
    ax.bar(x + width / 2, data["discharge_cv_N500"] * 100, width, color=colors["discharge"], label="放电")
    ax.set_title("B N=500 时的响应 CV")
    ax.set_ylabel("CV (%)")
    ax.set_xticks(x, names, rotation=35, ha="right")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(fontsize=8)

    ax = axes[1, 0]
    ax.bar(x - width / 2, data["charge_abs_response_per_resource_kw"], width, color=colors["charge"], label="充电")
    ax.bar(x + width / 2, data["discharge_abs_response_per_resource_kw"], width, color=colors["discharge"], label="放电")
    ax.set_title("C N=500 时的单位资源响应")
    ax.set_ylabel("|mean response| / N (kW/resource)")
    ax.set_xticks(x, names, rotation=35, ha="right")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(fontsize=8)

    ax = axes[1, 1]
    ax.bar(x, data["headroom_ratio"], color="#3d6c9e")
    ax.axhline(1.0, color="#555555", linestyle="--", linewidth=1)
    ax.set_title("D source-unique 规模余量")
    ax.set_ylabel("maximum_unique_fleet / N90")
    ax.set_xticks(x, names, rotation=35, ha="right")
    ax.grid(axis="y", alpha=0.25)
    ax.set_ylim(0, max(6.4, data["headroom_ratio"].max() * 1.12))

    fig.suptitle("E13 十个数据集的跨数据集指标对比（E7 协议）", fontsize=17, fontweight="bold")
    fig.savefig(E13_FIG, dpi=180, bbox_inches="tight")
    plt.close(fig)


def fmt(value, digits=3):
    return f"{float(value):.{digits}f}"


def fmt_pct(value, digits=2):
    return f"{float(value) * 100:.{digits}f}%"


def make_markdown(data):
    beta_min = min(data.charge_beta.min(), data.discharge_beta.min())
    beta_max = max(data.charge_beta.max(), data.discharge_beta.max())
    cv_min = min(data.charge_cv_N500.min(), data.discharge_cv_N500.min())
    cv_max = max(data.charge_cv_N500.max(), data.discharge_cv_N500.max())
    response_min = min(data.charge_abs_response_per_resource_kw.min(), data.discharge_abs_response_per_resource_kw.min())
    response_max = max(data.charge_abs_response_per_resource_kw.max(), data.discharge_abs_response_per_resource_kw.max())
    table_rows = []
    for row in data.itertuples():
        table_rows.append(
            "| {name} | {cb} [{cl}, {cu}] | {db} [{dl}, {du}] | {cc} | {dc} | {cr} | {dr} | {n90} | {fleet:.0f} | {hr:.3f} |".format(
                name=row.name,
                cb=fmt(row.charge_beta), cl=fmt(row.charge_beta_ci_lower), cu=fmt(row.charge_beta_ci_upper),
                db=fmt(row.discharge_beta), dl=fmt(row.discharge_beta_ci_lower), du=fmt(row.discharge_beta_ci_upper),
                cc=fmt_pct(row.charge_cv_N500), dc=fmt_pct(row.discharge_cv_N500),
                cr=fmt(row.charge_abs_response_per_resource_kw, 4),
                dr=fmt(row.discharge_abs_response_per_resource_kw, 4),
                n90=row.N90, fleet=row.maximum_unique_fleet, hr=row.headroom_ratio,
            )
        )
    e13_table = "\n".join(table_rows)
    return f"""# nc_excel_experiments_new 实验图片总结

本报告汇总当前目录下已完成的 E1、E6、E7、E10、E11、E12 实验，并加入 E13 跨十数据集比较。每张原始实验图对应 PPTX 中的一页；E13 图由 E7 的 `scaling_summary.csv` 和 `dataset_summary.csv` 重新汇总绘制。E7 的网络反馈被关闭，用于隔离 source-unique 设备分布本身的统计规模律。

## 总览

| 实验 | 研究问题 | 图 | 主要结论 |
|---|---|---|---|
| E1 | 规模和相关性如何共同决定可控边界 | [E1 图](E1_scale_boundary_new/Figs/figure_E1_scale_boundary.png) | 增大规模可降低波动，但相关性会造成有限边界和平台 |
| E6 | 行为可用性如何影响有效规模和误差 | [E6 图](E6_behaviour_availability_new/Figs/figure_E6_behaviour_availability.png) | 共同可用性会同时压缩有效规模并增加跟踪误差 |
| E7 | 十个真实数据分布是否保持 source-unique 规模律 | [E7 图](E7_empirical_scaling_new/Figs/figure_E7_empirical_scaling.png) | 20 个数据集-方向组合的斜率均接近 `N^-1/2` |
| E10 | IEEE33 网络限制如何影响响应交付 | [E10 图](E10_network_constraints_new/Figs/figure_E10_network_constraints.png) | 压力增加截断和线路风险，当前 zonal 未显示正增益 |
| E11 | 三类压力的失效边界是什么 | [E11 图](E11_failure_modes_new/Figs/figure_E11_failure_modes.png) | 三类机制存在不同的失效阈值，但横轴不具跨机制物理可比性 |
| E12 | 十个数据集的基础产物和证据是否完整 | [E12 图](E12_dataset_pipeline_integrity_new/Figs/figure_E12_dataset_pipeline_integrity.png) | 20/20 个模式有效，设备规模校准全部通过 |
| E13 | 十个数据集的对应指标如何比较 | [E13 图](E13_cross_dataset_comparison_new/Figs/figure_E13_cross_dataset_comparison.png) | 不同数据集的 beta、CV 和单位资源响应相近，规模余量差异较大 |

## E1：规模可控边界

图片：[figure_E1_scale_boundary.png](E1_scale_boundary_new/Figs/figure_E1_scale_boundary.png)

| 子图 | 横轴 | 纵轴/图形编码 | 指标含义 | 当前结果及其证明 |
|---|---|---|---|---|
| A | 设备数 `N` | 纵轴为相关条件，颜色为 `p_ctrl` | `p_ctrl` 是满足 `NRMSE<=0.10`、阻塞比例不超过 0.05 且符号一致率至少 0.95 的重复比例，范围 0-1 | `p_ctrl` 覆盖 0-1；低相关、大规模区域更容易达到 1，说明规模效应受相关性条件限制 |
| B | 名义相关系数 `rho_nominal` | 临界规模 `N*`，对数轴 | `N*` 是使 `p_ctrl` 的 95% Wilson 下界达到 0.90 的最小规模；`not reached` 表示扫描范围内未达到 | IID 的 `N*=300`，`rho=0.00312` 时为 500；更高四档相关性直到 `N=3000` 仍未达到，证明不能假设无限加设备即可消除相关性 |
| C | `N`，对数轴 | 固定相关条件下的条件 CV，对数轴 | `CV=Std(Y)/(abs(Mean(Y))+epsilon)`，反映重复聚合响应的相对波动 | CV 约从 0.247 降到 0.021；高相关曲线保持更高，说明相关性减慢波动收敛 |
| D | 注入的名义相关系数 | 实测条件残差相关系数；点颜色表示 `log10(N)` | 检验相关性是否实际进入设备残差；虚线 `y=x` 是名义值与实测值相等的参考 | 实测相关随名义相关增加但总体低于名义值，说明解释时必须区分注入参数和实测 residual correlation |

E1 按“`p_ctrl` 的 Wilson 95% 下界达到 0.90”这一预注册判据计算，可控网格占比为 `A_ctrl=0.2143`；这不是简单统计点估计 `p_ctrl>=0.90` 的网格比例。该结果支持“规模效应存在有限适用边界”，不支持“任意相关性都能靠增加 `N` 消除”。

### E1 的关键变量如何设置并进入仿真

- **设备规模 `N`**：扫描 `50、100、300、500、1000、2000、3000`。先从 SGSC 数据池按 `source_device_id` 无重复地抽取 3,000 个资源，再按六个 IEEE33 逻辑区域取平衡子集；因此 `N` 表示同一并发 fleet 中的 source-unique 资源数，不是把同一用户的多个日期重复计数。
- **相关变量**：代码真正输入仿真的变量是 `zone_correlation={0,0.0125,0.0289,0.05,0.1125,0.2425}`。每个重复中，每个区域抽取一个 `z_z ~ Normal(0, zone_correlation)` 的共同冲击，并把区域内资源的期望响应乘以 `1+z_z`；该冲击在该重复的 48 个条件内保持不变。图中的 `rho_nominal={0,0.00312,0.00721344,0.01248,0.02808,0.060528}` 是对应的标称残差相关档位，不是直接传入模拟器的方差参数。
- **实测相关 `rho_measured`**：保存 100 次测试重复的设备级响应后，先减去每个固定 EPS 条件下的重复均值，再计算设备对之间的平均残差相关。因此它检验的是去除共同条件均值后仍存在的依赖，不是原始功率曲线相关。
- **冻结目标与测试**：对每个 `N`，使用 IID 档的 10 次训练重复拟合线性条件均值目标；随后冻结该目标，在同一 `N` 的六个相关档位上各做 100 次测试。每次测试都包含 48 个随机 profile 条件，并同时生成 valley-filling 充电和 peak-shaving 放电广播。
- **可控事件**：单次测试只有在 `NRMSE<=0.10`、`network_scale<=0.05` 的条件占比不超过 0.05、且非零目标的响应方向一致率至少 0.95 时才记为成功。`p_ctrl` 是 100 次测试中的成功比例，`N*` 还要求其 Wilson 95% 下界达到 0.90。

`A_ctrl=0.2143` 是当前 `7×6=42` 个离散网格单元的等权通过比例，不是连续参数空间的几何面积。相关冲击也是合成的区域共同乘法扰动，不能解释为 SGSC 现场测得的用户相关性。

### E1 与 Reviewer #3-I.b、#3-II.c 的对应关系

Reviewer #3-I.b 质疑大规模确定性所依赖的弱相关假设，Reviewer #3-II.c 要求给出规模增加何时有效、何时失效的边界。E1 直接把区域共同相关作为测试压力，并用冻结的 IID 目标检查分布外性能：IID 时 `N*=300`、最低非零相关档时 `N*=500`，更高四档直到 `N=3000` 仍未达到。这说明“增加资源数”只在相关足够弱时能稳定提高可控概率，共同误差不会按独立样本的速度被平均掉。

E1 支持的是“在当前 SGSC 映射、EPS 规则、IEEE33 配置和合成区域冲击下存在有限规模-相关边界”。它不支持把 `N*=300/500` 当作其他数据集或现场系统的通用门槛，也不证明时间相关、用户共同参与、控制器同步和网络拥塞都能由同一个 `rho` 表示。

## E6：行为驱动可用性

图片：[figure_E6_behaviour_availability.png](E6_behaviour_availability_new/Figs/figure_E6_behaviour_availability.png)。三条线分别是 IID、Markov 时间持续和 Community 共同可用性。

| 子图 | 横轴 | 纵轴 | 指标含义 | 当前结果及其证明 |
|---|---|---|---|---|
| A | 参与率 0.3-1.0 | `N_eff_behavior` | 经可用性和行为相关性调整后的有效独立规模，用来判断在线设备数能否转化为独立统计样本 | 参与率 0.3 时 IID 约 802，而 Community 约 76.7；共同可用性会严重压缩有效规模 |
| B | 参与率 | NRMSE | 归一化均方根跟踪误差，越低越好 | 参与率从 0.3 增至 1.0 时 NRMSE 约从 0.640 降至 0.058，说明参与率下降直接恶化跟踪精度 |
| C | 参与率 | `Pi_beh = NRMSE_structure - NRMSE_IID` | 在同一参与率下，结构化行为相对 IID 的额外误差；0 表示没有额外惩罚 | Community 惩罚为正且明显高于 IID/Markov，证明共同用户行为带来的损失不只是在线数量减少 |
| D | 参与率 | 95% 备用功率需求（kW） | 为弥补目标与实际可交付响应差额而预留的功率，越低越好 | 范围约 118-833 kW；低参与率和 Community 条件需要更多备用 |

E6 证明“可用率下降”和“共同用户行为”是两个不同机制，但不证明所用 Markov/Community 参数来自真实社区现场估计。

### E6 的参与行为如何设置

参与率扫描为 `a={0.3,0.5,0.7,0.9,1.0}`。对资源 `i` 和条件 `t`，行为 mask `A_beh(t,i)=1` 表示允许设备尝试响应，`A_beh(t,i)=0` 表示本次不参与。三种 mask 保持相同边际参与率，但依赖结构不同：

- **IID Bernoulli**：每个资源、每个条件独立按 `Bernoulli(a)` 抽样，用作相同在线数量下的独立基准。
- **Markov**：资源状态受上一条件影响，持续性参数 `h=0.90`；`p11=h+(1-h)a`，`p01=(1-h)a`。例如 `a=0.3` 时，已参与者下一条件继续参与的概率为 0.93，未参与者转为参与的概率为 0.03。
- **Community**：每个 IEEE33 区域生成共同状态，资源-条件单元以 `shared_fraction=0.50` 的概率采用该共同状态，否则采用自己的独立状态。这会让同一区域的一部分用户同时参与或同时离线。

同一批 3,000 个资源和同一组 48 个 EPS/profile 条件用于全部组合。冻结目标由 10 次无额外行为 mask 的基线重复拟合，每个参与率-结构组合使用 100 次测试重复。目标不会随参与率同比下调，否则会人为消除用户不参与造成的功率缺额。

### E6 与 Reviewer #2-4 的对应关系

Reviewer #2-4 关注的是：设备并非始终在线，用户行为可能使“设备规模越大、聚合响应越确定”的假设失效。E6 中的“用户是否参与”特指用户是否允许设备进入当前 EPS 广播的功率响应流程，不是参与问卷、数据上传或市场交易。

当 `A_beh=1` 时，设备可以尝试执行广播要求的响应：低谷填充（`valley filling`）中的充电，或削峰（`peak shaving`）中的放电。当 `A_beh=0` 时，该设备在这个条件下被行为门控屏蔽，不执行本次广播响应。即使 `A_beh=1`，设备仍要经过区域可用性、通信门控、SOC、功率/能量限制和 IEEE33 网络约束，因此参与率不是最终响应率。

这项实验证明了三点：

1. **名义规模不等于有效规模。** 本实验名义使用 3,000 个资源，但参与率从 1.0 降到 0.3 时，有效行为规模会显著下降。所有结构合并看，`N_eff_behavior` 约为 76.7-2685，说明只报告注册设备数 `N` 会高估可交付确定性。
2. **低参与率会造成可测的性能损失。** NRMSE 约从 0.0580 增加到 0.6402，95% 备用功率需求约为 117.94-832.54 kW。参与不足会直接造成目标响应与实际响应之间的缺额，而不是可以忽略的随机噪声。
3. **共同参与行为比独立离线更危险。** 在名义参与率均为 0.3 时，IID 的 `N_eff_behavior` 约为 802.4、NRMSE 约为 0.6340，而 Community 的有效规模约为 76.7、NRMSE 约为 0.6402、备用需求约为 832.5 kW。也就是说，平均参与人数相同并不代表统计效果相同；区域内共同离线会进一步降低有效独立规模。

因此，E6 对 reviewer 的准确回复是：本文不再假设所有设备始终可用，而是把参与率和参与相关性作为额外行为压力显式加入广播响应前的门控。结果表明，参与率下降主要通过减少有效参与规模提高误差和备用需求；在相同平均参与率下，Community 相关行为还会产生额外惩罚。这说明规模确定性依赖有效且近似独立的参与资源数量，而不是名义注册设备数量。

E6 仍不能证明真实用户参与率或真实社区行为已经被现场验证。当前正式实现是在 SGSC 的 source-unique 资源上生成合成 IID/Markov/Community mask，`a=0.3-1.0` 是审稿压力测试档位，不是某个社区的统计估计。因此论文中应称为“合成用户可用性压力测试”，不能称为“真实用户行为验证”；真实 EV session、回家/离站时间、充电截止时间、价格响应和用户优先级仍需独立数据实验。

## E7：跨数据集 source-unique 规模律

图片：[figure_E7_empirical_scaling.png](E7_empirical_scaling_new/Figs/figure_E7_empirical_scaling.png)。

| 子图 | 横轴 | 纵轴/图形编码 | 指标含义 | 当前结果及其证明 |
|---|---|---|---|---|
| A | source-unique 设备数 `N`，对数轴 | 聚合响应 CV，对数轴；每条线为一个数据集（充电方向） | 检验真实负荷分布下重复聚合响应的相对波动是否随规模下降 | 十个数据集都完成充电方向扫描并达到各自 source-unique 上限，说明该协议在不同分布上可执行 |
| B | 数据集-方向组合 | `beta=slope(log N, log CV)`，误差线为 bootstrap 95% CI | `beta=-0.5` 是独立同分布聚合的基准；充放电分开估计以避免方向抵消 | 20 个组合的 beta 范围约为 -0.560 至 -0.447，均接近 -0.5，且所有组合 `N90=N95=500`；支持十个数据分布均保持近似 `N^-1/2` 律 |

E7 隔离了网络反馈，因此不能单独证明基础 Figure 2-5 已对所有数据集完整，也不能证明每个数据集含有真实电池、SOC、PV 或现场 EPS 响应。

### E7 的关键变量如何设置并进入实验

- **数据集和方向**：十个数据集分别运行 charge 与 discharge，共 20 个组合。charge 只生成 valley-filling 广播，discharge 只生成 peak-shaving 广播，避免正负方向在均值中相互抵消。
- **source-unique 规模 `N`**：候选为 `50、100、200、500、1000、2000、3000`，并追加不超过配置规模和六区域平衡唯一资源容量的最大值。每个数据集先抽取一个最大 source-unique master fleet，再从中取嵌套的平衡子集，因此同一并发 fleet 中不重复 source。
- **重复和目标**：每个数据集-方向-`N` 运行 110 次，每次包含 24 个 profile/EPS 条件。前 10 次响应的逐条件均值定义该 `N` 下的冻结目标，后 100 次用于 CV、可控概率和阈值统计；网络反馈被显式设为 `False`，IEEE33 不截断响应。
- **CV 与 `beta`**：先对每个测试重复的 24 条件响应取均值，再在 100 个重复均值之间计算 `CV=Std/abs(Mean)`。`beta` 是各 source-unique `N` 点上 `log(CV)` 对 `log(N)` 的普通线性斜率；95% CI 通过在每个 `N` 内重采样 100 个测试均值得到 1,000 条 bootstrap 斜率。
- **`N90/N95`**：沿离散 `N` 扫描，寻找 `p_ctrl` 的 Wilson 95% 下界首次达到 0.90/0.95 的规模。它是当前误差、阻塞和方向三重判据下的门槛，不是 CV 降到 10% 或 5% 的规模。

### E7 与 Reviewer #2-2、#2-3 的对应关系

Reviewer #2-2、#2-3 关注独立实体数量是否足以支撑规模结论，以及规模律能否在多个真实数据分布和充放电方向上复现。E7 使用 source-unique 并发 fleet 而不是重复用户切片，20 个数据集-方向组合的 `beta` 均落在约 `-0.560` 至 `-0.447`，且 `N90=N95=500`。这支持“在当前映射后的十种负荷分布中，相对波动具有接近平方根律的下降趋势”，并证明该趋势不是单一数据集或单一响应方向造成的。

E7 不支持“十个数据集都提供了真实电池控制响应”。真实数据主要驱动负荷/profile 分布，电池容量、SOC、响应概率和部分外部输入仍来自映射或 fallback；当前实现也没有执行设计文档中的 empirical-bootstrap、leave-one-dataset-out、参数弹性或随机效应 meta-analysis。因此它不能证明跨数据集零样本迁移能力，也不能替代 E10 的网络可交付性验证。

## E10：网络约束交付

图片：[figure_E10_network_constraints.png](E10_network_constraints_new/Figs/figure_E10_network_constraints.png)。横轴缩写由 `mode-layout-control` 组成，布局包括 balanced、feeder concentrated、end-bus concentrated，控制包括 global/zonal。

网络接受率定义为

`A_net = sum(min(abs(P_accepted), abs(P_desired)) * same_direction) / sum(abs(P_desired))`，截断指数为 `I_clip=1-A_net`。

| 子图 | 横轴 | 纵轴/图形编码 | 指标含义 | 当前结果及其证明 |
|---|---|---|---|---|
| A | 12 个 mode-layout-control 组合 | 网络接受率 `A_net` | 期望 EPS 响应中在 IEEE33 约束下实际可交付的比例，越高越好 | 范围 0.9510-0.9852，weak 组整体更高，说明网络压力会降低可交付响应 |
| B | 同 A | 截断指数 `I_clip` | 网络丢失的期望响应比例，越低越好 | 范围 0.0148-0.0490，stress/feeder/zonal 最大，证明压力和空间集中会增加截断 |
| C | `I_clip` | 电压、线路、变压器三者中的最小安全裕度；颜色区分 weak/stress | 衡量多交付是否以安全约束为代价；小于 0 表示至少一种约束越限 | 最小线路裕度约 -0.453，而电压和变压器裕度仍为正；stress 组存在支路过载风险 |
| D | mode-layout 配对 | `Delta A_net=A_net(zonal)-A_net(global)` | 检验区域化广播是否提升网络可交付比例；大于 0 才是 zonal 增益 | weak 增益约为 0，stress 三种布局为 -0.0058 至 -0.0080；当前实现不支持 zonal 稳定优于 global |

### E10 的关键变量如何设置并进入仿真

- **固定运行规模**：从 SGSC 配置抽取同一批 3,000 个 source-unique 资源，每个组合连续运行 288 个五分钟步长且 `reset_each_step=False`，因此 SOC 和设备状态会在一天内演化。每个组合运行 8 个随机重复。
- **网络模式**：`weak_correlation` 与 `network_stress` 分别加载弱约束和压力配置；模式改变网络容量/压力配置，并保持 EPS 与设备仿真框架一致。它不是 E1 的六档 `rho` 扫描。
- **物理布局**：balanced 保留原母线；end-bus-concentrated 把每个资源移到其逻辑区域的最末端母线；feeder-concentrated 把前 80% 资源的物理母线映射到最后一个 feeder/区域的母线。后两种布局只改 `bus_id`，不改资源的逻辑 `zone_id`，以隔离物理位置影响。
- **广播方式**：global 先汇总全 fleet 的负荷和外部输入，生成一条广播并复制给六个区域；zonal 按各逻辑区域自己的净供需生成独立广播。两者使用不同但固定的预生成 schedule，当前不是严格逐时同信号配对。
- **网络交付和安全**：设备先给出 `P_desired`，IEEE33 调度层计算网络缩放后得到 `P_accepted`。`A_net` 只计同方向且不超过 desired 绝对值的交付；`m_V`、`m_L=1-max branch loading`、`m_T=1-max transformer loading` 由 accepted 后的网络状态计算。任一裕度小于 0 都表示硬约束未满足；`unsafe_step_fraction` 当前只统计支路过载或电压越限步数，不含变压器越限。

### E10 与 Reviewer #4-3 的对应关系

Reviewer #4-3 关注 IEEE33 的线路、电压和变压器约束是否真正进入响应链，以及网络截断后还能交付多少广播响应。E10 中压力模式把接受率从 weak 的约 0.985 降到最低 0.951，并把截断提高到最高 0.049；feeder-concentrated/stress/zonal 的平均最小线路裕度约为 `-0.453`，stress 组合的 unsafe step 比例约为 11%-15%。这证明物理布局和网络压力确实改变交付结果，不能用 unconstrained desired response 代表最终可交付响应。

同时，负线路裕度意味着当前压力场景没有满足预设的硬网络安全判据。zonal 相对 global 的 `Delta A_net` 还为负，因此本图不能用于声称“区域广播提高了交付”或“IEEE33 安全约束已完全解决”。当前实现也没有 centralized-feasible 上界、严格同 schedule 配对和通信开销比较；Reviewer 回复应写成“识别并量化了网络瓶颈，同时暴露了 stress 下的剩余安全失效”，而不是宣称方法已在所有网络条件下安全。

## E11：失效模式

图片：[figure_E11_failure_modes.png](E11_failure_modes_new/Figs/figure_E11_failure_modes.png)。

| 项目 | 横轴/纵轴或图形编码 | 指标含义 | 当前结果及其证明 |
|---|---|---|---|
| 失效概率曲线 | 横轴为预注册压力强度，纵轴为失效概率 `p_F`（0-1）；三条线分别是 residual correlation、behavior unavailability、network stress；水平线为 `p_F=0.5` | `x50` 是曲线首次插值达到 0.5 的压力阈值；不同机制的压力强度在机制内部归一化 | `x50` 分别为 residual correlation 0.011452、behavior unavailability 0.057471、network stress 0.55，证明三类机制具有不同失效边界，但不能把这些横坐标当作统一物理量。`recovery_fraction` 尚未计算，故不能证明失效后的恢复能力 |

### E11 的三条压力轴如何构造

- **Residual correlation**：横轴直接使用 E1 的 `rho_nominal`；每个档位的失效概率是七个 `N` 上 `1-p_ctrl(N,rho)` 的算术平均。因此该曲线把规模维度边际化了，`x50` 不是任一固定 `N` 的相关阈值。零相关点的平均失效概率已约为 0.281，所以 `x10=0`；最高扫描档仍未达到 0.90，故 `x90` 未定义。
- **Behavior unavailability**：只使用 E6 的 Community 结构，横轴为 `1-participation`，即区域共同参与机制下的名义不参与率；纵轴沿用 E6 的不可控概率。它没有合并 IID 和 Markov，也不能解释成用户永久退出比例。
- **Network stress**：失效事件定义为一次 E10 重复中存在任意 unsafe step，或 `A_net<0.80`。横轴是六个 mode-layout 类别的人为有序编码：weak/balanced=0、weak/end-bus=0.25、weak/feeder=0.50、stress/balanced=0.60、stress/end-bus=0.80、stress/feeder=1.0；每个类别合并 global/zonal 和各重复。该数值只是场景排序，不是可测的线路负载率或容量缩放百分比。
- **阈值计算**：先对离散点的失效概率取单调累计上包络，再在首次跨过 0.10、0.50、0.90 的相邻点间作线性插值。图线不是 logistic 拟合，当前点数也不足以给出阈值置信区间。

### E11 与 Reviewer 意见的对应关系

E11 把 Reviewer #3-I.b/#3-II.c 的相关-规模失效、Reviewer #2-4 的用户不参与失效，以及 Reviewer #4-3 的网络失效放到同一张“机制内剂量-响应”图中。它支持三种质疑都能在当前仿真中产生可观测失效，并给出各自扫描范围内的过渡位置；它不支持比较 `0.011452 < 0.057471 < 0.55` 后宣称哪种机制更危险，因为三个横轴的单位、构造和聚合方式完全不同。

E11 还是 E1/E6/E10 的二次汇总，不是独立新实验。当前没有交互项、独立 hold-out 阈值验证、阈值 CI，也没有任何 registered recovery 策略；因此不能回答“失效后能恢复多少”，不能用本图声称方法具有故障恢复能力。

## E12：基础实验完整性与证据质量

图片：[figure_E12_dataset_pipeline_integrity.png](E12_dataset_pipeline_integrity_new/Figs/figure_E12_dataset_pipeline_integrity.png)。E12 不训练模型，只审计当前 `results/*_ieee33_real_load` 产物。

| 子图 | 横轴 | 纵轴/图形编码 | 指标含义 | 当前结果及其证明 |
|---|---|---|---|---|
| A | 十个数据集 | 完整度百分比；橙色为 network stress，蓝色为 weak | 检查运行状态、五张图、双向缩放文件和 Figure 4 独立性是否同时满足 | 20/20 个 mode 完整；BDG2 和 HEAPO 完整重跑后四个模式均达到 100% |
| B | 十个数据集 | 独立资源数；绿色为六区严格等量上限，紫色为原请求主规模，黑线为实际选择规模 | 检查主实验是否超过最小区域可提供的 source-unique 数量 | 所有实际选择均未超过容量；BDG2 取 1566，HEAPO 由请求 1362 下调为 1338 |
| C | 十个数据集 | 非零验证响应中的 charge 比例；方形为 stress、圆形为 weak、红色叉号为无效残留，虚线为 0.5 | 检查验证是否同时覆盖充放电；0.5 表示方向数量平衡 | 十个数据集的 20 个模式均为 0.5，证明当前验证产物全部覆盖平衡的充放电方向 |
| D | 十个数据集 | 弃能降低率；黑边区分 stress，圆形为 weak；绿色表示实测/混合输入，灰色表示反事实输入 | 检查 EPS 正向吸收是否降低外部输入盈余，并标明证据来源 | 多数有非零基线输入的数据集约 97%-100%；Danish 为 8.90%-14.25%；HEAPO 基线弃能为 0，因此降低率不定义而不是实验无效 |

E12 将两类判据分开报告：模式有效要求状态完成、五图齐全、双向缩放文件存在且两张 Figure 4 不同；设备规模是否合法由 fleet capacity audit 另行检查。当前有效基础模式为 20/20，设备规模校准也全部通过。

### E12 的审计变量如何定义

- **审计对象**：十个数据集各检查 `network_stress` 和 `weak_correlation`，共 20 个基础模式。E12 不重新训练模型、不生成新的 EPS 响应，只读取各模式已经落盘的状态、图片、估计输出和网络时间序列。
- **完整度**：每个模式共有 8 项检查，即 `run_status=completed`、5 张预期 Figure 2-5、`directional_n_scaling.json` 存在，以及两张 Figure 4 的 SHA-256 不同。完整度是通过项数除以 8；`valid_base_mode` 要求 8 项全部通过，而不是分数达到某个宽松阈值。
- **规模合法性**：主规模上限是六个区域中最少 source 数乘以 6，保证严格等量抽样；相关扫描规模另用允许区域数量相差至多 1 的最大平衡容量。图 B 的黑线是配置中实际选择的主规模，用来检查请求规模是否已按 source 容量下调。
- **方向覆盖**：从估计器测试集的 actual accepted response 中，以正值为 charge、负值为 discharge、近零为 blocked；图 C 只计算非零响应中的 charge 比例。0.5 证明测试输出数量方向平衡，不证明真实系统自然发生充放电的概率各为 50%。
- **弃能核算**：每个五分钟步先算 `baseline=max(energy_input-load,0)`，再算 `after_EPS=max(energy_input-load-max(accepted_control,0),0)`，乘 `5/60` 小时并由 kWh 转为 MWh。只有正向充电响应可减少该账本中的弃能；基线为 0 时 reduction 不定义。

### E12 与 Reviewer #2-2、#2-3 的对应关系

E12 回应的是证据链是否真的完整：每个数据集是否有独立实体容量、充放电是否都有验证样本、weak/stress 的基础图是否齐全，以及 Figure 4 是否误用重复图片。当前 20/20 模式有效，且 BDG2 取 1,566、HEAPO 取 1,338 后均未超过平衡唯一资源容量。这支持“本轮用于回复 Reviewer 的基础产物已完整且规模选择可审计”。

E12 是完整性审计，不是算法性能优越性的证据。图 D 也必须按 provenance 分层：SGSC 可提供 measured-input 驱动证据，BDG2 是 measured/counterfactual 混合，其余多数为 load-shape counterfactual；因此 97%-100% 的核算降低率不能统一写成十个现场数据集都观测到真实弃电减少。HEAPO 的基线弃能为 0，只表示该比例没有分母，不是运行无效。

## E13：十数据集指标对比

图片：[figure_E13_cross_dataset_comparison.png](E13_cross_dataset_comparison_new/Figs/figure_E13_cross_dataset_comparison.png)。该图不是新的实验协议，而是对 E7 已有结果的统一横向汇总。

| 子图 | 横轴 | 纵轴/图形编码 | 指标含义 | 当前结果及其证明 |
|---|---|---|---|---|
| A | 十个数据集；每个数据集含充电和放电两个点 | beta，误差线为 bootstrap 95% CI；虚线为 `beta=-0.5` | beta 是 `log(CV)` 对 `log(N)` 的斜率；-0.5 表示独立聚合的平方根收敛基准 | 20 个点落在约 {beta_min:.3f} 至 {beta_max:.3f}；整体围绕 -0.5 波动，说明规模律对数据集和方向具有较强一致性 |
| B | 十个数据集；分组柱为充电/放电 | `N=500` 的 CV（百分比） | CV 衡量同一规模下聚合响应的相对波动，越低越稳定 | 范围约 {cv_min * 100:.2f}% 至 {cv_max * 100:.2f}%，最大差约 {(cv_max - cv_min) * 100:.2f} 个百分点；没有出现数量级差异 |
| C | 十个数据集；分组柱为充电/放电 | `abs(mean_response_kw)/N`，单位 kW/resource | 单个 source-unique 资源平均贡献的绝对响应大小；用于比较响应尺度，不等同于可交付容量 | 范围约 {response_min:.4f}-{response_max:.4f} kW/resource，整体集中，说明 E7 协议下单位资源响应尺度相近 |
| D | 十个数据集 | `maximum_unique_fleet/N90` | source-unique 可用上限相对于达到 90% 可控概率所需规模的余量；大于 1 表示尚有协议内扩展空间 | 余量约为 {data.headroom_ratio.min():.3f}-{data.headroom_ratio.max():.3f}；BDG1 接近 1，而多数大数据集更高，说明数据覆盖规模而非 beta 是主要差异 |

### E13 的变量来源和计算关系

E13 没有再次运行仿真，所有点都由 E7 的 `scaling_summary.csv` 和 `dataset_summary.csv` 派生。A 图直接读取每个数据集/方向的 `beta` 及 bootstrap CI；B 图固定读取 `N=500` 的 CV；C 图把同一行的聚合平均响应绝对值除以 500；D 图把该数据集可用的最大 source-unique fleet 除以 `N90`。因此四幅子图共享 E7 的样本、目标和随机重复，不能视作四组独立证据。

### E13 与 Reviewer #2-2、#2-3 的对应关系

E13 把 Reviewer 关心的“跨数据集是否一致”和“独立实体规模是否充足”转换成可直接横向比较的四个量。A-C 显示当前协议下收敛斜率、`N=500` 的相对波动和单位资源响应尺度没有出现数量级差异；D 则显示 BDG1 的 source-unique 余量仅为 1.008，而多数组为 6.0，说明数据覆盖余量比斜率本身更不均衡。

这些结果支持 E7 结论在十种映射后数据分布和两个方向上的一致性，不支持数据集之间的因果排名。单位资源响应由共同 EPS/设备映射规则共同决定，不等于各数据集的实测柔性容量；`maximum_unique_fleet/N90` 也只是当前扫描协议的资源余量，不是部署成本或市场容量指标。由于 E13 完全复用 E7，它不能被计为独立复现实验。

### E13 数值表

CV 为百分比，单位资源响应为 kW/resource，beta 方括号内为 bootstrap 95% CI。

| 数据集 | 充电 beta [CI] | 放电 beta [CI] | 充电 CV@500 | 放电 CV@500 | 充电单位响应 | 放电单位响应 | N90 | 最大 unique fleet | 余量 |
|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|
{e13_table}

## 解释边界

1. E7/E13 使用独立 runner 并关闭网络反馈，结论针对 source-unique 设备分布的统计规模律，不等于 IEEE33 网络下的最终交付能力。
2. `N90`、`N95` 是当前预注册阈值下的扫描结果；它们不是所有运行条件下的普适设备数量。
3. BDG2 与 HEAPO 已于本轮按平衡唯一资源上限完整重跑；E12 当前不含无效或残留模式。
4. SGSC 的外部 gross generation 可形成实测输入驱动证据；BDG2 只有少量 solar 重叠列，属于实测与反事实混合证据；其他负荷型数据集的 Figure 4 不能称为真实 PV 自消费或现场弃电观测。
5. E11 的三个压力轴是机制内注册量，不能跨机制比较数值大小；恢复实验尚未实现。
"""


def add_textbox(slide, left, top, width, height, text, font_size, color, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.08)
    frame.margin_top = Inches(0.04)
    frame.margin_bottom = Inches(0.04)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "SimSun"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    return box


def add_slide(prs, image_path, title, conclusion):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(247, 248, 250)
    add_textbox(slide, Inches(0.45), Inches(0.18), Inches(12.35), Inches(0.48), title, 22, (31, 43, 55), True)
    add_textbox(slide, Inches(0.45), Inches(0.68), Inches(12.35), Inches(0.24), "实验结果图；本页仅放置一张图片", 8, (100, 110, 120))

    with Image.open(image_path) as img:
        width, height = img.size
    image_ratio = width / height
    area_left, area_top, area_width, area_height = 0.45, 1.02, 12.35, 5.75
    area_ratio = area_width / area_height
    if image_ratio >= area_ratio:
        pic_width = area_width
        pic_height = area_width / image_ratio
        pic_left = area_left
        pic_top = area_top + (area_height - pic_height) / 2
    else:
        pic_height = area_height
        pic_width = area_height * image_ratio
        pic_left = area_left + (area_width - pic_width) / 2
        pic_top = area_top
    slide.shapes.add_picture(str(image_path), Inches(pic_left), Inches(pic_top), Inches(pic_width), Inches(pic_height))

    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(6.95), Inches(12.35), Inches(0.45))
    strip.fill.solid()
    strip.fill.fore_color.rgb = RGBColor(225, 237, 238)
    strip.line.color.rgb = RGBColor(190, 213, 214)
    add_textbox(slide, Inches(0.58), Inches(6.96), Inches(12.1), Inches(0.42), conclusion, 10, (35, 67, 70), False)


def make_pptx():
    figures = [
        ("E1 规模可控边界", OUT / "E1_scale_boundary_new" / "Figs" / "figure_E1_scale_boundary.png", "回应 Reviewer #3-I.b/#3-II.c：IID 时 N*=300、最低非零相关时为 500；更高相关在 N<=3000 内未达到。"),
        ("E6 行为驱动可用性", OUT / "E6_behaviour_availability_new" / "Figs" / "figure_E6_behaviour_availability.png", "回应 Reviewer #2-4：参与率下降减少有效规模，Community 共同离线进一步提高 NRMSE 和备用需求。"),
        ("E7 真实数据 source-unique 规模律", OUT / "E7_empirical_scaling_new" / "Figs" / "figure_E7_empirical_scaling.png", "回应 Reviewer #2-2/#2-3：20 个数据集-方向组合的 beta 接近 -0.5；结论限于关闭网络反馈的分布规模律。"),
        ("E10 IEEE33 网络约束", OUT / "E10_network_constraints_new" / "Figs" / "figure_E10_network_constraints.png", "回应 Reviewer #4-3：stress 下线路裕度最低约 -0.453，硬安全判据未满足；当前 zonal 也没有正交付增益。"),
        ("E11 失效模式", OUT / "E11_failure_modes_new" / "Figs" / "figure_E11_failure_modes.png", "三条曲线复用 E1/E6/E10 且压力轴单位不同；当前只支持机制内失效边界，不支持跨机制危险度或恢复结论。"),
        ("E12 数据集流水线完整性", OUT / "E12_dataset_pipeline_integrity_new" / "Figs" / "figure_E12_dataset_pipeline_integrity.png", "回应证据完整性质疑：20/20 个基础模式有效且规模校准通过；该审计本身不证明算法性能优越。"),
        ("E13 十数据集指标对比", E13_FIG, "回应 Reviewer #2-2/#2-3：beta、CV@500 和单位响应整体接近，unique 余量差异较大；本图完全派生自 E7。"),
    ]
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    for title, image_path, conclusion in figures:
        if not image_path.exists():
            raise FileNotFoundError(image_path)
        add_slide(prs, image_path, title, conclusion)
    prs.save(PPTX_PATH)


def main():
    data = read_e13_data()
    E13.mkdir(parents=True, exist_ok=True)
    data.to_csv(E13 / "e13_cross_dataset_summary.csv", index=False, float_format="%.10g")
    make_e13_figure(data)
    make_pptx()
    MD_PATH.write_text(make_markdown(data), encoding="utf-8")
    print(f"Generated {E13_FIG}")
    print(f"Generated {PPTX_PATH}")
    print(f"Generated {MD_PATH}")


if __name__ == "__main__":
    main()
