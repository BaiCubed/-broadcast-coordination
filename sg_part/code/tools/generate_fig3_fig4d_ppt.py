"""生成 Figure 3 与最终 Figure 4D 的跨数据集汇总 PPT。"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import sys
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.extra.ieee33_device_day_simulation.figures.fig4d_extra_baselines import (
    FIG4D_COLORS,
    FIG4D_LABELS,
    FIG4D_ORDER,
)
from src.extra.ieee33_device_day_simulation.figures.fig4d_final_protocol import (
    AVAILABILITY_MODES,
    FLEET_MODES,
    NETWORK_MODES,
    discover_result_roots,
    output_stem,
)
from src.extra.ieee33_device_day_simulation.figures.plot_fig4d_final_overall import (
    _condition_payloads as load_condition_payloads,
)


DATASET_SUFFIX = "_ieee33_real_load"
BASELINE_KEYS = tuple(
    key for key in FIG4D_ORDER
    if key not in {"no_coordination", "eps_broadcast", "centralized_optimal"}
)


def _read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _dataset_name(result_root: Path) -> str:
    name = result_root.parents[1].name
    return name.removesuffix(DATASET_SUFFIX)


def _condition_title(fleet_mode: str, network_mode: str, availability_mode: str) -> str:
    fleet = "Fixed 5,000 logical devices" if fleet_mode == "fixed5000" else "Source-unique fleet limit"
    network = "IEEE-33 constraints" if network_mode == "ieee33" else "Aggregate 60% hosting"
    availability = "Data-driven availability" if availability_mode == "data_driven" else "Sim-default availability"
    return f"{fleet} / {network} / {availability}"


def _summarize_condition(
    payloads: list[tuple[str, dict[str, Any]]],
    fleet_mode: str,
    network_mode: str,
    availability_mode: str,
) -> dict[str, Any]:
    values = {
        strategy: np.asarray([
            float(payload["results"][strategy]["mean_reduction_pct"])
            for _, payload in payloads
        ])
        for strategy in FIG4D_ORDER
    }
    eps = values["eps_broadcast"]
    central = values["centralized_optimal"]
    best_baseline = np.max(np.column_stack([values[key] for key in BASELINE_KEYS]), axis=1)
    eps_to_central = np.divide(
        eps,
        central,
        out=np.ones_like(eps),
        where=central > 1e-12,
    )
    datasets = [dataset for dataset, _ in payloads]
    below_90 = np.flatnonzero(eps + 1e-9 < 0.9 * best_baseline)
    weakest_index = int(np.argmin(eps_to_central))
    return {
        "fleet_mode": fleet_mode,
        "network_mode": network_mode,
        "availability_mode": availability_mode,
        "title": _condition_title(fleet_mode, network_mode, availability_mode),
        "datasets": datasets,
        "dataset_count": len(payloads),
        "strategy_mean_pct": {key: float(np.mean(value)) for key, value in values.items()},
        "strategy_std_pct": {key: float(np.std(value)) for key, value in values.items()},
        "eps": {
            "mean_pct": float(np.mean(eps)),
            "min_pct": float(np.min(eps)),
            "max_pct": float(np.max(eps)),
            "beats_or_ties_best_baseline_count": int(np.sum(eps + 1e-9 >= best_baseline)),
            "within_90pct_best_baseline_count": int(np.sum(eps + 1e-9 >= 0.9 * best_baseline)),
            "mean_gap_to_best_baseline_pct_points": float(np.mean(eps - best_baseline)),
            "mean_centralized_ratio": float(np.mean(eps_to_central)),
            "min_centralized_ratio": float(np.min(eps_to_central)),
            "below_90pct_best_baseline_datasets": [datasets[index] for index in below_90],
            "weakest_centralized_ratio_dataset": datasets[weakest_index],
        },
    }


def _plot_overall(summary: dict[str, Any], output: Path) -> Path:
    means = np.asarray([summary["strategy_mean_pct"][key] for key in FIG4D_ORDER])
    stds = np.asarray([summary["strategy_std_pct"][key] for key in FIG4D_ORDER])
    fig, ax = plt.subplots(figsize=(13.2, 5.3), constrained_layout=True)
    x = np.arange(len(FIG4D_ORDER))
    bars = ax.bar(
        x,
        means,
        yerr=stds,
        capsize=4,
        color=[FIG4D_COLORS[key] for key in FIG4D_ORDER],
        edgecolor="#555555",
        linewidth=0.5,
    )
    for strategy, bar, value in zip(FIG4D_ORDER, bars, means):
        value_label = f"Mean\n{value:.1f}%"
        if strategy == "centralized_optimal":
            value_label = f"Upper bound\nMean\n{value:.1f}%"
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            max(1.0, value * 0.5),
            value_label,
            ha="center",
            va="center",
            fontsize=8,
            fontweight="bold",
            color="white" if value > 20 else "#222222",
        )
    ax.set_xticks(x)
    ax.set_xticklabels([FIG4D_LABELS[key] for key in FIG4D_ORDER], fontsize=8)
    ax.set_ylabel("Curtailment reduction (%)")
    ax.axhline(100, color="#dddddd", linewidth=0.8)
    ax.set_ylim(0, 120)
    ax.grid(alpha=0.2, axis="y")
    ax.text(0.0, 1.01, "Complexity: N = devices; H = MPC prediction horizon.", transform=ax.transAxes, fontsize=8, color="#444444")
    ax.set_title(
        f"Overall Figure 4D: {summary['title']}\n"
        f"mean +/- dataset SD, n={summary['dataset_count']} datasets",
        fontsize=11,
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return output


def _set_text_style(frame: Any, size: float, *, bold: bool = False, color: str = "222222") -> None:
    for paragraph in frame.paragraphs:
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color)


def _add_title(slide: Any, title: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.18), Inches(12.4), Inches(0.68))
    box.text_frame.text = title
    _set_text_style(
        box.text_frame,
        18 if len(title) > 60 else 23,
        bold=True,
        color="1F2937",
    )


def _add_footer(slide: Any, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.5), Inches(6.4), Inches(12.3), Inches(0.82))
    box.text_frame.text = text
    box.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    _set_text_style(box.text_frame, 9.5, color="374151")


def _add_figure3_slide(prs: Presentation, result_root: Path) -> None:
    dataset = _dataset_name(result_root)
    figure = result_root / "Figs" / "fig3_scaling_heterogeneity.png"
    threshold = _read_json(result_root / "data" / "n_threshold.json")
    n_values = [int(value) for value in threshold["N_values"]]
    missing = [value for value in (1, 5, 15, 30) if value not in n_values]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_title(slide, dataset)
    slide.shapes.add_picture(str(figure), Inches(0.35), Inches(0.9), width=Inches(12.65))
    n95 = threshold.get("threshold_N_95_interpolated")
    n95_text = "未达到" if n95 is None else f"{float(n95):.1f}"
    if dataset == "opsd_household_data" and not missing:
        audit = "新增 N=1,5,15,30 均已纳入；N=15/30 为 profile 复用审计"
    else:
        audit = "新增 N=1,5,15,30 均已纳入" if not missing else f"物理源限制，缺少 N={','.join(map(str, missing))}"
    _add_footer(slide, f"Figure 3 | 对数插值 N95={n95_text} | {audit}")


def _boundary_text(summary: dict[str, Any]) -> str:
    eps = summary["eps"]
    count = summary["dataset_count"]
    below = eps["below_90pct_best_baseline_datasets"]
    boundary = "无" if not below else ", ".join(below)
    return (
        f"内容：9 种算法的跨数据集弃电降低率均值，误差条为数据集间标准差（n={count}）。 "
        f"EPS 平均 {eps['mean_pct']:.1f}%（范围 {eps['min_pct']:.1f}%–{eps['max_pct']:.1f}%）；"
        f"在 {eps['beats_or_ties_best_baseline_count']}/{count} 个数据集不低于最佳非集中 baseline，"
        f"在 {eps['within_90pct_best_baseline_count']}/{count} 个达到其 90%；"
        f"相对集中式上界的平均比例为 {100 * eps['mean_centralized_ratio']:.1f}%。 "
        f"有效边界：低于最佳 baseline 90% 的数据集为 {boundary}；"
        f"最弱集中式比例出现在 {eps['weakest_centralized_ratio_dataset']}"
        f"（{100 * eps['min_centralized_ratio']:.1f}%）。"
    )


def _add_overall_slide(prs: Presentation, summary: dict[str, Any], figure: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_title(slide, f"Overall Figure 4D | {summary['title']}")
    slide.shapes.add_picture(str(figure), Inches(0.42), Inches(0.95), width=Inches(12.45))
    _add_footer(slide, _boundary_text(summary))


def generate(results_root: Path, output: Path) -> dict[str, Any]:
    roots = discover_result_roots(results_root)
    if len(roots) != 15:
        raise RuntimeError(f"预期 15 个数据集，实际发现 {len(roots)} 个")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for root in roots:
        _add_figure3_slide(prs, root)

    summaries = []
    figures = []
    for fleet_mode in FLEET_MODES:
        for network_mode in NETWORK_MODES:
            for availability_mode in AVAILABILITY_MODES:
                stem = output_stem(fleet_mode, network_mode, availability_mode)
                payloads = load_condition_payloads(results_root, stem)
                if not payloads:
                    raise RuntimeError(
                        f"缺少总体条件: {fleet_mode}/{network_mode}/{availability_mode}"
                    )
                summary = _summarize_condition(
                    payloads, fleet_mode, network_mode, availability_mode
                )
                figure = results_root / f"overall_fig4d_{stem}.png"
                _plot_overall(summary, figure)
                _add_overall_slide(prs, summary, figure)
                summaries.append(summary)
                figures.append(str(figure))

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    report = {
        "pptx": str(output),
        "slides": len(prs.slides),
        "figure3_slides": len(roots),
        "overall_figure4d_slides": len(summaries),
        "overall_figures": figures,
        "conditions": summaries,
    }
    report_path = results_root / "fig3_fig4d_ppt_summary.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--results-root", type=Path, default=ROOT / "results")
    parser.add_argument(
        "--output",
        type=Path,
        default=ROOT / "docs/fig3_fig4d_effective_boundary.pptx",
    )
    args = parser.parse_args()
    print(json.dumps(generate(args.results_root, args.output), ensure_ascii=False))


if __name__ == "__main__":
    main()
