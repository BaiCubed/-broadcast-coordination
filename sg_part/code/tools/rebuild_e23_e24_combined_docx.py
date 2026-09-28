"""根据 E23-E24 最终实验结构重建汇总 Word 文档。"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
import json


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "results/outputs/E23_E24_COMBINED.docx"


FIGURES = [
    ("E23连续压力实际效果", "results/E23/figures/e23_relative_effect_curves.png", "横坐标依次为请求强度 q、线路容量压力 λ_line 和空间集中比例 κ；纵坐标为可交付弃电降低率。曲线显示各算法在 IEEE-69 压力扫描下的实际执行效果。"),
    ("E23相对性能与物理风险", "results/E23/figures/e23_relative_advantage_curves.png", "横坐标为三条压力轴，纵坐标为 EPS 相对 baseline 的百分点差值。该图用于判断 EPS 何时相对落后。"),
    ("E23物理约束风险", "results/E23/figures/e23_physical_violation_curves.png", "纵坐标为安全裁剪前新增线路、电压或变压器违规比例；它与实际可交付效果共同描述网络风险。"),
    ("E23空间不平衡指标", "results/E23/figures/imbalance_metric_summary.png", "横坐标为 H0-H5，纵坐标为 I_N、I_C、I_T 和 I_H。该图先验证空间条件确实产生设计中的密度、容量、类型和网络耦合差异。"),
    ("E23密度与效果", "results/E23/figures/effect_vs_density_imbalance.png", "横坐标为设备密度不平衡 I_N，纵坐标为可交付弃电降低率；颜色表示算法，点形表示 H 条件。"),
    ("E23类型与效果", "results/E23/figures/effect_vs_type_heterogeneity.png", "横坐标为类型异质性 I_T，纵坐标为可交付弃电降低率；用于分析区域行为类型差异的影响。"),
    ("E23网络耦合与效果", "results/E23/figures/safety_vs_headroom_coupling.png", "横坐标为设备密度与低裕度的相关系数 I_H，纵坐标为可交付弃电降低率；用于比较 H4 与 H5 的不利和有利空间匹配。"),
    ("E23压力下响应跟踪 R²", "results/E23/figures/e23_r2_pressure_curves.png", "横坐标为压力轴，纵坐标为响应跟踪 R²；R² 衡量实际吸收曲线对目标弃电变化的解释程度。"),
    ("E23空间条件下响应跟踪 R²", "results/E23/figures/e23_r2_spatial_conditions.png", "横坐标为 H0-H5，纵坐标为响应跟踪 R²；用于分析空间异质性对时间跟踪的影响。"),
    ("E24 IEEE-123目标域训练 R²", "results/E24/trained_eps_ieee123_direct/figures/ieee123_direct_training_r2.png", "横坐标为数据集，纵坐标为 IEEE-123 目标域验证集 R²；虚线为 R²=0.95，审计直接训练模型质量。"),
    ("E24 IEEE-123各算法效果", "results/E24/figures/e24_ieee123_effect_comparison.png", "四个面板为 S0-S3，横坐标为算法，纵坐标为网络安全层执行后的可交付弃电降低率。"),
    ("E24 IEEE-123网络安全审计", "results/E24/figures/e24_ieee123_safety_audit.png", "红柱为新增违规请求比例，蓝柱为网络接受率；纵坐标为百分比，用于解释效果下降的网络原因。"),
]


def main() -> None:
    document = Document()
    heading = document.add_heading("E23-E24：IEEE-69 与 IEEE-123 配电网验证", 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("E23 和 E24 的主要实验都在目标网络上独立训练 EPS，并在同一目标网络上测试。E20/E21 pooled、mixed 和跨网络迁移模型单独作为附加实验，用于测量跨域泛化。")

    document.add_heading("一、实验结构", level=1)
    document.add_paragraph("E23 的主要网络是 IEEE-69，使用逐数据集直接训练 EPS，扫描请求强度、线路容量、空间集中和 H0-H5 空间异质性。E24 的主要网络是 IEEE-123，同样为每个数据集独立训练 EPS，并在目标网络上执行四个空间和容量场景。IEEE-123 完整三相 OpenDSS 校核由 run_ieee123_opendss_validation.py 独立执行。")

    document.add_heading("二、H0-H5 定义", level=1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "条件"
    table.rows[0].cells[1].text = "定义"
    definitions = [
        ("H0", "区域内完全混合；类型比例相同，设备数量按区域基础负荷分配。"),
        ("H1", "类型分区；不同区域由住宅、建筑、热负荷、配网/AMI、DER和充电类型主导。"),
        ("H2", "密度分区；类型比例固定，设备密度按25%、50%、75%向低裕度区域集中。"),
        ("H3", "类型和密度联合不均衡；远端低裕度区域集中住宅、热泵和充电。"),
        ("H4", "不利网络耦合；高功率类型位于最低裕度区域。"),
        ("H5", "有利网络耦合；高功率类型位于较高裕度区域。"),
    ]
    for condition, definition in definitions:
        cells = table.add_row().cells
        cells[0].text = condition
        cells[1].text = definition
    document.add_paragraph("H2 的设备密度指标为 I_N = sqrt(sum_r q_r * (x_r/q_r - 1)^2)。I_N=0 表示与基础负荷权重完全匹配；数值越大表示设备越集中。I_C、I_T、I_H 和 I_CH 分别衡量可控容量、类型分布、低裕度耦合和可控容量耦合。")

    document.add_heading("三、图片与指标", level=1)
    e24_direct_manifest = ROOT / "results/E24/trained_eps_ieee123_direct/manifest.json"
    direct_training_complete = False
    if e24_direct_manifest.is_file():
        direct_training_complete = json.loads(e24_direct_manifest.read_text(encoding="utf-8")).get("status") == "completed"
    for title, relative_path, explanation in FIGURES:
        if relative_path.startswith("results/E24/trained_eps_ieee123_direct/") and not direct_training_complete:
            continue
        path = ROOT / relative_path
        if not path.is_file():
            continue
        document.add_heading(title, level=2)
        document.add_picture(str(path), width=Inches(6.6))
        document.add_paragraph(explanation)

    document.add_heading("四、主结论范围", level=1)
    document.add_paragraph("E23 的主要结论来自 IEEE-69 直接训练 EPS；E24 的主要结论来自 IEEE-123 直接训练 EPS。可交付弃电降低率、网络接受率、最低电压、线路负载率、变压器负载率和 R² 分开报告。Centralized greedy UB 是设备层理论上界，不承担网络安全保证。IEEE-123 多相结果只有在安装 opendssdirect.py 并提供官方 IEEE123Master.dss 后生成。")
    document.add_heading("五、运行入口", level=1)
    document.add_paragraph("python -m src.extra.ieee33_device_day_simulation.figures.run_e24_ieee123_direct_training --seed-count 30 --workers 2\npython -m src.extra.ieee33_device_day_simulation.figures.run_e24_ieee123_audit --model-source results/E24/trained_eps_ieee123_direct/models --seed-count 30\npython -m src.extra.ieee33_device_day_simulation.figures.run_ieee123_opendss_validation --master path/to/IEEE123Master.dss")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)


if __name__ == "__main__":
    main()
