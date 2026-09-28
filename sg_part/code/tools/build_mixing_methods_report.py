from pathlib import Path
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'outputs'; FIG=OUT/'mixing_methods_figures'; FIG.mkdir(exist_ok=True)
def font(run,size=10,bold=False):
    run.font.name='Arial'; run._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial'); run.font.size=Pt(size); run.bold=bold
def para(doc,text,size=10,bold=False):
    p=doc.add_paragraph(); r=p.add_run(text); font(r,size,bold); p.paragraph_format.space_after=Pt(6); return p
def single_image(src,name,label):
    im=Image.open(ROOT/src).convert('RGB'); im.thumbnail((1200,760)); c=Image.new('RGB',(1240,820),'white'); c.paste(im,((1240-im.width)//2,20)); ImageDraw.Draw(c).text((25,785),label,fill='black'); out=FIG/name; c.save(out,dpi=(180,180)); return out
def caption(doc,text,ref):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(text+' Reference: '+ref); font(r,9); r.italic=True

doc=Document()
for s in doc.styles:
    if hasattr(s,'font'): s.font.name='Arial'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial'); s.font.size=Pt(10)
h=doc.add_heading('Mixed-Fleet Data Construction and Results',0)
for r in h.runs: font(r,14,True)
para(doc,'This report describes the two data-construction protocols used in the experiments, with emphasis on how heterogeneous datasets are matched, reweighted, and evaluated in one physical fleet.')
table=doc.add_table(rows=1,cols=4); table.style='Table Grid'; table.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,x in enumerate(['Scenario','Datasets and nominal shares','Mixing rule','Spatial or stress purpose']):
    c=table.rows[0].cells[i]; r=c.paragraphs[0].add_run(x); font(r,10,True)
scenarios=[
('Pairwise mixed','Any two of the 15 source datasets, 50% and 50%','Equal allocation across two sources; 105 unordered pairs','Robustness to source diversity and pair-specific composition'),
('S1-A','LCL 33.33%; CAMSL 33.33%; Irish 33.33%','Equal allocation across three sources','Similar residential profiles'),('S1-B','BDG2 33.33%; HEAPO 33.33%; SGSC 33.33%','Equal allocation across three sources','Building, heat-pump, and residential diversity'),('S2-A','EU-35297 40%; EU-8087 25%; EU-Rural 15%; Norway 10%; GoiEner 10%','Weighted allocation across five sources','Large-scale LV and AMI mixture'),('S2-B','BDG2 35%; LCL 25%; Danish 15%; HEAPO 15%; CEC 10%','Weighted allocation across five sources','Building, residential, thermal, and DER mixture'),('S3-A','LCL 18%; SGSC 14%; CAMSL 12%; Irish 10%; GoiEner 8%; Norway 8%; EU-Rural 7%; EU-35297 10%; EU-8087 8%; HEAPO 5%','Weighted allocation across ten sources','Broad electricity mix with thermal perturbation'),('S3-B','BDG1 10%; BDG2 15%; CEC 10%; Danish 10%; HEAPO 10%; LCL 10%; SGSC 10%; EU-Rural 8%; Norway 12%; OPSD 5%','Weighted allocation across ten sources','Multi-sector heterogeneity'),('S4-A','All 15 datasets, 6.67% each','Equal allocation across all sources','Maximum source heterogeneity'),('S4-B','Building 20%; residential 20%; thermal 20%; network 20%; DER 20%','Equal allocation by source type','Balanced type composition'),('S5-A','Residential 65%; network 20%; building 5%; thermal 5%; DER 5%','Long-tail weighted allocation','Residential-dominant population'),('S5-B','Building 40%; thermal 30%; residential 15%; network 10%; DER 5%','Long-tail weighted allocation','Building and thermal dominant population'),('S5-C','Network 45%; DER 25%; residential 20%; building 5%; thermal 5%','Long-tail weighted allocation','Network and DER dominant population'),('S6-A','EU-35297 35%; EU-8087 25%; EU-Rural 20%; Norway 20%','Weighted allocation plus zone placement','Spatially clustered LV networks'),('S6-B','BDG2 35%; HEAPO 20%; Danish 15%; LCL 20%; SGSC 10%','Weighted allocation plus zone placement','Cross-sector feeder congestion'),('S6-C','CEC 30%; SGSC 30%; OPSD 5%; Irish 20%; Norway 15%','Weighted allocation plus zone placement','Spatial DER and bidirectional mix')]
for row in scenarios:
    cells=table.add_row().cells
    for i,x in enumerate(row): cells[i].text=''; r=cells[i].paragraphs[0].add_run(x); font(r); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
para(doc,'How the two protocols construct data',10,True)
para(doc,'Pairwise mixed construction combines two source datasets at equal weight in one physical fleet. Across 105 unordered pairs, it exposes the controller to a broad range of source combinations and tests whether performance remains stable when the data distribution changes. This is the robustness arm of the study: it measures sensitivity to source diversity without introducing a new pair-specific training model. Each source contributes its own normalized daily load shape and data-driven availability records. The fleet generator converts the requested shares into integer device counts. Non-unique sampling allows bootstrap reuse when a source has too few independent records; source-unique sampling forbids reuse and reduces the feasible fleet size when necessary.')
para(doc,'The overall design is deliberately broad: pairwise enumeration covers systematic two-source diversity, while the fourteen multi-source scenarios cover focused, balanced, long-tail, sector-mixed, and spatially clustered populations. Together they provide evidence across source count, source type, mixture weight, dataset size, and feeder placement. The purpose is to assess algorithm robustness over a wide and heterogeneous operating envelope rather than to optimize for one representative dataset.')
para(doc,'The fourteen predefined scenarios use the same device-level construction but replace the equal two-source rule with explicit multi-source weights. Training fleets come from training partitions, validation fleets from validation partitions, and test fleets from test partitions. The controller receives common signal features, not dataset labels or source identifiers. For S6, source groups are additionally assigned to specified IEEE-33 zones and buses so that identical source proportions can produce different feeder interactions.')
para(doc,'Matching heterogeneous inputs',10,True)
para(doc,'The datasets do not all measure the same channel. The common matching layer retains each source profile’s temporal demand shape, availability information, day identifier, and location metadata. A shared simulation layer supplies comparable device capacity, power-rate, initial state of charge, state of health, and dispatch rules. Residential meters, building electricity, heat-demand proxies, low-voltage and AMI profiles, and integrated DER records therefore enter the same fleet without claiming that every source contains measured photovoltaic curtailment.')
para(doc,'Dataset-size differences are handled through explicit integer allocation and audit fields for available sources, selected unique sources, bootstrap reuse, and duplicate counts. Small sources, especially OPSD, are capped at five percent in most non-equal stress scenarios so that repeated records cannot dominate the mixture.')
atlas=single_image('results/E21/curtailment_baseline_supplement/figures/scenario_baseline_atlas.png','scenario_baseline_atlas.png','Scenario-level algorithm results'); doc.add_picture(str(atlas),width=Inches(6.5)); caption(doc,'Figure 1. Performance across all fourteen source compositions.','Repository scenario atlas description'); para(doc,'The atlas shows that source composition changes the achievable curtailment reduction, and that feeder constraints can materially change the magnitude and ranking of results. Averaging single-source results would hide these composition effects.')
pen=single_image('results/E21/curtailment_baseline_supplement/figures/scenario_network_penalty.png','scenario_network_penalty.png','Network penalty by mixed scenario'); doc.add_picture(str(pen),width=Inches(6.5)); caption(doc,'Figure 2. Loss of deliverable benefit after IEEE-33 feeder screening.','Repository scenario network-penalty description'); para(doc,'The network-penalty figure demonstrates the role of spatial matching. Clustered high-power or DER groups can lose much of their aggregate benefit after feeder screening, while other compositions are less affected. Spatial placement is therefore part of the data construction.')
path=OUT/'data_mixing_methods_report.docx'; doc.save(path); print(path)
