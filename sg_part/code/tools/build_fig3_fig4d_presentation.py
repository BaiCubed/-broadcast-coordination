"""生成 Figure 3 与最终 Figure 4D 的实验汇报 PPTX。"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


DATASET_SUFFIX = "_ieee33_real_load"
FIG3_RELATIVE = Path("coverage_fix/network_constrained_new/Figs/fig3_scaling_heterogeneity.png")
THRESHOLD_RELATIVE = Path("coverage_fix/network_constrained_new/data/n_threshold.json")


def _add_title(slide, title: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.2), Inches(12.4), Inches(0.55))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(25)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor(32, 43, 56)


def _add_text(slide, text: str, left: float, top: float, width: float, height: float, size: int = 14) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()
    for index, line in enumerate(text.splitlines()):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = RGBColor(48, 55, 64)
        paragraph.space_after = Pt(7)


def _r2_summary(payload: dict[str, Any]) -> str:
    per_n = payload["per_N"]
    requested = [1, 5, 15, 30]
    values = []
    for fleet_size in requested:
        row = per_n.get(str(fleet_size))
        values.append(f"N={fleet_size}: R²={row['r2']:.3f}" if row else f"N={fleet_size}: 不适用")
    threshold = payload.get("threshold_N_95_interpolated")
    threshold_text = "未达到 R²=0.95" if threshold is None else f"插值 N₉₅={threshold:.1f}"
    return "\n".join(values + [threshold_text])


def _fig3_boundary_text(payload: dict[str, Any]) -> str:
    threshold = payload.get("threshold_N_95_interpolated")
    maximum = max(int(value) for value in payload["N_values"])
    if threshold is None:
        return f"有效边界：在已审计 N≤{maximum} 范围内，聚合响应尚未达到 R²≥0.95。"
    return (
        f"有效边界：N 约达到 {threshold:.1f} 后，固定条件下的聚合响应进入 R²≥0.95 可预测区。"
        "该边界只证明响应可估计，不等同于 EPS 在弃电指标上必然最优。"
    )


def _add_dataset_slide(prs: Presentation, dataset_root: Path) -> None:
    dataset = dataset_root.name.removesuffix(DATASET_SUFFIX)
    figure = dataset_root / FIG3_RELATIVE
    threshold = json.loads((dataset_root / THRESHOLD_RELATIVE).read_text(encoding="utf-8"))
    population = json.loads((
        dataset_root / "coverage_fix/network_constrained_new/data/population_summary.json"
    ).read_text(encoding="utf-8"))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_title(slide, dataset)
    slide.shapes.add_picture(
        str(figure), Inches(0.5), Inches(0.82), width=Inches(12.3)
    )
    _add_text(slide, _r2_summary(threshold), 0.65, 5.08, 5.7, 1.95, 12)
    source_devices = int(population["source_devices"])
    boundary_text = _fig3_boundary_text(threshold)
    if source_devices < 30:
        boundary_text += (
            f"\n审计说明：仅 {source_devices} 个实测源；超过该数量的 N 是逻辑设备曲线复用，"
            "不代表独立实测设备。"
        )
    _add_text(slide, boundary_text, 6.55, 5.08, 6.1, 1.95, 12)


def _condition_description(condition: dict[str, Any], dataset_count: int) -> str:
    fleet = condition["fleet_mode"]
    network = condition["network_mode"]
    availability = condition["availability_mode"]
    fleet_text = (
        "固定 5,000 个逻辑设备，允许真实负荷曲线自助采样"
        if fleet == "fixed5000"
        else "唯一源审计，N=min(5,000, 数据集唯一源数)"
    )
    network_text = "聚合 60% hosting 约束" if network == "aggregate" else "IEEE-33 节点与线路约束"
    availability_text = "纯仿真默认离线/丢包率" if availability == "sim_default" else "数据驱动可用率"
    return f"样本：{dataset_count} 个适用数据集\n设备：{fleet_text}\n网络：{network_text}\n可用率：{availability_text}"


def _overall_boundary_text(summary: dict[str, Any]) -> str:
    means = summary["mean_reduction_pct"]
    eps = means["eps_broadcast"]
    centralized = means["centralized_optimal"]
    ranking = sorted(
        ((value, strategy) for strategy, value in means.items()), reverse=True
    )
    rank = 1 + ranking.index((eps, "eps_broadcast"))
    better = sum(
        row["eps_broadcast"] >= row["local_rules"]
        for row in summary["per_dataset_reduction_pct"].values()
    )
    return (
        f"EPS 平均弃电降低 {eps:.1f}%，九种算法中均值排名第 {rank}；"
        f"与 centralized optimal 相差 {centralized - eps:.1f} 个百分点。\n"
        f"EPS 在 {better}/{summary['dataset_count']} 个数据集上不低于 Local SOC rules。\n"
        "有效边界由两部分共同限定：Figure 3 的响应可预测性，以及本条件下网络约束和可用率造成的可实现性。"
    )


def _add_overall_slide(prs: Presentation, summary: dict[str, Any], results_root: Path) -> None:
    condition = summary["condition"]
    title = "Overall Figure 4D | " + " | ".join(
        str(condition[key]) for key in ("fleet_mode", "network_mode", "availability_mode")
    )
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_title(slide, title)
    figure = Path(summary["figure"])
    if not figure.is_absolute():
        figure = results_root.parent / figure
    slide.shapes.add_picture(
        str(figure), Inches(0.45), Inches(1.05), width=Inches(8.65)
    )
    _add_text(slide, _condition_description(condition, summary["dataset_count"]), 9.35, 1.0, 3.35, 2.5, 13)
    _add_text(slide, _overall_boundary_text(summary), 9.35, 3.65, 3.35, 2.8, 13)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--results-root", type=Path, default=Path("results"))
    parser.add_argument("--output", type=Path, default=Path("docs/fig3_fig4d_eps_boundary.pptx"))
    args = parser.parse_args()
    manifest_path = args.results_root / "overall_fig4d_final_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for dataset_root in sorted(args.results_root.glob(f"*{DATASET_SUFFIX}")):
        if (dataset_root / FIG3_RELATIVE).is_file():
            _add_dataset_slide(prs, dataset_root)
    for summary in manifest["conditions"]:
        _add_overall_slide(prs, summary, args.results_root)

    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.LEFT
    args.output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
