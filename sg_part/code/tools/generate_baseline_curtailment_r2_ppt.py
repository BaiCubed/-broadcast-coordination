"""生成 baseline、弃电降低与 R2 实验结果汇总 PPT。"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]


@dataclass(frozen=True)
class SlideSpec:
    title: str
    image: Path


def _specs() -> list[SlideSpec]:
    specs = [
        SlideSpec("全数据集算法弃电降低对比", ROOT / "results/overall_algorithm_curtailment_grouped_bar.png"),
        SlideSpec("全数据集算法弃电降低热图", ROOT / "results/overall_algorithm_curtailment_heatmap.png"),
        SlideSpec("E20 Aggregate baseline 对比", ROOT / "results/E20/figures/e20_algorithm_comparison_aggregate.png"),
        SlideSpec("E20 IEEE-33 baseline 对比", ROOT / "results/E20/figures/e20_algorithm_comparison_ieee33.png"),
        SlideSpec("E20 baseline 总体对比", ROOT / "results/E20/figures/e20_algorithm_comparison_overall.png"),
        SlideSpec("E20 迁移算法 R2 总图", ROOT / "results/E20/fig3a_transfer/fig3a_transfer_r2.png"),
        SlideSpec("E20 CAMSL 迁移 R2", ROOT / "results/E20/fig3a_transfer/fig3a_transfer_r2_camsl_japan_smart_meters.png"),
        SlideSpec("E20 EU-8k 迁移 R2", ROOT / "results/E20/fig3a_transfer/fig3a_transfer_r2_european_lv_urban_8087.png"),
        SlideSpec("E20 Irish 迁移 R2", ROOT / "results/E20/fig3a_transfer/fig3a_transfer_r2_irish_domestic_smart_meters.png"),
        SlideSpec("E20 OPSD 迁移 R2", ROOT / "results/E20/fig3a_transfer/fig3a_transfer_r2_opsd_household_data.png"),
        SlideSpec("E20 SGSC 迁移 R2", ROOT / "results/E20/fig3a_transfer/fig3a_transfer_r2_smart_grid_smart_city.png"),
        SlideSpec("E20 各算法 N95 对比", ROOT / "results/E20/fig3a_transfer/fig3a_transfer_n95.png"),
        SlideSpec("E20 N95 迁移规模代价", ROOT / "results/E20/fig3a_transfer/fig3a_transfer_n95_inflation.png"),
        SlideSpec("E21 Aggregate 弃电降低热图", ROOT / "results/E21/figures/e21_algorithm_heatmap_aggregate.png"),
        SlideSpec("E21 IEEE-33 弃电降低热图", ROOT / "results/E21/figures/e21_algorithm_heatmap_ieee33.png"),
        SlideSpec("E21 算法弃电降低总体对比", ROOT / "results/E21/figures/e21_algorithm_overall.png"),
        SlideSpec("E21 混合训练与池化训练对比", ROOT / "results/E21/figures/e21_eps_mix_vs_pool.png"),
        SlideSpec("E21 两两组合弃电降低矩阵", ROOT / "results/E21/pairwise_curtailment/figures/pairwise_curtailment_reduction.png"),
        SlideSpec("E21 两两组合弃电降低矩阵（仅混合）", ROOT / "results/E21/pairwise_curtailment/figures/pairwise_curtailment_reduction_mixed_only.png"),
        SlideSpec("E21 数据集配对效果分布", ROOT / "results/E21/pairwise_curtailment/figures/dataset_curtailment_reduction_distribution.png"),
        SlideSpec("E21 弃电机会与降低率", ROOT / "results/E21/pairwise_curtailment/figures/curtailment_opportunity_vs_reduction.png"),
        SlideSpec("E21 两两组合控制前后弃电", ROOT / "results/E21/pairwise_curtailment/figures/pairwise_curtailment_before_after_ranked.png"),
        SlideSpec("E21 设计场景 R2 曲线总图", ROOT / "results/E21/gamma_mixed_boundary/figures/e21_mixed_r2_group.png"),
    ]
    scenario_ids = (
        "S1-A", "S1-B", "S2-A", "S2-B", "S3-A", "S3-B", "S4-A",
        "S4-B", "S5-A", "S5-B", "S5-C", "S6-A", "S6-B", "S6-C",
    )
    specs.extend(
        SlideSpec(
            f"E21 设计场景 {scenario_id} 的 R2 曲线",
            ROOT / f"results/E21/gamma_mixed_boundary/figures/r2_curves/e21_mixed_r2_{scenario_id}.png",
        )
        for scenario_id in scenario_ids
    )
    specs.append(
        SlideSpec(
            "E21 两两组合 N95 汇总",
            ROOT / "results/E21/pairwise_r2/figures/pairwise_r2_n95_overall.png",
        )
    )
    specs.extend(
        SlideSpec(
            f"E21 两两组合 R2 曲线（{page}/9）",
            ROOT / f"results/E21/pairwise_r2/figures/r2_curve_pages/pairwise_r2_curves_page_{page:02d}.png",
        )
        for page in range(1, 10)
    )
    return specs


def _add_title(slide: object, title: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.38), Inches(0.16), Inches(12.57), Inches(0.55))
    frame = box.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = title
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.font.name = "Microsoft YaHei"
    paragraph.font.size = Pt(22 if len(title) <= 24 else 19)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor(31, 41, 55)


def _add_image(slide: object, path: Path) -> None:
    area_left = Inches(0.28)
    area_top = Inches(0.82)
    area_width = Inches(12.77)
    area_height = Inches(6.40)
    with Image.open(path) as image:
        width_px, height_px = image.size
    ratio = min(area_width / width_px, area_height / height_px)
    width = int(width_px * ratio)
    height = int(height_px * ratio)
    left = int(area_left + (area_width - width) / 2)
    top = int(area_top + (area_height - height) / 2)
    slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def generate(output: Path) -> int:
    specs = _specs()
    missing = [str(spec.image) for spec in specs if not spec.image.is_file()]
    if missing:
        raise FileNotFoundError("缺少 PPT 输入图片：\n" + "\n".join(missing))
    if len({spec.image for spec in specs}) != len(specs):
        raise RuntimeError("PPT 输入图片存在重复")

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = "Baseline、弃电降低与 R2 实验结果"
    presentation.core_properties.subject = "Results、E20 与 E21 实验图片汇总"
    for spec in specs:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _add_title(slide, spec.title)
        _add_image(slide, spec.image)

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    return len(specs)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=ROOT / "docs/baseline_curtailment_r2_results.pptx",
    )
    args = parser.parse_args()
    slide_count = generate(args.output)
    print(f"已生成 {args.output}，共 {slide_count} 页")


if __name__ == "__main__":
    main()
