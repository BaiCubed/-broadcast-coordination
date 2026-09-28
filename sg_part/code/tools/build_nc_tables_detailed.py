from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]

def set_run(r, size=8.5, bold=False, font='Times New Roman'):
    r.font.name=font; r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),font); r.font.size=Pt(size); r.bold=bold

def add_omml(p, expr, size=10):
    """Insert a native Word OMML equation (Cambria Math) into paragraph p."""
    oMath = OxmlElement('m:oMath')
    for token in expr.split(' '):
        mr = OxmlElement('m:r'); mt = OxmlElement('m:t'); mt.text = token; mr.append(mt); oMath.append(mr)
    p._p.append(oMath)
    for mr in oMath:
        for mt in mr:
            mt.set(qn('xml:space'),'preserve')

def equation(doc, expr, label=None):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(4)
    add_omml(p,expr)
    if label:
        r=p.add_run('   '+label); set_run(r,8.5)

def body(doc, text, boldlead=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.08
    if boldlead:
        r=p.add_run(boldlead); set_run(r,9.2,True)
    r=p.add_run(text); set_run(r,9.2)
    return p

def cell(cell,text,bold=False,size=6.8):
    cell.text=''; p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
    r=p.add_run(str(text)); set_run(r,size,bold); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tcPr=cell._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
    for side,val in [('top','55'),('start','55'),('bottom','55'),('end','55')]:
        e=OxmlElement('w:'+side); e.set(qn('w:w'),val); e.set(qn('w:type'),'dxa'); mar.append(e)
    tcPr.append(mar); bd=OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'8'); e.set(qn('w:color'),'000000'); bd.append(e)
    tcPr.append(bd)

def table(doc,title,heads,rows,widths,size=6.8):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(9); p.paragraph_format.space_after=Pt(5)
    r=p.add_run(title); set_run(r,12.5,True)
    t=doc.add_table(rows=1,cols=len(heads)); t.autofit=False
    tr=t.rows[0]; tr.height=Inches(.52); tr.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
    hp=tr._tr.get_or_add_trPr(); e=OxmlElement('w:tblHeader'); e.set(qn('w:val'),'true'); hp.append(e)
    for i,h in enumerate(heads): cell(tr.cells[i],h,True,size); tr.cells[i].width=Inches(widths[i])
    for row in rows:
        cells=t.add_row().cells; t.rows[-1].height=Inches(.38); t.rows[-1].height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
        for i,v in enumerate(row): cell(cells[i],v,False,size); cells[i].width=Inches(widths[i])
    doc.add_paragraph()

def doc_base():
    d=Document(); s=d.sections[0]; s.orientation=WD_ORIENT.LANDSCAPE; s.page_width,s.page_height=s.page_height,s.page_width; s.top_margin=s.bottom_margin=Inches(.52); s.left_margin=s.right_margin=Inches(.5); return d

datasets=[
('P00','NextGen','device-day battery','5 min','p_{i,t},q_{i,t},s_{i,t},E_i,P_i^max,a_{i,t},z_i,t','100 source devices; 3,000 device-days','battery and feeder reference'),
('P00b','data2','charging session','5 min','p_{i,t},E_i,P_i^max,a_{i,t},z_i,t','measured sessions; g_{i,t} counterfactual','bidirectional charging stress'),
('P01','BDG1','building electricity','1 h','p_{i,t},b_{i,t},z_i,t','507 buildings; no measured g_{i,t},s_{i,t}','building heterogeneity'),
('P02','Low Carbon London','household electricity','30 min','p_{i,t},b_{i,t},z_i,t','5,000 households; tariff metadata','large residential scale'),
('P03','BDG2','building electricity + solar','1 h','p_{i,t},g_{i,t},b_{i,t},z_i,t','1,570 buildings; g_{i,t} only for aligned subset','measured/counterfactual mixture'),
('P04','Danish heat meters','thermal demand','1 h','h_{i,t},b_{i,t},z_i,t','2,400 meters; h_{i,t} is thermal proxy','thermal response diversity'),
('P05','Smart Grid Smart City','load + generation','30 min','p_{i,t},g_{i,t},b_{i,t},z_i,t','600 customers; general/controlled load and gross/net generation','measured self-consumption'),
('P06','HEAPO','heat-pump electricity','15 min / daily','p_{i,t},b_{i,t},w_t,z_i,t','1,362 households; weather w_t; no measured storage state','heat-pump heterogeneity'),
('P07','GoiEner','smart-meter electricity','1 h','p_{i,t},b_{i,t},z_i,t,c_i','5,000 users; c_i is contract-power metadata','large-user diversity'),
('P08','European LV urban-8k','LV active-power profile','24/168 points','p_{i,t},q_{i,t},z_i,t,n_i','5,000 customers; GIS/feeder context n_i','urban feeder diversity'),
('P09','European LV rural','LV active/reactive profile','24/168 points','p_{i,t},q_{i,t},z_i,t,n_i','2,731 customers; rural topology','rural spatial diversity'),
('P10','European LV urban-35k','LV active/reactive profile','24/168 points','p_{i,t},q_{i,t},z_i,t,n_i','12,000 customers; urban topology','large urban network'),
('P11','Norway AMI','MV/LV AMI profile','1 h','p_{i,t},q_{i,t},g_{i,t},z_i,t,n_i,w_t','2,994 sources; measured import/export and network model','network-linked measured data'),
('P12','CAMSL-JP','smart-meter electricity','30 min','p_{i,t},b_{i,t},z_i,t,tariff_i','1,423 households; tariff groups and temperature','tariff/temporal response'),
('P13','Irish CER','import/export household','30 min','p_{i,t},g_{i,t},b_{i,t},z_i,t','2,904 consumers; active import and export','bidirectional household profiles'),
('P14','OPSD','household multi-channel','15 min','p_{i,t},g_{i,t},b_{i,t},s_{i,t},z_i,t','11 households; bootstrap limitation','small-source audit'),
('P15','COMPLETE-EC','community load/PV/BESS/EV','15 min','p_{i,t},g_{i,t},b_{i,t},s_{i,t},E_i,z_i,t','250 members; PV, BESS and EV fields','integrated prosumer test'),]
SOURCE_COUNTS={'NextGen':'100','data2':'proxy','BDG1':'507','Low Carbon London':'5000','BDG2':'1570','Danish heat meters':'2400','Smart Grid Smart City':'600','HEAPO':'1362','GoiEner':'5000','European LV urban-8k':'5000','European LV rural':'2731','European LV urban-35k':'12000','Norway AMI':'2994','CAMSL-JP':'1423','Irish CER':'2904','OPSD':'11','COMPLETE-EC':'250'}

urls=[('NextGen','project archive / local source documentation','parameter and feeder reference'),('data2','project archive / local source documentation','charging-session source'),('BDG1','https://github.com/buds-lab/the-building-data-genome-project','Miller and Meggers (2017), DOI 10.1016/j.egypro.2017.07.400'),('Low Carbon London','https://data.london.gov.uk/dataset/smartmeter-energy-use-data-in-london-households','London Datastore / UK Power Networks'),('BDG2','https://github.com/buds-lab/building-data-genome-project-2','Miller et al. (2020), DOI 10.1038/s41597-020-00712-x'),('Danish heat meters','https://zenodo.org/records/6563114','Schaffer et al. (2022), DOI 10.1038/s41597-022-01502-3'),('Smart Grid Smart City','https://data.gov.au/data/dataset/4e21dea3-9b87-4610-94c7-15a8a77907ef','Australian SGSC project'),('HEAPO','https://zenodo.org/records/15056919','arXiv:2503.16993'),('GoiEner','https://zenodo.org/records/7362094','WHY project / GoiEner'),('European LV urban-8k','https://data.mendeley.com/datasets/685vgp64sm/1','DOI 10.17632/685vgp64sm.1'),('European LV rural / urban-35k','https://data.mendeley.com/datasets/gspyzvvrhm/2','DOI 10.17632/gspyzvvrhm.2'),('Norway AMI','https://data.mendeley.com/datasets/jv3rz8k35r/1','Maree et al., AMI smart-meter dataset'),('CAMSL-JP','https://data.mendeley.com/datasets/cmpsyncmmk/1','Kiguchi et al., CAMSL dataset'),('Irish CER','https://figshare.com/articles/dataset/6_150_years_1_month_17_days_and_8_hours_of_anonymised_electricity_smart_meter_data/31851922','DOI 10.6084/m9.figshare.31851922'),('OPSD','https://open-power-system-data.org/data-sources','Open Power System Data household data'),('COMPLETE-EC','https://zenodo.org/records/7602546','Faia et al. (2023), DOI 10.1016/j.dib.2023.109218')]

dataset_detail_en={
'NextGen':'NextGen is the device-day calibration population. It supplies battery response trajectories, state variables, power limits and feeder assignments needed to test the complete simulation chain. Its source count is approximately 100, whereas the 3,000 device-days are repeated source-day records. We therefore use NextGen for parameter, snapshot and network-mechanism validation, not as evidence that 3,000 independent field devices were observed.',
'data2':'data2 contains five-minute charging sessions. The session index is treated as i, the requested or delivered charging power as p_{i,t}, and session availability as a_{i,t}. It is useful for bidirectional charging stress tests because session start and end times create explicit availability boundaries. External generation g_{i,t} is counterfactual in this adapter and is not reported as measured PV.',
'BDG1':'BDG1 provides hourly building electricity trajectories p_{i,t} for 507 buildings. The source has strong cross-building load-shape diversity but does not provide measured battery SOC or PV generation. Battery variables s_{i,t}, E_i and P_i^max are therefore bounded model parameters when storage dispatch is simulated, while b_{i,t} is derived from the building load.',
'Low Carbon London':'Low Carbon London contributes half-hourly household electricity and tariff metadata. It is a large residential population with repeated daily structure and household-level heterogeneity. The dataset does not provide a measured battery state, so s_{i,t}, E_i and P_i^max are simulation-side storage variables rather than observations.',
'BDG2':'BDG2 adds measured solar channels to building electricity. Only buildings with aligned electricity and solar records receive measured g_{i,t}; other runs retain the load trajectory and use an explicitly labelled counterfactual input. This distinction allows measured-input sensitivity to be separated from the broader building-load experiment.',
'Danish heat meters':'The Danish heat-meter population measures thermal energy h_{i,t}. It is valuable for testing whether population-scale regularity extends to thermal demand, but h_{i,t} is not relabelled as electric power. An electrical interpretation requires a stated heat-pump or COP mapping, and the present table keeps that distinction visible.',
'Smart Grid Smart City':'SGSC contains general load, controlled load, gross generation and net generation fields. It therefore supports a direct comparison between b_{i,t}, g_{i,t} and the resulting net response. Its smaller source inventory makes it especially useful for provenance-sensitive measured-input and self-consumption analyses.',
'HEAPO':'HEAPO records household heat-pump electricity at 15-minute and daily resolutions and includes weather covariates w_t. The weather field can explain common temporal forcing, while p_{i,t} remains the measured electrical demand. Storage state and battery capacity are not measured and must remain explicit model variables.',
'GoiEner':'GoiEner supplies hourly smart-meter demand for a large user population together with contract-power metadata c_i and a self-generation type flag. The flag is categorical metadata and is not treated as a generation trace g_{i,t}. This makes GoiEner a load-diversity source rather than a measured PV source.',
'European LV urban-8k':'The European LV urban-8k source provides active-power profiles p_{i,t} together with feeder or GIS context n_i. It tests whether urban customer diversity remains visible after spatial assignment. Reactive power is not the primary EPS input in this adapter but is retained where available for network audit.',
'European LV rural':'The rural European LV source provides paired p_{i,t} and q_{i,t} profiles on a rural topology. The joint P/Q structure matters for voltage and branch loading, so the source is used both for customer diversity and for spatially heterogeneous network realization.',
'European LV urban-35k':'The urban-35k source supplies the largest customer inventory in the table and paired active/reactive profiles. Its value is not only N: it tests whether a large urban source with explicit network context changes the relationship between aggregate predictability and network acceptance.',
'Norway AMI':'Norway AMI combines hourly measured active import/export, reactive power, weather and an anonymized MV/LV network model. It is the clearest source for connecting p_{i,t}, q_{i,t}, g_{i,t} and n_i in a network-linked measured setting. The measured-only restriction prevents counterfactual input from being presented as observation.',
'CAMSL-JP':'CAMSL-JP supplies half-hourly household consumption with tariff-group and temperature context. It is used to examine temporal response differences under tariff exposure. No external generation trace is assumed, so g_{i,t} remains absent unless explicitly introduced as a counterfactual experiment.',
'Irish CER':'Irish CER records import and export energy separately at 30-minute resolution. This makes the sign and direction of p_{i,t} explicit and supports bidirectional household-profile tests. The data are anonymized consumer files; gaps and outage periods are retained as provenance conditions rather than silently imputed into a new source population.',
'OPSD':'OPSD contains multiple household channels, including load, PV, storage, EV and heat-pump fields, but only 11 independent households. It is therefore a variable-rich, source-poor stress case. Its role is to expose the difference between a rich feature vector and a statistically adequate number of independent sources.',
'COMPLETE-EC':'COMPLETE-EC describes an integrated community with load, PV, battery and EV fields and explicit capacity or SOC variables for many assets. It is useful for multi-technology coupling and spatial community tests, but its 250-member scale means that results should be interpreted as an integrated-community case rather than a large-population source inventory.'}

def make(lang):
    d=doc_base(); title='Supplementary Tables S1, S4 and S5' if lang=='en' else '补充表 S1、S4 和 S5'
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(title); set_run(r,16,True)
    if lang=='en':
        body(d,'This Supplementary Information provides the detailed data, control and network definitions used to interpret the population-scale experiments. The tables are deliberately compact; the accompanying text defines the symbols, provenance boundaries, algorithmic assumptions and physical limits required for reproducibility.')
        rows=[[a,b,c,e,f,g,h] for a,b,c,e,f,g,h in datasets]
        table(d,'Supplementary Table S1. Empirical population characteristics, standardized variables and data features.', ['ID','Dataset','Resource / signal','Resolution','Variables','Provenance / limit','N_src','Data feature'],[[a,b,c,e,f,g,SOURCE_COUNTS[b],h] for a,b,c,e,f,g,h in rows],[.34,.92,1.0,.55,2.1,1.45,.55,1.35],5.4)
        body(d,'The notation is common to all datasets. The index i identifies a resource or session, t identifies a dispatch interval, and τ is the sampling interval. The active and reactive powers are p_{i,t} and q_{i,t}; s_{i,t} is state of charge; E_i is usable energy capacity; P_i^max is the maximum active-power magnitude; a_{i,t}∈{0,1} is availability; z_i is the assigned feeder or zone; n_i is network or GIS context; g_{i,t} is external generation or charging input; b_{i,t} is baseline demand; and w_t is a common weather covariate. The population aggregate is defined by the following Word equation.')
        equation(d,'P_t = Σ_i p_{i,t}   ,   Q_t = Σ_i q_{i,t}   ,   N_src = |{source_i}|','aggregate active/reactive response and independent-source count')
        for name in ['NextGen','data2','BDG1','Low Carbon London','BDG2','Danish heat meters','Smart Grid Smart City','HEAPO','GoiEner','European LV urban-8k','European LV rural','European LV urban-35k','Norway AMI','CAMSL-JP','Irish CER','OPSD','COMPLETE-EC']:
            body(d,dataset_detail_en[name],boldlead=name+'. ')
        body(d,'Across the table, N_src is intentionally separated from the number of rows after resampling. A source can generate multiple device-days or sessions, but those records inherit source-level dependence. The comparison therefore distinguishes the breadth of observed variables from the breadth of independent empirical replication. This is the basis for interpreting the scale transition as a cross-source result rather than as a consequence of one unusually large file.')
        s4=[['Open-loop','p_{i,t}=u_t','r_t,τ','u_t only','O(1)','No adaptation to s_{i,t},a_{i,t}'],['Local SOC','p_{i,t}=clip(u_t;s_{i,t},P_i^max)','s_{i,t},P_i^max,a_{i,t}','local s_{i,t}','O(1)/i','threshold and heterogeneity sensitivity'],['MPC','min Σ_t(P_t-r_t)^2 s.t. g(x)≤0','r_t,s_i,E_i,P_i^max,q_i,V_k,L_l,T','states + forecasts','O(HN); solver','telemetry and model burden'],['Mean-field','p_{i,t}=π(x_{i,t},m_t), m_t=N^-1Σ_iφ(x_i)','x_i=(s_i,a_i,p_i,q_i),φ','moments + local x_i','O(N)+O(1)','closure under correlation/drift'],['Virtual battery','P_t=clip(r_t,P_min(m_t),P_max(m_t))','s_i,E_i,P_i^max,a_i,t','aggregate envelope','O(1) after fit','envelope mismatch'],['Packetized','A_{i,t}~Bernoulli(ρ_t|s_i,a_i)','s_i,a_i,e_i,r_t','request + acceptance','O(N)+O(1)','delay/randomness'],['Transactive','p_i=argmax_p[v_i(p;x_i)-λ_tp], Σ_ip_i=r_t','x_i,λ_t,r_t','price + clearing','O(N log N)','bid/convergence dependence'],['Hybrid / EPS','p_{i,t}=π_i(b_t,x_i); b_t=f(r_t,P_t^agg)','r_t,P_t^agg,π_i','b_t + aggregate feedback','O(1) routine; O(N) fit','calibration/drift/network layer'],['Centralized greedy UB','max Σ_iv_i(p_i) s.t. g(x)≤0','all x_i,q_i,V_k,L_l,T','all device + feeder states','O(N)-O(N log N)+solver','theoretical upper bound; not deployable']]
        table(d,'Supplementary Table S4. Control strategies, required inputs, complexity and limitations.', ['Strategy','Equation / rule','Required inputs','Online signal','Complexity','Limitation'],s4,[1.2,2.25,2.25,1.35,1.15,2.25],5.8)
        body(d,'The control request is r_t and the broadcast command is u_t or b_t. The standard protocol fixes τ=300 s, prediction horizon H=12 intervals, R²_min=0.95, p_ctrl,min=0.90 and N_eff^*=500. The state vector x_{i,t} can contain s_{i,t}, a_{i,t}, p_{i,t}, q_{i,t}, E_i and P_i^max, but each benchmark receives only the subset shown in the table. The term g(x)≤0 denotes device, feeder and voltage feasibility inequalities; it is not assumed available to low-information methods.')
        equation(d,'τ = 300 s   ,   H = 12   ,   R²_min = 0.95   ,   p_ctrl,min = 0.90   ,   N_eff^* = 500','protocol hyperparameters')
        body(d,'Open-loop control establishes a lower-information reference by sending u_t without feedback. The local SOC rule uses s_{i,t}, P_i^max and a_{i,t} at each resource and therefore scales in total device work even though its operator command is scalar. MPC minimizes aggregate tracking error while enforcing g(x)≤0, so it needs forecasts, device states and network variables such as q_i, V_k, L_l and T. Mean-field control compresses the population to moments m_t; its closure can fail when common signals create strong dependence or when the response law drifts. The virtual-battery model replaces the fleet by an admissible envelope, which can hide individual saturation if the envelope is stale. Packetized control randomizes acceptance and can introduce delay variance. Transactive control requires bids v_i and a clearing price λ_t, so convergence depends on bid regularity and communication rounds. Hybrid/EPS uses a calibrated local response π_i and routine broadcast b_t with aggregate feedback P_t^agg; its O(1) claim applies to routine operator-side coordination, not registration, calibration, drift detection or network screening.')
        body(d,'The centralized greedy controller is intentionally defined as a theoretical upper bound. It observes all x_i, q_i and network states and solves the constrained allocation without imposing a communication budget. It is included to quantify the information-deliverability ceiling, not as a field-deployable implementation, and not as a universal optimum for every solver or constraint realization.')
        s5=[['IEEE-33','33 buses / 32 branches','P_l,Q_l,V_k,L_l=S_l/S_l^max,T=T_f/T^max','L_l≤1; V_min≤V_k≤V_max; T≤1','small radial bottlenecks','first clipping constraint'],['IEEE-69','69 / 68','same + z_i,c_f,c_n','c_f=0.5; c_n=0.8; derating; M0-M6','transfer/spatial stress','topology and placement'],['IEEE-123','123 / 122','same + A_k=P_{k,acc}/P_{k,req}','voltage/transformer stress; direct audit','largest radial case','scalability of feasibility']]
        table(d,'Supplementary Table S5. Physical variables, mathematical limits and questions.', ['Environment','Scale','Variables / equations','Limits','Feature','Question'],s5,[1.1,1.15,2.55,2.15,1.75,1.8],5.8)
        body(d,'The physical layer maps each p_{i,t} through z_i to a bus and computes branch active/reactive flows P_l and Q_l, bus voltage magnitude V_k, branch loading L_l and transformer loading T. The apparent branch flow is S_l=(P_l^2+Q_l^2)^{1/2}; feasibility requires L_l≤1. The accepted response P_acc is the response remaining after projection or clipping onto the feasible network set. The resulting acceptance and loss are defined by the following Word equations.')
        equation(d,'S_l = (P_l² + Q_l²)^{1/2}   ,   L_l = S_l / S_l^max   ,   T = T_f / T^max','branch and transformer loading')
        equation(d,'V_min ≤ V_k ≤ V_max   ,   A_net = P_acc / P_req   ,   C_loss = 1 - A_net','voltage feasibility, network acceptance and physical loss')
        body(d,'The reference feeder is a comparator for the aggregate threshold and deliberately carries no topology-specific claim. IEEE-33 is small enough to identify whether a branch, voltage or transformer inequality becomes active first. IEEE-69 introduces a larger radial path and explicit placement stress: c_f=0.5 concentrates half of the resources on one feeder, while c_n=0.8 places 80% at distal nodes. These settings test shared upstream loading and end-of-feeder voltage sensitivity. IEEE-123 is the largest tested representation and is used for a direct-training and safety audit, where the question is whether a reduced-observability strategy retains feasible effect after spatial mapping.')
        body(d,'The constraint-sweep environment varies one physical quantity, such as S_l^max, T^max, V_min or V_max, while holding the remaining protocol values fixed. This separates high aggregate response fidelity R² from network acceptance A_net. A response can therefore be statistically predictable while C_loss is large. Uniform placement, feeder concentration and distal-node concentration are physical realizations of the same population request, not new independent empirical datasets. These results establish constrained-simulation feasibility on the listed IEEE models; they do not certify field safety.')
        table(d,'Data-source index used in Table S1.', ['Dataset','Public source','Citation / use'],urls,[1.6,5.25,3.4],6.6)
    else:
        body(d,'本补充材料详细定义用于解释群体规模实验的数据、控制和网络环境。表格采用紧凑布局，配套正文给出统一符号、数据来源边界、算法假设和物理约束，便于直接审阅和后续改写为 NC Supplementary Information。')
        rows=[[a,b,c,e,f,g,h] for a,b,c,e,f,g,h in datasets]
        table(d,'补充表 S1。经验数据集特征、统一变量与数据特色。',['编号','数据集','资源/信号','分辨率','变量','来源/限制','N_src','数据特色'],[[a,b,c,e,f,g,SOURCE_COUNTS[b],h] for a,b,c,e,f,g,h in rows],[.34,.92,1.0,.55,2.1,1.45,.55,1.35],5.4)
        body(d,'全文统一使用以下符号：i 表示资源或会话，t 表示调度时刻，τ 表示采样间隔；p_{i,t} 和 q_{i,t} 分别表示有功和无功功率；s_{i,t} 表示荷电状态；E_i 表示可用容量；P_i^max 表示最大有功功率；a_{i,t}∈{0,1} 表示可用性；z_i 表示馈线或区域位置；n_i 表示网络或 GIS 信息；g_{i,t} 表示外部发电或充电输入；b_{i,t} 表示基线需求；w_t 表示共同天气变量。聚合响应定义如下。')
        equation(d,'P_t = Σ_i p_{i,t}   ,   Q_t = Σ_i q_{i,t}   ,   N_src = |{source_i}|','聚合有功/无功响应与独立来源数')
        detail_zh={'NextGen':'NextGen 是电池参数、snapshot 和馈线机制的设备日校准群体，提供电池响应、状态、功率限制和馈线位置。其独立来源约为 100 个，而 3,000 个 device-day 是重复来源日记录，不能当作 3,000 个独立现场设备。','data2':'data2 提供 5 分钟充电会话。会话索引作为 i，充电功率作为 p_{i,t}，会话可用性作为 a_{i,t}。会话起止时间形成明确的可用性边界，适合双向充电压力测试；g_{i,t} 在当前适配器中是反事实输入。'}
        for n,_,_,_,_,_,_ in datasets:
            if n=='P00': body(d,detail_zh['NextGen'],boldlead='NextGen. ')
            elif n=='P00b': body(d,detail_zh['data2'],boldlead='data2. ')
            else:
                enname=next(x[1] for x in datasets if x[0]==n); body(d,{'BDG1':'BDG1 提供 507 个建筑的小时用电 p_{i,t}，具有显著负荷形状差异，但不含实测电池 SOC 或光伏发电。s_{i,t}、E_i 和 P_i^max 因此属于仿真变量。','Low Carbon London':'Low Carbon London 提供半小时家庭用电和费率信息，体现大规模住宅负荷的日内结构和家庭差异，但电池状态仍是仿真变量。','BDG2':'BDG2 在建筑用电之外提供太阳能字段。只有时间对齐的子集使用实测 g_{i,t}，其余明确标记为反事实输入。','Danish heat meters':'Danish 数据观测的是热能 h_{i,t}，用于热需求群体规律测试，不能改写成实测电功率；需要 COP 或热泵映射才能进行电气解释。','Smart Grid Smart City':'SGSC 同时提供一般负荷、受控负荷、总发电和净发电字段，是实测自消费分析的主要数据源。','HEAPO':'HEAPO 提供 15 分钟和日尺度的热泵家庭用电以及天气变量 w_t。电池状态和容量不属于实测字段。','GoiEner':'GoiEner 提供大规模小时智能电表负荷和合同功率元数据 c_i。self-generation 标志是类别变量，不是 g_{i,t} 发电轨迹。','European LV urban-8k':'European LV urban-8k 同时提供有功轮廓和馈线/GIS 信息 n_i，用于检验城市客户差异经过空间分配后的表现。','European LV rural':'European LV rural 提供农村拓扑上的 p_{i,t} 和 q_{i,t}，可同时用于客户差异和网络电压/支路压力。','European LV urban-35k':'European LV urban-35k 的价值不仅是最大的 N_src，还在于其大规模城市客户和 P/Q 网络信息可以检验可预测性与网络接受率的分离。','Norway AMI':'Norway AMI 将实测输入/输出、有功/无功、天气和中低压网络模型结合，是网络关联实测证据。','CAMSL-JP':'CAMSL-JP 提供半小时家庭用电、费率组和温度信息，用于费率暴露下的时间响应分析，不假定实测外部发电。','Irish CER':'Irish CER 分开记录输入和输出电量，使 p_{i,t} 的方向明确，适合双向家庭轮廓测试；缺失和停电应作为来源条件保留。','OPSD':'OPSD 具有负荷、光伏、储能、电动车和热泵等丰富字段，但只有 11 个独立家庭，因此是变量丰富而来源不足的压力案例。','COMPLETE-EC':'COMPLETE-EC 包含社区负荷、光伏、电池和电动车及容量/SOC 字段，适合多技术耦合，但 250 个成员不应解释成大规模来源库存。'}[enname],boldlead=enname+'. ')
        body(d,'N_src 与重采样后的行数严格分开。一个来源可以生成多个 device-day 或会话，但这些记录继承来源层面的相关性。因此，跨数据集比较同时考察变量覆盖范围和独立实证复制范围，而不是只比较 N。')
        body(d,'S4 中 r_t 是目标聚合请求，u_t 或 b_t 是广播命令。协议超参数为 τ=300 s、H=12、R²_min=0.95、p_ctrl,min=0.90、N_eff^*=500。x_{i,t} 可以包含 s_{i,t}、a_{i,t}、p_{i,t}、q_{i,t}、E_i 和 P_i^max，但每个算法只接收表中列出的子集。g(x)≤0 表示设备、馈线和电压可行性不等式。集中式贪心方法是忽略通信限制、读取全部设备与网络状态的理论上限，不是真实部署算法，也不声称对所有约束实现普适最优。')
        table(d,'补充表 S4。控制策略、所需输入、复杂度与局限。',['策略','公式/规则','所需输入','在线信号','复杂度','局限'],[['开环','p_{i,t}=u_t','r_t,τ','仅 u_t','O(1)','无法适应 s_{i,t},a_{i,t}'],['本地 SOC','p_{i,t}=clip(u_t;s_{i,t},P_i^max)','s_i,P_i^max,a_i','本地 s_i','O(1)/i','阈值与异质性敏感'],['MPC','min Σ_t(P_t-r_t)^2, g(x)≤0','r_t,s_i,E_i,P_i^max,q_i,V_k,L_l,T','状态+预测','O(HN)+求解器','遥测和模型负担'],['均值场','p_i=π(x_i,m_t), m_t=N^-1Σ_iφ(x_i)','x_i,φ','矩+本地 x_i','O(N)+O(1)','相关/漂移下闭合误差'],['虚拟电池','P_t=clip(r_t,P_min,P_max)','s_i,E_i,P_i^max,a_i','聚合包络','拟合后 O(1)','包络失配'],['分组能量','A_i~Bernoulli(ρ|s_i,a_i)','s_i,a_i,e_i,r_t','请求+接受','O(N)+O(1)','延迟和随机性'],['交易式','p_i=argmax[v_i(p;x_i)-λ_tp], Σ_ip_i=r_t','x_i,λ_t,r_t','价格+清算','O(N log N)','竞价/收敛依赖'],['混合/EPS','p_i=π_i(b_t,x_i), b_t=f(r_t,P_t^agg)','r_t,P_t^agg,π_i','b_t+聚合反馈','日常 O(1)；拟合 O(N)','需标定/漂移/网络层'],['集中式贪心 UB','max Σ_iv_i(p_i), g(x)≤0','全部 x_i,q_i,V_k,L_l,T','全部设备+馈线状态','O(N)-O(N log N)+求解器','理论上限，不可部署']], [1.15,2.2,2.15,1.3,1.1,2.35],5.8)
        body(d,'对于 S5，S_l=(P_l²+Q_l²)^{1/2}，L_l=S_l/S_l^max≤1；V_min≤V_k≤V_max；T=T_f/T^max≤1；A_net=P_acc/P_req；C_loss=1-A_net。IEEE-33 用于识别最先激活的支路、电压或变压器限制；IEEE-69 使用 c_f=0.5 的馈线集中和 c_n=0.8 的末端节点集中检验拓扑与空间布置；IEEE-123 用于最大规模径向网络的直接训练和安全审计。约束扫描一次改变 S_l^max、T^max、V_min 或 V_max 中的一个，其余条件固定，用于区分高 R² 与低 A_net。')
        table(d,'补充表 S5。物理变量、数学限制和研究问题。',['环境','规模','变量/公式','限制','环境特点','问题'],[['IEEE-33','33/32','P_l,Q_l,V_k,L_l,T','L_l≤1;V_min≤V_k≤V_max;T≤1','小型径向瓶颈','哪类约束先截断'],['IEEE-69','69/68','同上+z_i,c_f,c_n','c_f=.5;c_n=.8;降额;M0-M6','迁移/空间压力','拓扑和位置影响'],['IEEE-123','123/122','同上+A_k','电压/变压器压力','最大径向环境','大规模可行性']], [1.05,1.1,2.5,2.1,1.7,1.8],5.8)
        table(d,'S1 使用的数据来源索引。',['数据集','公开来源','引用/用途'],urls,[1.55,5.25,3.45],6.5)
    out=ROOT/('outputs/Supplementary_Tables_S1_S4_S5_English.docx' if lang=='en' else 'outputs/Supplementary_Tables_S1_S4_S5_Chinese.docx'); d.save(out); print(out)

if __name__=='__main__':
    make('en'); make('zh')
