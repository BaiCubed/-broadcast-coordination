from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "Supplementary_Tables_S1_S4_S5_English_tables_only.docx"


def borders(cell, size="8", color="000000"):
    tc_pr = cell._tc.get_or_add_tcPr()
    b = tc_pr.first_child_found_in("w:tcBorders")
    if b is None:
        b = OxmlElement("w:tcBorders")
        tc_pr.append(b)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = b.find(qn("w:" + edge))
        if e is None:
            e = OxmlElement("w:" + edge)
            b.append(e)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), size)
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)


def margins(cell, top=80, start=85, bottom=80, end=85):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    m = tc_pr.first_child_found_in("w:tcMar")
    if m is None:
        m = OxmlElement("w:tcMar")
        tc_pr.append(m)
    for side, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        e = m.find(qn("w:" + side))
        if e is None:
            e = OxmlElement("w:" + side)
            m.append(e)
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    tr_pr.append(e)


def set_cell_text(cell, text, bold=False, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(str(text))
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r.font.size = Pt(size)
    r.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margins(cell)
    borders(cell)


def add_table(doc, title, headers, rows, widths, font_size=9.2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(title)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r.font.size = Pt(13)
    r.bold = True
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    hdr = table.rows[0]
    repeat_header(hdr)
    hdr.height = Inches(0.55)
    hdr.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    for i, h in enumerate(headers):
        set_cell_text(hdr.cells[i], h, bold=True, size=font_size)
        hdr.cells[i].width = Inches(widths[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        table.rows[-1].height = Inches(0.40)
        table.rows[-1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for i, val in enumerate(row):
            set_cell_text(cells[i], val, size=font_size)
            cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.top_margin = Inches(0.55)
    sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.55)
    sec.right_margin = Inches(0.55)

    s1 = [
        ["P01", "BDG1", "Building", "Measured building electricity", "507", "N-scale anchor; smallest building inventory", "S1-01"],
        ["P02", "Low Carbon London", "Meter", "Measured household electricity", "5,000", "Large household-load population", "S1-02"],
        ["P03", "BDG2", "Building", "Measured electricity; partial measured solar", "1,570", "Mixed measured/counterfactual sensitivity", "S1-03"],
        ["P04", "Danish heat meters", "Thermal", "Measured district-heating demand", "2,400", "Thermal-demand proxy; not direct electric load", "S1-04"],
        ["P05", "Smart Grid Smart City", "Meter", "Measured load and generation fields", "600", "Measured-input and self-consumption evidence", "S1-05"],
        ["P06", "HEAPO", "Thermal", "Measured household heat-pump electricity", "1,362", "Heat-pump demand heterogeneity", "S1-06"],
        ["P07", "GoiEner", "Meter", "Measured smart-meter electricity", "5,000", "Large-user population; load diversity", "S1-07"],
        ["P08", "European LV urban-8k", "Feeder", "Measured P profiles with LV context", "5,000", "Urban feeder customer profiles", "S1-08"],
        ["P09", "European LV rural", "Feeder", "Measured P/Q profiles and topology", "2,731", "Rural network customer profiles", "S1-09"],
        ["P10", "European LV urban-35k", "Feeder", "Measured P/Q profiles and topology", "12,000", "Large urban feeder population", "S1-10"],
        ["P11", "Norway AMI", "Feeder", "Measured AMI active/reactive power", "2,994", "Measured-only network-linked profiles", "S1-11"],
        ["P12", "CAMSL-JP", "Meter", "Measured half-hourly smart-meter load", "1,423", "Tariff-response and temporal diversity", "S1-12"],
        ["P13", "Irish CER", "Meter", "Measured import/export smart-meter load", "2,904", "Bidirectional household profiles", "S1-13"],
        ["P14", "OPSD", "Meter", "Measured household channels", "11", "Small-source stress test; explicitly limited", "S1-14"],
        ["P15", "COMPLETE-EC", "Building", "Measured community load, PV, BESS and EV fields", "250", "Integrated community stress test", "S1-15"],
    ]
    add_table(doc, "Supplementary Table S1. Empirical population characteristics and mixed-population construction summary.",
              ["Population ID", "Dataset", "Resource category", "Observed data", "Available source count", "Role in the study", "Reference index"], s1,
              [0.55, 1.25, 1.0, 2.25, 1.0, 2.55, 0.8], 8.0)

    add_table(doc, "Supplementary Table S4. Control strategies and information assumptions.",
              ["Strategy", "Control type", "Information required for action", "Calibration / planning information", "Purpose", "Interpretation boundary"], [
                  ["Open-loop baseline", "Predefined signal", "No feedback during dispatch", "None online", "Lower reference for response variability", "No adaptation to state or network"],
                  ["Local SOC rule", "Device-local heuristic", "Device SOC and local limits", "Device state only", "Device-level control baseline", "No fleet-wide state required"],
                  ["MPC", "Optimization", "Predictions, constraints and state estimates", "Forecasts plus device/network model", "Optimization reference", "Information-intensive and horizon-dependent"],
                  ["Hybrid / EPS broadcast", "Broadcast plus local response", "Broadcast signal and aggregate measurements", "Offline calibration; aggregate online adaptation", "Practical reduced-observability strategy", "Routine coordination is fleet-size independent"],
                  ["Centralized greedy reference", "Full-information dispatch", "All device states and network constraints", "Complete fleet and feeder state", "Upper reference for deliverability", "Reference implementation, not a universal bound"],
              ], [1.45, 1.25, 2.25, 2.0, 2.0, 2.1], 8.1)

    add_table(doc, "Supplementary Table S5. Physical realization and network evaluation settings.",
              ["Network", "Nodes", "Branches", "Evaluation purpose", "Interpretation boundary"], [
                  ["Reference feeder", "Reference case", "Reference representation", "Aggregate threshold robustness", "No topology-specific performance claim"],
                  ["IEEE-33", "33", "32", "Network validation and branch/voltage screening", "Radial feeder with explicit line, voltage and transformer limits"],
                  ["IEEE-69", "69", "68", "Transfer and spatial-placement evaluation", "Larger radial feeder for topology and concentration stress"],
                  ["IEEE-123", "123", "122", "Safety audit and direct-training evaluation", "Highest-resolution tested feeder representation"],
              ], [1.45, 1.0, 1.0, 3.0, 4.6], 8.6)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
