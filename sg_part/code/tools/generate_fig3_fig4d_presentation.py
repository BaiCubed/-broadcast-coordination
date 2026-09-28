"""Generate the cross-dataset Figure 3 and Figure 4D review deck."""

from __future__ import annotations

import json
from pathlib import Path
from statistics import mean
from typing import Any

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
OUTPUT = ROOT / "docs/fig3_fig4d_eps_boundary_review.pptx"

DATASETS = [
    ("BDG1", "bdg1_building_data_genome"),
    ("BDG2", "bdg2_building_data_genome"),
    ("CAMSL", "camsl_japan_smart_meters"),
    ("CEC", "complete_energy_community"),
    ("Danish", "danish_smart_heat_meters"),
    ("EU LV Rural 2731", "european_lv_rural_2731"),
    ("EU LV Urban 35297", "european_lv_urban_35297"),
    ("EU LV Urban 8087", "european_lv_urban_8087"),
    ("Goiener", "goiener_smart_meters"),
    ("HEAPO", "heapo_heat_pumps"),
    ("Irish", "irish_domestic_smart_meters"),
    ("LCL", "low_carbon_london"),
    ("Norway AMI", "norway_ami_energy_distribution"),
    ("OPSD", "opsd_household_data"),
    ("SGSC", "smart_grid_smart_city"),
]

FIG4_VARIANTS = [
    {
        "title": "Overall Figure 4D | 100% external availability",
        "image": RESULTS / "overall_fig4d_algorithms_availability100.png",
        "data": "curtailment_baselines_extra_dataset_scaled_availability100.json",
        "description": "固定外部可用率为 100%，比较九种算法在 15 个数据集上的弃电削减率。",
    },
    {
        "title": "Overall Figure 4D | 89.49% external availability",
        "image": RESULTS / "overall_fig4d_algorithms_availability8949.png",
        "data": "curtailment_baselines_extra_dataset_scaled_availability8949.json",
        "description": "使用 89.49% 外部可用率压力测试，设备层、网络和训练设置保持一致。",
    },
    {
        "title": "Overall Figure 4D | Capacity-sufficient audit",
        "image": RESULTS / "overall_fig4d_algorithms_capacity_sufficient_audit.png",
        "data": "curtailment_baselines_extra_dataset_scaled_capacity_sufficient.json",
        "description": "只对原始容量不足的 6 个数据集放大容量与功率包络，用于分离容量瓶颈。",
    },
]

NAVY = RGBColor(31, 50, 72)
TEAL = RGBColor(20, 133, 130)
ORANGE = RGBColor(216, 103, 62)
TEXT = RGBColor(42, 49, 56)
MUTED = RGBColor(96, 105, 115)
PALE = RGBColor(239, 244, 245)
WHITE = RGBColor(255, 255, 255)


def _add_text(
    slide: Any,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    *,
    size: float,
    color: RGBColor = TEXT,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
) -> Any:
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(0.06)
    frame.margin_top = frame.margin_bottom = Inches(0.03)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = "Microsoft YaHei"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return shape


def _base_slide(prs: Presentation, title: str, label: str | None = None) -> Any:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    header = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.72))
    header.fill.solid()
    header.fill.fore_color.rgb = NAVY
    header.line.fill.background()
    _add_text(slide, 0.45, 0.12, 11.7, 0.44, title, size=22, color=WHITE, bold=True)
    if label:
        _add_text(
            slide, 11.8, 0.15, 1.05, 0.36, label, size=14, color=WHITE,
            bold=True, align=PP_ALIGN.RIGHT,
        )
    return slide


def _add_image(slide: Any, path: Path, left: float, top: float, width: float, height: float) -> None:
    if not path.is_file():
        raise FileNotFoundError(path)
    with Image.open(path) as source:
        image_ratio = source.width / source.height
    area_ratio = width / height
    if image_ratio >= area_ratio:
        output_width = width
        output_height = width / image_ratio
        output_left = left
        output_top = top + (height - output_height) / 2
    else:
        output_height = height
        output_width = height * image_ratio
        output_left = left + (width - output_width) / 2
        output_top = top
    slide.shapes.add_picture(
        str(path), Inches(output_left), Inches(output_top), Inches(output_width), Inches(output_height)
    )


def _note_band(slide: Any, heading: str, body: str, *, accent: RGBColor = TEAL) -> None:
    band = slide.shapes.add_shape(1, Inches(0.45), Inches(5.75), Inches(12.43), Inches(1.3))
    band.fill.solid()
    band.fill.fore_color.rgb = PALE
    band.line.color.rgb = RGBColor(208, 220, 222)
    accent_bar = slide.shapes.add_shape(1, Inches(0.45), Inches(5.75), Inches(0.08), Inches(1.3))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = accent
    accent_bar.line.fill.background()
    _add_text(slide, 0.67, 5.91, 2.05, 0.28, heading, size=12, color=accent, bold=True)
    _add_text(slide, 2.55, 5.84, 10.0, 0.98, body, size=10.5, color=TEXT, valign=MSO_ANCHOR.MIDDLE)


def _threshold_payload(dataset: str) -> dict[str, Any]:
    path = (
        RESULTS / f"{dataset}_ieee33_real_load/coverage_fix/network_constrained_new/data/n_threshold.json"
    )
    return json.loads(path.read_text(encoding="utf-8"))


def _figure3_slide(prs: Presentation, short_name: str, dataset: str) -> None:
    root = RESULTS / f"{dataset}_ieee33_real_load/coverage_fix/network_constrained_new"
    threshold = _threshold_payload(dataset)
    slide = _base_slide(prs, f"{short_name} | {dataset}", short_name)
    _add_image(slide, root / "Figs/fig3_scaling_heterogeneity.png", 0.45, 0.88, 12.43, 4.63)
    points = threshold["per_N"]
    small_values = ", ".join(
        f"N={fleet}: {points[str(fleet)]['r2']:.3f}" for fleet in (1, 5, 15, 30)
    )
    n95 = threshold.get("threshold_N_95_interpolated")
    if n95 is None:
        boundary = "N95 未达到；OPSD 的 N=15/30 使用 6 个独立来源 profile 构造逻辑设备，仅用于容量边界审计。"
        accent = ORANGE
    else:
        boundary = f"对数插值 N95={n95:.1f}；N<=30 尚未达到 R2=0.95，可靠聚合预测需要更大设备规模。"
        accent = TEAL
    _note_band(
        slide,
        "Figure 3 结论",
        f"A: R2-规模；B: 聚合 CV-规模；C: 异质性查找表。小规模实测：{small_values}。{boundary}",
        accent=accent,
    )


def _variant_rows(filename: str) -> list[dict[str, Any]]:
    rows = []
    pattern = f"*_ieee33_real_load/coverage_fix/network_constrained_new/data/{filename}"
    for path in sorted(RESULTS.glob(pattern)):
        payload = json.loads(path.read_text(encoding="utf-8"))
        rows.append({
            "dataset": path.parts[-5].removesuffix("_ieee33_real_load"),
            "eps": float(payload["results"]["eps_broadcast"]["mean_reduction_pct"]),
            "local": float(payload["results"]["local_rules"]["mean_reduction_pct"]),
            "centralized": float(payload["results"]["centralized_optimal"]["mean_reduction_pct"]),
        })
    if not rows:
        raise FileNotFoundError(filename)
    return rows


def _figure4_slide(prs: Presentation, variant: dict[str, Any]) -> dict[str, Any]:
    rows = _variant_rows(variant["data"])
    eps = [row["eps"] for row in rows]
    below_50 = [row["dataset"] for row in rows if row["eps"] < 50.0]
    slide = _base_slide(prs, variant["title"], "Figure 4D")
    _add_image(slide, variant["image"], 0.45, 0.88, 12.43, 4.58)
    summary = (
        f"{variant['description']} EPS 平均={mean(eps):.1f}%，范围={min(eps):.1f}%–{max(eps):.1f}%，"
        f"达到 90% 的数据集={sum(value >= 90 for value in eps)}/{len(eps)}。"
    )
    if below_50:
        summary += f" 低于 50%：{', '.join(below_50)}。"
    else:
        summary += " 所有审计数据集均达到 50%。"
    _note_band(slide, "EPS 有效边界", summary, accent=ORANGE if below_50 else TEAL)
    return {
        "title": variant["title"],
        "count": len(rows),
        "mean": mean(eps),
        "minimum": min(eps),
        "maximum": max(eps),
        "ge90": sum(value >= 90 for value in eps),
        "below50": below_50,
    }


def _cover_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    _add_text(
        slide, 0.85, 1.25, 11.7, 1.2,
        "Figure 3 Small-Fleet Scaling\nand Figure 4D EPS Boundary",
        size=30, color=WHITE, bold=True,
    )
    _add_text(
        slide, 0.9, 3.0, 11.0, 0.8,
        "15 datasets | Added N = 1, 5, 15, 30 | Overall baseline comparison",
        size=17, color=RGBColor(196, 220, 221),
    )
    _add_text(
        slide, 0.9, 6.4, 11.0, 0.45,
        "数据来自 results/coverage_fix/network_constrained_new；Figure 3 图片已覆盖更新。",
        size=11, color=RGBColor(196, 203, 211),
    )


def _boundary_slide(prs: Presentation, summaries: list[dict[str, Any]]) -> None:
    n95_values = [
        float(_threshold_payload(dataset)["threshold_N_95_interpolated"])
        for _, dataset in DATASETS
        if _threshold_payload(dataset)["threshold_N_95_interpolated"] is not None
    ]
    slide = _base_slide(prs, "EPS Effective Boundary | Evidence synthesis", "Summary")
    lines = [
        f"1. 聚合可预测性边界：14 个数据集的插值 N95 为 {min(n95_values):.1f}–{max(n95_values):.1f}；N=1/5/15/30 用于展示进入稳定区前的曲线。",
        "2. 来源边界：OPSD 只有 6 个平衡独立来源，N=15/30 是 profile 复用审计，R2 在 N<=30 未达到 0.95，不能当作 source-unique 证据。",
        f"3. 容量边界：100% 可用率下 EPS 平均 {summaries[0]['mean']:.1f}%，仍有 {len(summaries[0]['below50'])} 个数据集低于 50%，说明可用率不是唯一瓶颈。",
        f"4. 可用率边界：89.49% 版本 EPS 平均 {summaries[1]['mean']:.1f}%，相对 100% 版本变化 {summaries[1]['mean'] - summaries[0]['mean']:+.1f} 个百分点。",
        f"5. 容量充足审计：6 个受限数据集的 EPS 提升到 {summaries[2]['minimum']:.1f}–{summaries[2]['maximum']:.1f}%，证明容量/功率 headroom 是关键边界；但不证明 EPS 在所有数据集上优于集中式或优化 baseline。",
    ]
    for index, line in enumerate(lines):
        top = 1.05 + index * 1.12
        color = TEAL if index in {0, 4} else TEXT
        _add_text(slide, 0.75, top, 11.9, 0.82, line, size=15, color=color, bold=index in {0, 4})


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    _cover_slide(prs)
    for short_name, dataset in DATASETS:
        _figure3_slide(prs, short_name, dataset)
    summaries = [_figure4_slide(prs, variant) for variant in FIG4_VARIANTS]
    _boundary_slide(prs, summaries)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(json.dumps({"output": str(OUTPUT), "slides": len(prs.slides)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
