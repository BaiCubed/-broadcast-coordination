"""将 E22 子图组合为期刊式 PPTX 和 PNG。"""

from __future__ import annotations

import argparse
import csv
import tempfile
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

try:
    import nckeys as K
except ModuleNotFoundError:
    from tools import nckeys as K


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "outputs" / "figs" / "subpanel"
DIRECT_TRAIN_DIR = ROOT / "outputs" / "figs" / "direct_train"
DEFAULT_PPTX = SOURCE_DIR / "e22_grid_constrained_panels.pptx"
DEFAULT_PNG = SOURCE_DIR / "e22_grid_constrained_panels.png"
DEFAULT_07_PPTX = SOURCE_DIR / "07_physical_to_computational_boundary.pptx"
DEFAULT_07_PNG = SOURCE_DIR / "07_physical_to_computational_boundary_page.png"
PAGE_WIDTH = 12.0
PAGE_HEIGHT = 15.0
MARGIN = 0.25
GAP = 0.025
HEADER_HEIGHT = 0.50
FULL_WIDTH = PAGE_WIDTH - 2 * MARGIN
HALF_WIDTH = (FULL_WIDTH - GAP) / 2


@dataclass(frozen=True)
class Panel:
    label: str
    filename: str
    description: str


PANELS = (
    Panel("a", "00_ieee69_network_acceptance.png", ""),
    Panel("c", "01_ieee69_constraint_audit.png", ""),
    Panel("f", "02_algorithm_effect_raw_vs_pairwise_mixed.png", ""),
    Panel("b", "03_ieee69_spatial_retention.png", ""),
    Panel("e", "04_topology_effect_raw_vs_pairwise_mixed.png", "IEEE-33  |  IEEE-69  |  IEEE-123"),
    Panel("d", "05_ieee69_dataset_algorithm_retention.png", ""),
    Panel("g", "06_ieee69_all_algorithms_m0_m6.png", ""),
)
PANEL_07 = Panel("h", "07_physical_to_computational_boundary.png", "")
ALGORITHM_KEY = (
    "A  No coordination  O(0)    B  Local SOC rules  O(1) per device    C  MPC  O(HN)\n"
    "D  Mean-field control  O(N)    E  Virtual battery  O(N)    F  Packetized Energy Management  O(N log N)\n"
    "G  Transactive control  O(N log N)    H  EPS (ours)  O(1)    I  Centralized greedy upper bound  O(N), reference only"
)
PANEL_GUIDE_ROWS = (
    ("a", "Seven stress modes", "Algorithm code A-I", "Cell = network acceptance (%)"),
    ("b", "Retention (%) by topology and deployment", "Algorithm code A-I", "Dot = median; whisker = 10-90% across source datasets"),
    ("c", "Constraint value", "Curtailment reduction (left); R² / acceptance (right)", "Line and transformer use remaining capacity; voltage uses p.u."),
    ("d", "Response tracking R²", "Algorithm code A-H", "Colored points = original datasets; gray = pairwise mixed"),
    ("e", "Curtailment reduction (%)", "Algorithm code A-H", "Colored points = original datasets; gray = pairwise mixed"),
    ("f", "Spatial heterogeneity scenarios", "Curtailment reduction (%)", "Dashed line = centralized greedy UB"),
    ("g", "Deployment / capacity scenarios", "Curtailment reduction (%)", "Dashed line = centralized greedy UB"),
)


def _f_comparison_text() -> str:
    source = SOURCE_DIR / "source_data" / "05_response_fidelity_r2.csv"
    values: dict[str, list[float]] = defaultdict(list)
    with source.open(encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            try:
                values[row["algorithm"]].append(float(row["response_r2"]))
            except (KeyError, TypeError, ValueError):
                continue
    eps_values = values.get("eps_ieee69_fused", [])
    coordinated = [
        "mpc_optimal",
        "mean_field_control",
        "virtual_battery",
        "packetized_energy_management",
        "transactive_control",
    ]
    coordinated_means = [sum(values[name]) / len(values[name]) for name in coordinated if values.get(name)]
    eps_mean = sum(eps_values) / len(eps_values)
    return (
        f"EPS (ours): R² {eps_mean:.2f}, O(1)\n"
        f"Coordinated baselines: R² {min(coordinated_means):.2f}-{max(coordinated_means):.2f}\n"
        "R² measures target-response fidelity\n"
        "EPS keeps constant communication and fast response"
    )


def _set_textbox_text(box, lines: list[str], *, size: int, bold_first: bool = False) -> None:
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold_first and index == 0
        paragraph.font.color.rgb = RGBColor(26, 26, 26)
        paragraph.alignment = PP_ALIGN.LEFT


def _add_panel_guide_slide(presentation: Presentation) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(MARGIN), Inches(0.18), Inches(FULL_WIDTH), Inches(0.35))
    _set_textbox_text(title, ["Panel guide"], size=18, bold_first=True)

    subtitle = slide.shapes.add_textbox(Inches(MARGIN), Inches(0.50), Inches(FULL_WIDTH), Inches(0.22))
    _set_textbox_text(
        subtitle,
        ["Panels are ordered left-to-right, top-to-bottom as a, b, c, d, e, f, g."],
        size=10,
    )

    table_left = Inches(MARGIN)
    table_top = Inches(0.84)
    table_width = Inches(FULL_WIDTH)
    table_height = Inches(6.15)
    rows = len(PANEL_GUIDE_ROWS) + 1
    cols = 4
    table = slide.shapes.add_table(rows, cols, table_left, table_top, table_width, table_height).table
    widths = [0.55, 2.70, 3.05, 4.70]
    for idx, width in enumerate(widths):
        table.columns[idx].width = Inches(width)

    headers = ["Panel", "X-axis", "Y-axis / value", "Meaning"]
    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = header
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.name = "Arial"
            paragraph.font.size = Pt(11)
            paragraph.font.bold = True
            paragraph.font.color.rgb = RGBColor(26, 26, 26)
            paragraph.alignment = PP_ALIGN.CENTER

    for row_index, row in enumerate(PANEL_GUIDE_ROWS, start=1):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = value
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "Arial"
                paragraph.font.size = Pt(9.5)
                paragraph.font.color.rgb = RGBColor(26, 26, 26)
                paragraph.alignment = PP_ALIGN.CENTER if col_index == 0 else PP_ALIGN.LEFT

    for row in range(rows):
        for col in range(cols):
            cell = table.cell(row, col)
            cell.margin_left = Inches(0.03)
            cell.margin_right = Inches(0.03)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            if row == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(240, 243, 248)

    row_height = 6.15 / rows
    for idx in range(rows):
        table.rows[idx].height = Inches(row_height)

    lower_left = slide.shapes.add_textbox(Inches(MARGIN), Inches(7.28), Inches(5.55), Inches(3.15))
    _set_textbox_text(
        lower_left,
        [
            "Algorithm key",
            "A  No coordination  O(0)",
            "B  Local SOC rules  O(1) per device",
            "C  MPC  O(HN)",
            "D  Mean-field control  O(N)",
            "E  Virtual battery  O(N)",
            "F  Packetized Energy Management  O(N log N)",
            "G  Transactive control  O(N log N)",
            "H  EPS (ours)  O(1)",
            "I  Centralized greedy upper bound  O(N), reference only",
        ],
        size=10,
        bold_first=True,
    )

    lower_right = slide.shapes.add_textbox(Inches(6.05), Inches(7.28), Inches(5.70), Inches(3.15))
    _set_textbox_text(
        lower_right,
        [
            "Legend",
            "Colored points = original datasets",
            "Gray points = pairwise mixed",
            "Dashed line = centralized greedy UB",
            "0.0% = no-coordination baseline anchor",
            "Panel d uses R²; panel e uses curtailment reduction (%)",
        ],
        size=10,
        bold_first=True,
    )

    footer = slide.shapes.add_textbox(Inches(MARGIN), Inches(10.55), Inches(FULL_WIDTH), Inches(0.38))
    _set_textbox_text(
        footer,
        ["This slide replaces in-chart x-axis labels for panels d and e."],
        size=9,
    )


def _font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = (
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    )
    for candidate in candidates:
        if Path(candidate).is_file():
            return ImageFont.truetype(candidate, size=size)
    return ImageFont.load_default()


def _fit(image: Image.Image, width: float, height: float) -> tuple[float, float]:
    scale = min(width / image.width, height / image.height)
    return image.width * scale, image.height * scale


def _source_path(source_dir: Path, panel: Panel) -> Path:
    candidate = source_dir / panel.filename
    if candidate.is_file():
        return candidate
    return SOURCE_DIR / panel.filename


def _panels_for_source(source_dir: Path) -> tuple[Panel, ...]:
    return PANELS


def _main_boxes(source_dir: Path, panels: tuple[Panel, ...]) -> list[tuple[float, float, float, float]]:
    row_heights = (3.05, 3.60, 2.85, 2.85)
    tops = []
    current = 0.14
    for height in row_heights:
        image_top = current + HEADER_HEIGHT
        tops.append(image_top)
        current = image_top + height + GAP

    specs = [
        (PANELS[0], MARGIN, tops[0], HALF_WIDTH, row_heights[0]),
        (PANELS[3], MARGIN + HALF_WIDTH + GAP, tops[0], HALF_WIDTH, row_heights[0]),
        (PANELS[1], MARGIN, tops[1], FULL_WIDTH, row_heights[1]),
        (PANELS[5], MARGIN, tops[2], HALF_WIDTH, row_heights[2]),
        (PANELS[4], MARGIN + HALF_WIDTH + GAP, tops[2], HALF_WIDTH, row_heights[2]),
        (PANELS[2], MARGIN, tops[3], HALF_WIDTH, row_heights[3]),
        (PANELS[6], MARGIN + HALF_WIDTH + GAP, tops[3], HALF_WIDTH, row_heights[3]),
    ]
    boxes_by_filename: dict[str, tuple[float, float, float, float]] = {}
    for panel, left, top, slot_width, slot_height in specs:
        boxes_by_filename[panel.filename] = (left, top, slot_width, slot_height)
    return [boxes_by_filename[panel.filename] for panel in PANELS]


def _standalone_07_box(source_dir: Path) -> tuple[float, float, float, float]:
    with Image.open(_source_path(source_dir, PANEL_07)) as image:
        width, height = _fit(image, FULL_WIDTH, 6.5)
    return MARGIN + (FULL_WIDTH - width) / 2, 0.14 + HEADER_HEIGHT, width, height


def _paste_panel(page: Image.Image, draw: ImageDraw.ImageDraw, panel: Panel, box: tuple[float, float, float, float], dpi: int, source_dir: Path) -> None:
    left, top, width, height = box
    with Image.open(_source_path(source_dir, panel)) as source:
        source = source.convert("RGB")
        target_size = (round(width * dpi), round(height * dpi))
        fitted_size = tuple(round(value) for value in _fit(source, *target_size))
        fitted = source.resize(fitted_size, Image.Resampling.LANCZOS)
        image = Image.new("RGB", target_size, "white")
        image.paste(fitted, ((target_size[0] - fitted.width) // 2, (target_size[1] - fitted.height) // 2))
    x, y = round(left * dpi), round(top * dpi)
    page.paste(image, (x, y))


def _draw_panel_header(draw: ImageDraw.ImageDraw, panel: Panel, box: tuple[float, float, float, float], dpi: int) -> None:
    left, top, width, _ = box
    x = round(left * dpi)
    y = round((top - HEADER_HEIGHT + 0.02) * dpi)
    label_font = _font(round(0.21 * dpi), bold=True)
    text_font = _font(round(0.155 * dpi))
    draw.text((x, y), panel.label, font=label_font, fill="#1A1A1A")
    if panel.description and panel.label != "e":
        draw.multiline_text(
            (x + round(0.27 * dpi), y + round(0.015 * dpi)),
            panel.description,
            font=text_font,
            fill="#1A1A1A",
            spacing=round(0.035 * dpi),
        )



def _extra_header_items(panel, box):
    """Header text stays outside the fixed image rectangle (inches, points)."""
    left, top, width, _ = box
    if panel.label == "b":
        axis_width = 0.89 / (6 + 5 * 0.12)
        axis_gap = axis_width * 0.12
        for index, topology in enumerate(("IEEE-33", "IEEE-69", "IEEE-123")):
            center_fraction = 0.075 + axis_width + axis_gap / 2 + index * 2 * (axis_width + axis_gap)
            center = left + width * center_fraction
            yield center - 0.65, top - HEADER_HEIGHT + 0.045, 1.30, 0.22, topology, 12, None, True
    if panel.label == "e":
        for index, topology in enumerate(("IEEE-33", "IEEE-69", "IEEE-123")):
            center = left + width * (0.23 + index * 0.30)
            yield center - 0.65, top - HEADER_HEIGHT + 0.035, 1.30, 0.24, topology, 12, None, True
    if panel.label == "b":
        labels = ("50%", "80%") * 3
        # Centers of the six axes after make_spatial_retention's margins and
        # wspace are not evenly spaced across the full image rectangle.
        for index, label in enumerate(labels):
            center_fraction = 0.075 + axis_width / 2.0 + index * (axis_width + axis_gap)
            center = left + width * center_fraction
            yield center - 0.46, top - HEADER_HEIGHT + 0.29, 0.92, 0.17, label, 11, None, True
    if panel.label == "g":
        items = [
            ("Uniform", "#365A7C"),
            ("50% feeder", "#009E73"),
            ("80% distal node", "#E69F00"),
            ("50% feeder; 70% cap.", "#D55E00"),
        ]
        col_width = (width - 0.34) / len(items)
        for index, (label, colour) in enumerate(items):
            yield left + 0.34 + index * col_width, top - 0.23, col_width, 0.22, label, 9.5, colour, False


def _draw_extra_headers(draw, panel, box, dpi):
    if panel.label == "b":
        left, top, width, _ = box
        axis_width = 0.89 / (6 + 5 * 0.12)
        axis_gap = axis_width * 0.12
        band_top = round((top - HEADER_HEIGHT + 0.01) * dpi)
        band_bottom = round((top - 0.005) * dpi)
        colours = ("#DCEAF7", "#DDF2E8", "#FBE7D3")
        for group_index, colour in enumerate(colours):
            x0 = left + width * (0.075 + group_index * 2 * (axis_width + axis_gap))
            x1 = left + width * (0.075 + (group_index * 2 + 2) * axis_width + (group_index * 2 + 1) * axis_gap)
            draw.rectangle((round(x0 * dpi), band_top, round(x1 * dpi), band_bottom), fill=colour)
    for left, top, width, height, label, size, colour, centered in _extra_header_items(panel, box):
        x, y = round(left * dpi), round(top * dpi)
        font = _font(round(size / 72 * dpi))
        if colour:
            radius = round(0.023 * dpi)
            cx, cy = x + radius, y + round(0.065 * dpi)
            draw.ellipse((cx-radius, cy-radius, cx+radius, cy+radius), fill=colour)
            x += round(0.10 * dpi)
        if centered and "\n" in label:
            line_height = max(1, round(size / 72 * dpi * 1.02))
            for line_index, line in enumerate(label.splitlines()):
                line_width = draw.textlength(line, font=font)
                line_x = round(left * dpi + (width * dpi - line_width) / 2)
                draw.text((line_x, y + line_index * line_height), line, font=font, fill="#1A1A1A")
        else:
            if centered:
                bounds = draw.textbbox((0, 0), label, font=font)
                x += round((width * dpi - (bounds[2] - bounds[0])) / 2)
            draw.text((x, y), label, font=font, fill="#1A1A1A")


def _add_extra_headers(slide, panel, box):
    if panel.label == "b":
        left, top, width, _ = box
        axis_width = 0.89 / (6 + 5 * 0.12)
        axis_gap = axis_width * 0.12
        colours = ("#DCEAF7", "#DDF2E8", "#FBE7D3")
        for group_index, colour in enumerate(colours):
            x0 = left + width * (0.075 + group_index * 2 * (axis_width + axis_gap))
            x1 = left + width * (0.075 + (group_index * 2 + 2) * axis_width + (group_index * 2 + 1) * axis_gap)
            band = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(x0), Inches(top - HEADER_HEIGHT + 0.01),
                Inches(x1 - x0), Inches(HEADER_HEIGHT - 0.015),
            )
            band.fill.solid()
            band.fill.fore_color.rgb = _hex_rgb(colour)
            band.line.fill.background()
            band._element.getparent().remove(band._element)
            slide.shapes._spTree.insert(2, band._element)
    for left, top, width, height, label, size, colour, centered in _extra_header_items(panel, box):
        if colour:
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top+0.045), Inches(0.046), Inches(0.046))
            dot.fill.solid()
            dot.fill.fore_color.rgb = _hex_rgb(colour)
            dot.line.fill.background()
            left += 0.10
            width -= 0.10
        text = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = text.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.word_wrap = False
        for index, line in enumerate(label.splitlines()):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = "Arial"
            paragraph.font.size = Pt(size)
            paragraph.space_before = paragraph.space_after = Pt(0)
            paragraph.alignment = PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT


def _draw_f_comparison(draw: ImageDraw.ImageDraw, box: tuple[float, float, float, float], dpi: int) -> None:
    left, top, width, height = box
    text_value = _f_comparison_text()
    font = _font(round(0.115 * dpi))
    x = round((left + width * 0.47) * dpi)
    y = round((top + height * 0.18) * dpi)
    spacing = round(0.025 * dpi)
    bounds = draw.multiline_textbbox((x, y), text_value, font=font, spacing=spacing)
    padding = round(0.055 * dpi)
    draw.rounded_rectangle(
        (bounds[0] - padding, bounds[1] - padding, bounds[2] + padding, bounds[3] + padding),
        radius=round(0.035 * dpi),
        fill="#FFFFFF",
        outline="#B9C2CB",
        width=max(1, round(0.008 * dpi)),
    )
    draw.multiline_text((x, y), text_value, font=font, fill="#1A1A1A", spacing=spacing)


def _draw_algorithm_key(draw: ImageDraw.ImageDraw, left: int, top: int, width: int, dpi: int) -> None:
    font = _font(round(0.19 * dpi))
    draw.text((left, top), "Algorithm key", font=_font(round(0.21 * dpi), bold=True), fill="#1A1A1A")
    line_height = round(0.24 * dpi)
    for index, line in enumerate(ALGORITHM_KEY.splitlines(), start=1):
        draw.text((left, top + index * line_height), line, font=font, fill="#1A1A1A")


def _draw_dataset_key(draw: ImageDraw.ImageDraw, left: int, top: int, width: int, dpi: int) -> None:
    datasets = list(K.DATASET_LABELS)
    columns = 6
    column_width = width / columns
    title_font = _font(round(0.17 * dpi), bold=True)
    item_font = _font(round(0.14 * dpi))
    draw.text((left, top), "Dataset points", font=title_font, fill="#1A1A1A")
    item_top = top + round(0.24 * dpi)
    row_height = round(0.22 * dpi)
    radius = round(0.035 * dpi)
    for index, dataset in enumerate(datasets):
        row, column = divmod(index, columns)
        x = round(left + column * column_width)
        y = item_top + row * row_height
        center_x = x + round(0.045 * dpi)
        center_y = y + round(0.065 * dpi)
        colour = K.DATASET_COLOURS[dataset]
        draw.ellipse((center_x - radius, center_y - radius, center_x + radius, center_y + radius), fill=colour)
        draw.text((x + round(0.105 * dpi), y), K.DATASET_LABELS[dataset], font=item_font, fill="#1A1A1A")


def _write_png(
    output: Path,
    panels: tuple[Panel, ...],
    boxes: list[tuple[float, float, float, float]],
    height: float,
    algorithm_key_top: float,
    dataset_key_top: float | None = None,
    source_dir: Path = SOURCE_DIR,
) -> None:
    dpi = 600
    page = Image.new("RGB", (round(PAGE_WIDTH * dpi), round(height * dpi)), "white")
    draw = ImageDraw.Draw(page)
    for panel, box in zip(panels, boxes):
        _draw_panel_header(draw, panel, box, dpi)
        _paste_panel(page, draw, panel, box, dpi, source_dir)
        _draw_extra_headers(draw, panel, box, dpi)
    if dataset_key_top is not None:
        _draw_dataset_key(draw, round(MARGIN * dpi), round(dataset_key_top * dpi), round(FULL_WIDTH * dpi), dpi)
    _draw_algorithm_key(draw, round(MARGIN * dpi), round(algorithm_key_top * dpi), round(FULL_WIDTH * dpi), dpi)
    output.parent.mkdir(parents=True, exist_ok=True)
    page.save(output, dpi=(dpi, dpi), compress_level=6)
    page.save(output.with_suffix(".pdf"), "PDF", resolution=dpi)


def _hex_rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def _add_dataset_key_pptx(slide, top: float) -> None:
    title = slide.shapes.add_textbox(Inches(MARGIN), Inches(top), Inches(FULL_WIDTH), Inches(0.24))
    title_frame = title.text_frame
    title_frame.clear()
    title_paragraph = title_frame.paragraphs[0]
    title_paragraph.text = "Dataset points"
    title_paragraph.font.name = "Arial"
    title_paragraph.font.size = Pt(12)
    title_paragraph.font.bold = True
    title_paragraph.font.color.rgb = RGBColor(26, 26, 26)
    datasets = list(K.DATASET_LABELS)
    columns = 6
    column_width = FULL_WIDTH / columns
    for index, dataset in enumerate(datasets):
        row, column = divmod(index, columns)
        left = MARGIN + column * column_width
        item_top = top + 0.25 + row * 0.22
        marker = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(item_top + 0.03), Inches(0.07), Inches(0.07))
        marker.fill.solid()
        marker.fill.fore_color.rgb = _hex_rgb(K.DATASET_COLOURS[dataset])
        marker.line.fill.background()
        label = slide.shapes.add_textbox(Inches(left + 0.10), Inches(item_top), Inches(column_width - 0.12), Inches(0.18))
        frame = label.text_frame
        frame.clear()
        paragraph = frame.paragraphs[0]
        paragraph.text = K.DATASET_LABELS[dataset]
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(10)
        paragraph.font.color.rgb = RGBColor(26, 26, 26)


def _write_pptx(
    output: Path,
    panels: tuple[Panel, ...],
    boxes: list[tuple[float, float, float, float]],
    height: float,
    algorithm_key_top: float,
    dataset_key_top: float | None = None,
    source_dir: Path = SOURCE_DIR,
) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(PAGE_WIDTH)
    presentation.slide_height = Inches(height)
    presentation.core_properties.title = "E22 grid-constrained performance panels"
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    with tempfile.TemporaryDirectory(prefix="e22-panel-ppt-") as temporary_directory:
        temporary = Path(temporary_directory)
        for panel, box in zip(panels, boxes):
            left, top, width, panel_height = box
            with Image.open(_source_path(source_dir, panel)) as source:
                source = source.convert("RGB")
                target_size = (1800, round(1800 * panel_height / width))
                fitted_size = tuple(round(value) for value in _fit(source, *target_size))
                fitted = source.resize(fitted_size, Image.Resampling.LANCZOS)
                asset = Image.new("RGB", target_size, "white")
                asset.paste(fitted, ((target_size[0] - fitted.width) // 2, (target_size[1] - fitted.height) // 2))
            asset_path = temporary / panel.filename
            asset.save(asset_path, optimize=True)
            slide.shapes.add_picture(str(asset_path), Inches(left), Inches(top), width=Inches(width), height=Inches(panel_height))
            header = slide.shapes.add_textbox(Inches(left), Inches(top - HEADER_HEIGHT + 0.01), Inches(width), Inches(HEADER_HEIGHT - 0.01))
            header.fill.background()
            header.line.fill.background()
            frame = header.text_frame
            frame.clear()
            frame.word_wrap = True
            paragraph = frame.paragraphs[0]
            label_run = paragraph.add_run()
            label_run.text = f"{panel.label}  "
            label_run.font.name = "Arial"
            label_run.font.size = Pt(15)
            label_run.font.bold = True
            label_run.font.color.rgb = RGBColor(26, 26, 26)
            _add_extra_headers(slide, panel, box)
            if panel.description and panel.label != "e":
                description_run = paragraph.add_run()
                description_run.text = panel.description.replace("\n", " ")
                description_run.font.name = "Arial"
                description_run.font.size = Pt(11)
                description_run.font.color.rgb = RGBColor(26, 26, 26)
    if dataset_key_top is not None:
        _add_dataset_key_pptx(slide, dataset_key_top)
    key = slide.shapes.add_textbox(Inches(MARGIN), Inches(algorithm_key_top), Inches(FULL_WIDTH), Inches(1.05))
    frame = key.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraphs = frame.paragraphs
    for index, line in enumerate(("Algorithm key", *ALGORITHM_KEY.splitlines())):
        paragraph = paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(14 if index else 16)
        paragraph.font.bold = index == 0
        paragraph.font.color.rgb = RGBColor(26, 26, 26)
    _add_panel_guide_slide(presentation)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pptx", type=Path, default=DEFAULT_PPTX, help="主组合 PPTX 输出路径")
    parser.add_argument("--png", type=Path, default=DEFAULT_PNG, help="主组合 PNG 输出路径")
    parser.add_argument("--standalone-07-pptx", type=Path, default=DEFAULT_07_PPTX, help="07 独立 PPTX 输出路径")
    parser.add_argument("--standalone-07-png", type=Path, default=DEFAULT_07_PNG, help="07 独立 PNG 输出路径")
    parser.add_argument("--source-dir", type=Path, default=SOURCE_DIR, help="子图输入目录")
    args = parser.parse_args()
    source_dir = args.source_dir.resolve()
    panels = _panels_for_source(source_dir)
    missing = [str(_source_path(source_dir, panel)) for panel in (*panels, PANEL_07) if not _source_path(source_dir, panel).is_file()]
    if missing:
        raise FileNotFoundError("缺少输入子图：\n" + "\n".join(missing))

    main_boxes = _main_boxes(source_dir, panels)
    main_panel_bottom = max(top + panel_height for _, top, _, panel_height in main_boxes)
    main_dataset_key_top = main_panel_bottom + 0.18
    main_algorithm_key_top = main_dataset_key_top + 0.92
    main_height = main_algorithm_key_top + 1.02
    _write_png(args.png, panels, main_boxes, main_height, main_algorithm_key_top, main_dataset_key_top, source_dir)
    _write_pptx(args.pptx, panels, main_boxes, main_height, main_algorithm_key_top, main_dataset_key_top, source_dir)

    standalone_box = _standalone_07_box(source_dir)
    standalone_panel_bottom = standalone_box[1] + standalone_box[3]
    standalone_key_top = standalone_panel_bottom + 0.18
    standalone_height = standalone_key_top + 1.05
    _write_png(args.standalone_07_png, (PANEL_07,), [standalone_box], standalone_height, standalone_key_top, source_dir=source_dir)
    _write_pptx(args.standalone_07_pptx, (PANEL_07,), [standalone_box], standalone_height, standalone_key_top, source_dir=source_dir)
    print(f"已生成 {args.png}")
    print(f"已生成 {args.pptx}")
    print(f"已生成 {args.standalone_07_png}")
    print(f"已生成 {args.standalone_07_pptx}")


if __name__ == "__main__":
    main()
